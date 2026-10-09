"""Conferir a planilha do fornecedor antes de subir, e criar só o que foi marcado.

O PEDIDO DO DONO DO PRODUTO, nas duas metades: o veículo sem cadastro nasce
junto com a subida, E a conta aparece antes, com caixa de seleção marcada em
todos para se poder desmarcar. A segunda metade é o que torna a primeira segura
— medido contra o export de 08–09/2026, são 2.631 criações numa subida só.

O QUE ESTES TESTES PROTEGEM é a divisão: `conferir` lê, `ingerir` grava, e o
padrão de `ingerir` é não criar nada. Esse padrão não é timidez: a tela de
Calibração já tem um botão "Importar planilha" que chama `ingerir` direto, e um
padrão "criar tudo" faria quem clicasse ali criar 2.631 instituições sem ter
visto nada.
"""

import io

import pytest
from sqlalchemy import create_engine, func, select
from sqlalchemy.orm import Session

from app.banco.tabelas_score import Mencao, ScoreFonte
from app.banco.tabelas_stakeholders import Instituicao
from app.casos_de_uso import ingerir_mencoes
from app.dominio.texto import normalizar
from tests.test_e2e_postgres import URL

_engine = create_engine(URL, pool_pre_ping=True)


@pytest.fixture
def sessao():
    conexao = _engine.connect()
    transacao = conexao.begin()
    sessao = Session(bind=conexao, expire_on_commit=False)
    try:
        yield sessao
    finally:
        sessao.close()
        transacao.rollback()
        conexao.close()


#: As colunas que o cadastro da Clipei espera, na ordem em que o export as manda.
_CABECALHO = [
    "Data",
    "Classificação",
    "Aegea Tier",
    "Veículo",
    "Público-alvo",
    "Subcategoria",
    "Atributo",
    "ID",
    "Arquivo/Link",
    "Título",
    "Estado do Veículo",
    "Empresa",
    "Abrangência",
]


def _planilha(linhas: list[dict]) -> bytes:
    """Um `.xlsx` na forma que o cadastro da Clipei lê."""
    from openpyxl import Workbook

    pasta = Workbook()
    pasta.remove(pasta.active)
    aba = pasta.create_sheet("Clipping")
    aba.append(_CABECALHO)
    for linha in linhas:
        aba.append([linha.get(coluna, "") for coluna in _CABECALHO])
    arquivo = io.BytesIO()
    pasta.save(arquivo)
    return arquivo.getvalue()


def _linha(veiculo: str, **extra) -> dict:
    base = {
        "Data": "2026-08-10",
        "Classificação": "POSITIVA",
        "Aegea Tier": "Relevante",
        "Veículo": veiculo,
        "Público-alvo": "Opinião Pública",
        "Subcategoria": "Tarifa",
        "Atributo": "1. Governança",
        "ID": f"id-{veiculo}",
        "Arquivo/Link": "https://exemplo/1",
        "Título": "Uma matéria",
        "Estado do Veículo": "Santa Catarina",
        "Empresa": "Águas de Teste",
        "Abrangência": "Local",
    }
    base.update(extra)
    return base


@pytest.fixture
def clipei(sessao):
    return sessao.scalar(select(ScoreFonte).where(ScoreFonte.codigo == "clipei"))


def _quantos_veiculos(sessao) -> int:
    return sessao.scalar(
        select(func.count()).select_from(Instituicao).where(Instituicao.tipo == "veiculo")
    )


# ============================================= conferir: le, e NAO grava
def test_conferir_NAO_cria_veiculo(sessao, clipei):
    """A CONTA QUE A TELA MOSTRA não pode ser o fato consumado.

    Se a conferência criasse, o "mostre antes" do dono não existiria: a pessoa
    veria 2.631 e eles já estariam no banco.
    """
    antes = _quantos_veiculos(sessao)

    ingerir_mencoes.conferir(
        sessao, clipei, _planilha([_linha("Veículo Da Conferência 42")])
    )

    assert _quantos_veiculos(sessao) == antes


def test_conferir_NAO_grava_mencao(sessao, clipei):
    antes = sessao.scalar(select(func.count()).select_from(Mencao))

    ingerir_mencoes.conferir(sessao, clipei, _planilha([_linha("Veículo Qualquer 42")]))

    assert sessao.scalar(select(func.count()).select_from(Mencao)) == antes


def test_conferir_diz_o_que_seria_ingerido_em_cada_fonte_irma(sessao, clipei):
    """UMA LINHA POR FONTE, como a subida: o export da Clipei alimenta Imprensa
    e, recortado, Mercado. Conferir uma e subir as duas seria prometer uma coisa
    e fazer outra."""
    previsao, _ = ingerir_mencoes.conferir(
        sessao, clipei, _planilha([_linha("Veículo Irmãs 42")])
    )

    assert {r.fonte for r in previsao} == {"clipei", "clipei_investidores"}
    assert next(r for r in previsao if r.fonte == "clipei").ingeridas == 1


def test_conferir_lista_os_veiculos_que_nasceriam(sessao, clipei):
    previsao, reco = ingerir_mencoes.conferir(
        sessao,
        clipei,
        _planilha(
            [
                _linha("Veículo Novo A 42"),
                _linha("Veículo Novo A 42", ID="id-2"),
                _linha("Veículo Novo B 42"),
            ]
        ),
    )

    assert [v.nome for v in reco.novos] == ["Veículo Novo A 42", "Veículo Novo B 42"]
    # O MAIS CITADO PRIMEIRO: a tela mostra os primeiros, e 2.631 nomes não
    # cabem numa lista.
    assert [v.mencoes for v in reco.novos] == [2, 1]
    assert reco.novos[0].uf == "Santa Catarina"
    assert reco.novos[0].esfera == "Municipal"


def test_conferir_recusa_o_arquivo_que_nao_abre(sessao, clipei):
    """AS RECUSAS ESTRUTURAIS APARECEM NA CONFERÊNCIA, antes de a pessoa
    escolher nada. Conferir com um leitor diferente do que grava seria conferir
    outra coisa."""
    from app.dominio.erros import RegraViolada

    with pytest.raises(RegraViolada):
        ingerir_mencoes.conferir(sessao, clipei, b"isto nao e um xlsx")


# ============================================= ingerir: cria SO o autorizado
def test_ingerir_sem_autorizacao_NAO_cria_nada(sessao, clipei):
    """O PADRÃO É NÃO CRIAR, e isto protege a tela de Calibração.

    Ela tem um botão "Importar planilha" por fonte que chama `ingerir` direto.
    Com padrão "criar tudo", quem clicasse ali criaria 2.631 instituições sem
    ter visto nada — por uma mudança de assinatura.
    """
    antes = _quantos_veiculos(sessao)

    resumos = ingerir_mencoes.ingerir(
        sessao, clipei, _planilha([_linha("Veículo Não Autorizado 42")])
    )

    assert _quantos_veiculos(sessao) == antes
    assert all(r.veiculos_criados == 0 for r in resumos)
    assert all(r.mencoes_ligadas == 0 for r in resumos)


def test_ingerir_cria_so_os_nomes_marcados(sessao, clipei):
    conteudo = _planilha(
        [_linha("Veículo Marcado 42"), _linha("Veículo Desmarcado 42", ID="id-2")]
    )

    ingerir_mencoes.ingerir(sessao, clipei, conteudo, ["Veículo Marcado 42"])

    assert sessao.scalar(
        select(func.count())
        .select_from(Instituicao)
        .where(Instituicao.nome == "Veículo Marcado 42")
    )
    assert not sessao.scalar(
        select(func.count())
        .select_from(Instituicao)
        .where(Instituicao.nome == "Veículo Desmarcado 42")
    )


def test_a_mencao_fica_ligada_ao_veiculo_criado(sessao, clipei):
    """O PONTO DE TUDO ISSO: o `instituicao_id` deixa de ser nulo, e a menção
    passa a somar pela mesma régua do CRM."""
    conteudo = _planilha([_linha("Veículo Ligado 42")])

    resumos = ingerir_mencoes.ingerir(sessao, clipei, conteudo, ["Veículo Ligado 42"])

    veiculo = sessao.scalar(
        select(Instituicao).where(Instituicao.nome == "Veículo Ligado 42")
    )
    mencao = sessao.scalar(
        select(Mencao).where(Mencao.fonte_id == clipei.id, Mencao.veiculo == "Veículo Ligado 42")
    )
    assert mencao.instituicao_id == veiculo.id
    assert next(r for r in resumos if r.fonte == "clipei").mencoes_ligadas == 1


def test_a_mencao_liga_ao_veiculo_que_JA_estava_cadastrado(sessao, clipei):
    """Sem autorizar criação nenhuma: o que já existe liga sozinho, e é o que
    faz a cobertura subir sem precisar recadastrar."""
    sessao.add(
        Instituicao(
            nome="Veículo Já Existia 42",
            nome_normalizado=normalizar("Veículo Já Existia 42"),
            tipo="veiculo",
        )
    )
    sessao.flush()

    resumos = ingerir_mencoes.ingerir(
        sessao, clipei, _planilha([_linha("Veículo Já Existia 42")])
    )

    assert next(r for r in resumos if r.fonte == "clipei").mencoes_ligadas == 1
    assert next(r for r in resumos if r.fonte == "clipei").veiculos_criados == 0


def test_nome_autorizado_que_a_planilha_NAO_tem_e_ignorado(sessao, clipei):
    """A ROTA NÃO PODE SER UM JEITO DE CADASTRAR INSTITUIÇÃO ARBITRÁRIA.

    Os nomes vêm de um campo de formulário. Sem conferir contra a planilha,
    mandar "Banco Qualquer S/A" criaria um veículo que nenhum fornecedor
    mencionou.
    """
    ingerir_mencoes.ingerir(
        sessao,
        clipei,
        _planilha([_linha("Veículo Da Planilha 42")]),
        ["Veículo Inventado No Formulário 42"],
    )

    assert not sessao.scalar(
        select(func.count())
        .select_from(Instituicao)
        .where(Instituicao.nome == "Veículo Inventado No Formulário 42")
    )


def test_os_veiculos_nascem_UMA_vez_para_as_duas_fontes_irmas(sessao, clipei):
    """E O NÚMERO NÃO PODE SER SOMADO.

    As irmãs leem o mesmo arquivo e veriam os mesmos veículos. Criar por fonte
    faria a segunda achar tudo cadastrado (inofensivo) e contaria a criação duas
    vezes — a tela mentiria o número que a pessoa acabou de autorizar.
    """
    antes = _quantos_veiculos(sessao)

    resumos = ingerir_mencoes.ingerir(
        sessao,
        clipei,
        _planilha([_linha("Veículo Uma Vez 42", **{"Público-alvo": "Investidores"})]),
        ["Veículo Uma Vez 42"],
    )

    assert _quantos_veiculos(sessao) == antes + 1, "nasceu uma vez"
    assert len(resumos) == 2, "as duas fontes responderam"
    # O MESMO VALOR nos dois, e é da SUBIDA: somar daria 2 onde nasceu 1.
    assert {r.veiculos_criados for r in resumos} == {1}


def test_a_irma_com_recorte_VAZIO_nao_derruba_a_subida(sessao, clipei):
    """DEFEITO QUE VEM DO COMMIT ORIGINAL DA INGESTÃO, e que ia piorar.

    A lente Mercado é a Clipei recortada por público investidor. Um mês sem
    nenhuma menção de investidor é normal — e o erro "nenhuma linha virou
    menção" derrubava A SUBIDA INTEIRA: as 25.457 menções de Imprensa não
    entravam porque o recorte ficou vazio.

    E ia piorar: quando o critério do Mercado virar a subcategoria do veículo
    cadastrado, a primeira subida terá zero veículos classificados. O recorte
    nasce vazio por construção, e nenhuma carga da Clipei passaria.
    """
    # Só público geral: o recorte de investidores fica em zero.
    resumos = ingerir_mencoes.ingerir(
        sessao, clipei, _planilha([_linha("Veículo Sem Investidor 42")])
    )

    por_fonte = {r.fonte: r for r in resumos}
    assert por_fonte["clipei"].ingeridas == 1, "a fonte da subida entrou"
    assert por_fonte["clipei_investidores"].ingeridas == 0, "o recorte ficou vazio"
    assert por_fonte["clipei_investidores"].linhas == 1, "mas ele LEU a planilha"


def test_zero_na_fonte_da_subida_CONTINUA_sendo_erro(sessao, clipei):
    """O CONTRAPESO. Afrouxar o vazio da irmã não pode afrouxar o da fonte pela
    qual se subiu: ali zero significa que o arquivo não serve, e aceitar em
    silêncio apagaria o mês trocando-o por nada.
    """
    from app.dominio.erros import RegraViolada

    # Data que não é data: nenhuma linha vira menção em nenhuma fonte.
    with pytest.raises(RegraViolada, match="Nenhuma linha"):
        ingerir_mencoes.ingerir(
            sessao, clipei, _planilha([_linha("Veículo Sem Data 42", Data="ontem")])
        )


def _planilha_com_titulo(linhas: list[dict], titulo: list) -> bytes:
    """Um `.xlsx` com uma linha ACIMA do cabeçalho, como a Clipei manda."""
    from openpyxl import Workbook

    pasta = Workbook()
    pasta.remove(pasta.active)
    aba = pasta.create_sheet("Clipping")
    aba.append(titulo)
    aba.append(_CABECALHO)
    for linha in linhas:
        aba.append([linha.get(coluna, "") for coluna in _CABECALHO])
    arquivo = io.BytesIO()
    pasta.save(arquivo)
    return arquivo.getvalue()


def test_o_cabecalho_e_PROCURADO_e_nao_assumido_na_linha_1(sessao, clipei):
    """O EXPORT DA CLIPEI TEM UMA LINHA ACIMA DO CABEÇALHO.

    São três marcadores de nível (`N1`, `N2`, `N3`) sobre as colunas Atributo,
    Categoria e Subcategoria. O leitor assumia a linha 1 e recusava o arquivo
    inteiro com "a planilha não tem as colunas que o cadastro desta fonte
    espera" — a mensagem certa para um arquivo errado, e a errada para um
    arquivo CERTO com uma linha de título.

    Era isso que impedia a planilha de subir como o fornecedor a manda.
    """
    conteudo = _planilha_com_titulo(
        [_linha("Veículo Sob Título 42")], ["N2", "N3", "N1"]
    )

    previsao, _ = ingerir_mencoes.conferir(sessao, clipei, conteudo)

    assert next(r for r in previsao if r.fonte == "clipei").ingeridas == 1


def test_arquivo_que_REALMENTE_nao_tem_as_colunas_segue_recusado(sessao, clipei):
    """O CONTRAPESO. Procurar o cabeçalho não pode virar aceitar qualquer coisa:
    a recusa com o nome das colunas que faltam é o que diz à pessoa que ela
    pegou o arquivo errado.
    """
    from openpyxl import Workbook

    from app.dominio.erros import RegraViolada

    pasta = Workbook()
    pasta.remove(pasta.active)
    aba = pasta.create_sheet("Clipping")
    aba.append(["Assunto", "Grupo", "Outra Coisa"])
    aba.append(["x", "y", "z"])
    arquivo = io.BytesIO()
    pasta.save(arquivo)

    with pytest.raises(RegraViolada, match="não tem as colunas"):
        ingerir_mencoes.conferir(sessao, clipei, arquivo.getvalue())


def test_a_rota_de_conferencia_esta_isenta_do_teto_de_1_MB():
    """O DEFEITO QUE A PESSOA ENCONTROU AO USAR A TELA.

    Eu criei a rota e esqueci de isentá-la do teto genérico de 1 MB. O export da
    Clipei tem 3,8 MB, e o sintoma NÃO foi um 413 legível: o middleware recusa
    pelo `Content-Length` e fecha a conexão enquanto o navegador ainda envia,
    então o `fetch` falha em nível de rede e a tela mostrou "não foi possível
    falar com o servidor" — mandando procurar um backend derrubado quando o
    problema era o tamanho do arquivo.

    Este teste é estrutural de propósito: a próxima rota de upload que nascer
    sob este prefixo não vai descobrir isso pela tela de alguém.
    """
    from app.seguranca.protecao_http import _fora_do_limite_de_corpo

    assert _fora_do_limite_de_corpo("/api/score/fontes/clipei/conferencia")
    assert _fora_do_limite_de_corpo("/api/score/fontes/clipei/planilha")
    # E o que NÃO recebe arquivo continua com teto.
    assert not _fora_do_limite_de_corpo("/api/score/fontes")
