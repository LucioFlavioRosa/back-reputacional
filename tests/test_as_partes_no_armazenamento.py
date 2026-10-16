"""A montagem em partes contra o armazenamento DE VERDADE.

POR QUE ESTE ARQUIVO EXISTE, e é a correção de um ponto cego meu. Os testes da
subida em partes (`test_a_planilha_sobe_em_partes.py`) trocam o módulo `blob`
inteiro por um fake — é o certo lá, porque o que eles afirmam é a orquestração
da rota. Mas isso deixava a lógica de BLOCO sem rede nenhuma: o harness de
mutação mostrou que inverter a ordem da lista de blocos, ou tirar a largura fixa
do identificador, não quebrava teste nenhum.

Um fake prova o que eu ACREDITO que o Blob faz. Só o Blob prova o que ele faz.

PULA SEM AZURITE, como o resto da suíte pula sem Postgres: o CI não sobe
armazenamento, e um teste que falhasse ali viraria ruído que alguém desliga.
"""

from __future__ import annotations

from uuid import uuid4

import pytest

from app.armazenamento import blob
from app.configuracao import Configuracao, obter_configuracao

#: A cadeia pública do Azurite, a mesma do `docker-compose.pilha.yml`. O nome
#: `blob` resolve dentro da rede da pilha; fora dela o teste pula.
_AZURITE = (
    "DefaultEndpointsProtocol=http;AccountName=devstoreaccount1;"
    "AccountKey=Eby8vdM02xNOcqFlqUwJPLlmEtlCDXJ1OUzFT50uSRZ6IFsuFq2UVErCz4I6tq/"
    "K1SZFPTOtr/KBHBeksoGMGw==;BlobEndpoint=http://blob:10000/devstoreaccount1;"
)


@pytest.fixture
def armazenamento_real(monkeypatch):
    """Aponta o `blob` para o Azurite da pilha, num contêiner só deste teste."""
    padrao = obter_configuracao()
    como = Configuracao(
        **{**padrao.model_dump(), "blob_url": _AZURITE, "blob_contenedor": "testes-partes"}
    )
    monkeypatch.setattr(blob, "obter_configuracao", lambda: como)
    try:
        blob._contenedor()
    except Exception as erro:  # noqa: BLE001
        pytest.skip(f"sem Azurite alcançável: {type(erro).__name__}")
    return como


def _caminho() -> str:
    return blob.caminho_da_planilha_em_partes(fonte="clipei", planilha_id=uuid4())


# -- o identificador do bloco -------------------------------------------------


def test_o_identificador_do_bloco_tem_LARGURA_FIXA():
    """O Blob recusa a lista quando os blocos têm identificadores de tamanhos
    diferentes. `2` e `10` sem largura fixa produzem ids de tamanhos diferentes,
    e o erro aparece na montagem — longe de onde o id foi feito.

    É FUNÇÃO PURA, então este teste não precisa de armazenamento: ele prende a
    propriedade, e o teste de ida e volta abaixo prende a consequência.
    """
    tamanhos = {len(blob._id_do_bloco(i)) for i in (0, 1, 9, 10, 99, 100, 999, 1000)}

    assert len(tamanhos) == 1, f"ids de tamanhos diferentes: {tamanhos}"


def test_o_identificador_do_bloco_e_DIFERENTE_por_indice():
    """Dois índices com o mesmo id fariam uma parte sobrescrever a outra, e o
    arquivo montado sairia com um pedaço repetido e outro faltando."""
    ids = [blob._id_do_bloco(i) for i in range(50)]

    assert len(set(ids)) == 50


# -- a ida e volta ------------------------------------------------------------


def test_as_partes_montam_NA_ORDEM_DA_LISTA_e_nao_da_chegada(armazenamento_real):
    """O cliente pode subir em paralelo. A ordem é a dos índices.

    ESTE É O TESTE QUE FALTAVA: ele falha se `juntar_partes` montar na ordem de
    chegada, e foi inverter essa ordem que o harness de mutação não conseguiu
    detectar com o fake no lugar.
    """
    caminho = _caminho()
    pedacos = [b"primeiro-", b"segundo-", b"terceiro"]

    #: DE TRÁS PARA FRENTE de propósito.
    for indice in reversed(range(len(pedacos))):
        blob.guardar_parte(caminho, indice, pedacos[indice])
    blob.juntar_partes(caminho, len(pedacos), "application/octet-stream")

    assert blob.ler(caminho) == b"".join(pedacos)
    blob.apagar(caminho)


def test_a_parte_NAO_EXISTE_para_leitura_antes_da_montagem(armazenamento_real):
    """Bloco não commitado não é arquivo.

    É o que faz uma subida interrompida não deixar nada para alguém ler — e o
    que torna desnecessário limpar as partes de uma subida abandonada.
    """
    from app.dominio.erros import NaoEncontrado

    caminho = _caminho()
    blob.guardar_parte(caminho, 0, b"um pedaco qualquer")

    with pytest.raises(NaoEncontrado):
        blob.ler(caminho)


def test_a_parte_reenviada_SUBSTITUI_no_armazenamento(armazenamento_real):
    """`PUT` promete idempotência por índice, e é o que deixa o cliente repetir
    a parte que falhou sem recomeçar."""
    caminho = _caminho()
    blob.guardar_parte(caminho, 0, b"a versao que falhou")
    blob.guardar_parte(caminho, 0, b"a versao certa....")
    blob.juntar_partes(caminho, 1, "application/octet-stream")

    assert blob.ler(caminho) == b"a versao certa...."
    blob.apagar(caminho)


def test_montar_com_uma_parte_FALTANDO_falha(armazenamento_real):
    """O serviço recusa a lista inteira, e é o desfecho certo: montar sem um
    pedaço daria um `.xlsx` truncado que só falharia na leitura."""
    caminho = _caminho()
    blob.guardar_parte(caminho, 0, b"tem")
    blob.guardar_parte(caminho, 2, b"falta a do meio")

    with pytest.raises(Exception):  # noqa: B017 - o SDK escolhe o tipo
        blob.juntar_partes(caminho, 3, "application/octet-stream")

    #: E NADA FOI PUBLICADO: a recusa não pode deixar arquivo pela metade.
    from app.dominio.erros import NaoEncontrado

    with pytest.raises(NaoEncontrado):
        blob.ler(caminho)


def test_o_arquivo_montado_PRESERVA_os_bytes_de_um_xlsx(armazenamento_real):
    """Um `.xlsx` é um zip, e zip não perdoa um byte trocado.

    Afirmar sobre texto curto não provaria isso: o teste monta um arquivo de
    verdade, em pedaços que cortam no meio da estrutura do zip.
    """
    import io

    from openpyxl import Workbook

    pasta = Workbook()
    pasta.active.append(["Data", "Veículo"])
    pasta.active.append(["2026-08-10", "Jornal de Teste"])
    saida = io.BytesIO()
    pasta.save(saida)
    conteudo = saida.getvalue()

    caminho = _caminho()
    pedacos = [conteudo[i : i + 700] for i in range(0, len(conteudo), 700)]
    assert len(pedacos) > 2, "o teste precisa de cortes no meio do zip"
    for indice, pedaco in enumerate(pedacos):
        blob.guardar_parte(caminho, indice, pedaco)
    blob.juntar_partes(caminho, len(pedacos), "application/octet-stream")

    montado = blob.ler(caminho)
    assert montado == conteudo
    #: E ABRE: o teste de igualdade já basta, mas abrir prende o que importa
    #: para quem lê — que o arquivo montado É uma planilha.
    from openpyxl import load_workbook

    assert load_workbook(io.BytesIO(montado)).active["B2"].value == "Jornal de Teste"
    blob.apagar(caminho)
