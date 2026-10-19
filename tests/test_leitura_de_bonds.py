"""O leitor do export BDH: as regras que, erradas, gravam dado errado calado.

POR QUE ESTE ARQUIVO É O MAIOR DO BLOCO. O leitor é o único lugar da pilha de
bonds onde um defeito não falha alto: a migration levanta exceção, a rota
devolve erro, o gráfico não desenha. Aqui, `0` lido como preço grava 1.363 dias
de cotação zero para o `AEGEBZ36` e o gráfico fica com um precipício — sem
nenhum erro em lugar nenhum.

AS LINHAS SÃO TUPLAS, não um `.xlsx`. O módulo é domínio puro de propósito
(mesma decisão de `importacao_de_agendas`), e é isso que permite montar aqui
cada caso-limite em três linhas em vez de gerar uma pasta de trabalho por teste.

O ÚLTIMO TESTE LÊ O ARQUIVO DE VERDADE, e é o que impede os de cima de
concordarem entre si sobre um formato que não existe. Ele é marcado e PULA
quando o arquivo não está na máquina — o `.xlsx` vem de terminal Bloomberg e
não entra no repositório.
"""

from __future__ import annotations

import os
from datetime import date
from decimal import Decimal
from pathlib import Path

import pytest

from app.dominio.leitura_de_bonds import ler

#: O cabeçalho mínimo que o formato exige, nas cinco primeiras linhas.
#: Colunas: 0..4 catálogo, 5 vazia, 6 DATES, 7+ os pares.
CABECALHO = [
    ["Bond Description", None, None, "ID", "AMT_OUTSTANDING", None, None, "AAA29", "AAA29", "BBB31", "BBB31"],
    [None, None, None, None, None, None, None, "Preço", "Yield", "Preço", "Yield"],
    [None, None, None, None, None, None, None, None, None, None, None],
    [None, None, None, None, None, None, None, "ID-A", "ID-A", "ID-B", "ID-B"],
    [None, None, None, None, None, None, "DATES", "px_last(", "yield(", "px_last(", "yield("],
]


def _com_catalogo(*papeis):
    """O cabeçalho com o catálogo `A..E` preenchido a partir da linha 2."""
    linhas = [list(l) for l in CABECALHO]
    for posicao, (descricao, id_bb, ticker, amt) in enumerate(papeis, start=1):
        while len(linhas) <= posicao:
            linhas.append([None] * 11)
        linhas[posicao][0] = descricao
        linhas[posicao][1] = id_bb
        linhas[posicao][2] = ticker
        linhas[posicao][3] = id_bb
        linhas[posicao][4] = amt
    return linhas


def _dia(linhas, dia, *valores):
    """Acrescenta uma linha de série: data na coluna 6, valores de 7 em diante."""
    linha = [None] * 7 + list(valores)
    linha[6] = dia
    linhas.append(linha)
    return linhas


# -- a regra do zero -----------------------------------------------------------


def test_o_ZERO_vira_nulo_e_nao_preco_zero():
    """A regra que mais importa.

    As datas são de CALENDÁRIO: sábado, domingo, feriado e todo dia anterior à
    emissão do papel vêm `0`. Gravar o zero não é só feio no gráfico — a média e
    a variação do período passariam a incluir dias que não existiram.
    """
    linhas = _com_catalogo(("AAA 1 01/01/29", "ID-A", "AAA29", 100))
    _dia(linhas, date(2026, 1, 5), 0, 0, 95.5, 9.1)

    leitura = ler(linhas)

    por_ticker = {c.ticker: c for c in leitura.cotacoes}
    assert "AAA29" not in por_ticker, "o par com 0/0 não devia gerar linha"
    assert por_ticker["BBB31"].preco_pct_par == Decimal("95.5")


def test_zero_como_texto_e_como_decimal_tambem_viram_nulo():
    """`0`, `0.0`, `"0"` e `"0,00"` são o mesmo "sem cotação".

    A planilha vem do Bloomberg, mas passa por Excel — e Excel devolve o mesmo
    zero como inteiro, como float ou como texto dependendo da formatação da
    célula. Tratar só o `int` deixaria os outros três virarem preço zero.
    """
    linhas = _com_catalogo(("AAA 1 01/01/29", "ID-A", "AAA29", 100))
    _dia(linhas, date(2026, 1, 5), "0", "0,00", Decimal("0.0"), 0.0)

    assert ler(linhas).cotacoes == []


def test_preco_valido_proximo_de_zero_NAO_e_descartado():
    """O controle negativo da regra do zero.

    `0,0001` é cotação (um papel em default profundo), e tratar "perto de zero"
    como ausência jogaria fora exatamente o dado mais dramático. A regra é
    IGUAL A ZERO, não "pequeno".
    """
    linhas = _com_catalogo(("AAA 1 01/01/29", "ID-A", "AAA29", 100))
    _dia(linhas, date(2026, 1, 5), Decimal("0.0001"), 0, 0, 0)

    assert ler(linhas).cotacoes[0].preco_pct_par == Decimal("0.0001")


# -- a coluna das datas --------------------------------------------------------


def test_a_coluna_das_datas_e_PROCURADA_e_nao_contada():
    """Uma coluna a mais à esquerda não pode deslocar preço para yield."""
    linhas = [[None] + list(l) for l in CABECALHO]
    linhas[1][3] = "AAA29"  # o catálogo também andou uma casa
    linhas[1][1] = "AAA 1 01/01/29"
    serie = [None] * 8 + [95.5, 9.1, 80.0, 7.0]
    serie[7] = date(2026, 1, 5)
    linhas.append(serie)

    leitura = ler(linhas)

    por_ticker = {c.ticker: c for c in leitura.cotacoes}
    assert por_ticker["AAA29"].preco_pct_par == Decimal("95.5")
    assert por_ticker["AAA29"].yield_pct == Decimal("9.1")


def test_sem_a_palavra_DATES_devolve_o_catalogo_e_nenhuma_cotacao():
    """Arquivo da forma errada não levanta — quem chama vira isso em mensagem."""
    linhas = _com_catalogo(("AAA 1 01/01/29", "ID-A", "AAA29", 100))
    linhas[4][6] = "ALGO OUTRO"
    _dia(linhas, date(2026, 1, 5), 95.5, 9.1, 80.0, 7.0)

    leitura = ler(linhas)

    assert leitura.cotacoes == []
    assert [p.ticker for p in leitura.papeis] == ["AAA29"]


# -- preço e yield na mesma linha ---------------------------------------------


def test_preco_e_yield_do_mesmo_dia_viram_UMA_cotacao():
    """Duas colunas, uma linha — é o que a chave `(ticker, data)` exige."""
    linhas = _com_catalogo(("AAA 1 01/01/29", "ID-A", "AAA29", 100))
    _dia(linhas, date(2026, 1, 5), 95.5, 9.1, 0, 0)

    cotacoes = ler(linhas).cotacoes

    assert len(cotacoes) == 1
    assert (cotacoes[0].preco_pct_par, cotacoes[0].yield_pct) == (
        Decimal("95.5"),
        Decimal("9.1"),
    )


def test_so_o_yield_tambem_grava_a_linha():
    """Medido no arquivo real: vários papéis têm yield desde 2022 e preço só
    desde 10/2023. Exigir os dois perderia 639 dias de yield do `ELEBRA30`."""
    linhas = _com_catalogo(("AAA 1 01/01/29", "ID-A", "AAA29", 100))
    _dia(linhas, date(2026, 1, 5), 0, 9.1, 0, 0)

    cotacoes = ler(linhas).cotacoes

    assert len(cotacoes) == 1
    assert cotacoes[0].preco_pct_par is None
    assert cotacoes[0].yield_pct == Decimal("9.1")


# -- o que o leitor reporta ----------------------------------------------------


def test_par_sem_valor_nenhum_e_reportado_e_nao_gravado():
    """O caso `AEGEBZ29`: o papel que o mockup mostra e o arquivo não tem."""
    linhas = _com_catalogo(
        ("AAA 1 01/01/29", "ID-A", "AAA29", 100),
        ("BBB 2 01/01/31", "ID-B", "BBB31", 200),
    )
    _dia(linhas, date(2026, 1, 5), 0, 0, 80.0, 7.0)

    leitura = ler(linhas)

    assert leitura.tickers_sem_valor == ["AAA29"]
    assert {c.ticker for c in leitura.cotacoes} == {"BBB31"}


def test_ticker_de_serie_fora_do_catalogo_e_reportado():
    """Papel novo emitido chega na série antes de alguém cadastrá-lo."""
    linhas = _com_catalogo(("AAA 1 01/01/29", "ID-A", "AAA29", 100))
    _dia(linhas, date(2026, 1, 5), 95.5, 9.1, 80.0, 7.0)

    assert ler(linhas).tickers_fora_do_catalogo == ["BBB31"]


def test_o_catalogo_acaba_na_primeira_linha_sem_ticker():
    linhas = _com_catalogo(
        ("AAA 1 01/01/29", "ID-A", "AAA29", 100),
        ("BBB 2 01/01/31", "ID-B", "BBB31", 200),
    )
    linhas.append([None] * 11)
    linhas.append(["LIXO", None, None, None, None] + [None] * 6)

    assert [p.ticker for p in ler(linhas).papeis] == ["AAA29", "BBB31"]


def test_o_montante_zero_do_catalogo_NAO_vira_nulo():
    """A regra do zero é da SÉRIE, não do catálogo.

    Montante em circulação de zero é um fato (papel inteiramente recomprado), e
    aplicar a regra da série aqui apagaria esse fato.
    """
    linhas = _com_catalogo(("AAA 1 01/01/29", "ID-A", "AAA29", 0))

    assert ler(linhas).papeis[0].amt_outstanding == Decimal("0")


# -- o fim do arquivo ----------------------------------------------------------


def test_linha_sem_data_ENCERRA_a_leitura():
    """O arquivo real declara 54.254 linhas e só 1.750 têm data.

    E o ponto não é só desempenho: uma data perdida no rodapé fantasma seria
    lida como série válida e esticaria o eixo do gráfico por anos vazios.
    """
    linhas = _com_catalogo(("AAA 1 01/01/29", "ID-A", "AAA29", 100))
    _dia(linhas, date(2026, 1, 5), 95.5, 9.1, 0, 0)
    linhas.append([None] * 11)
    _dia(linhas, date(2030, 1, 1), 50.0, 20.0, 0, 0)

    leitura = ler(linhas)

    assert [c.data for c in leitura.cotacoes] == [date(2026, 1, 5)]
    assert leitura.ultima_data == date(2026, 1, 5)


def test_arquivo_curto_demais_nao_levanta():
    assert ler([]).cotacoes == []
    assert ler(CABECALHO).cotacoes == []


def test_as_datas_extremas_saem_na_leitura():
    linhas = _com_catalogo(("AAA 1 01/01/29", "ID-A", "AAA29", 100))
    _dia(linhas, date(2026, 3, 2), 95.5, 9.1, 0, 0)
    _dia(linhas, date(2026, 1, 5), 94.0, 9.3, 0, 0)

    leitura = ler(linhas)

    assert (leitura.primeira_data, leitura.ultima_data) == (
        date(2026, 1, 5),
        date(2026, 3, 2),
    )


# -- o arquivo de verdade ------------------------------------------------------

#: O `.xlsx` vem de terminal Bloomberg e NÃO entra no repositório — a licença de
#: redistribuição é ponto aberto da especificação. O teste lê do disco quando
#: está lá e pula quando não está, para o CI não depender dele.
ARQUIVO_REAL = Path(
    os.environ.get(
        "PLANILHA_DE_BONDS",
        Path.home()
        / "OneDrive - Peers Consulting"
        / "Área de Trabalho"
        / "aegea"
        / "painel reputacional"
        / "Aegea_Bonds_e_Reputacao"
        / "bonds-reputacao"
        / "Cópia de BASE DE DADOS_HISTORICO BONDS_Yieldv1.xlsx",
    )
)


@pytest.mark.skipif(not ARQUIVO_REAL.exists(), reason="planilha de bonds fora da máquina")
def test_o_arquivo_real_bate_com_o_que_foi_medido():
    """O teste que impede os de cima de concordarem sobre um formato fictício.

    Os números são a medição de 16/10/2026. Se a planilha da semana mudar de
    forma, é AQUI que se descobre — e não no gráfico.
    """
    from openpyxl import load_workbook

    pasta = load_workbook(ARQUIVO_REAL, read_only=True, data_only=True)
    folha = pasta[pasta.sheetnames[0]]
    linhas = [linha for linha in folha.iter_rows(values_only=True)]
    pasta.close()

    leitura = ler(linhas)

    assert len(leitura.papeis) == 31
    assert leitura.primeira_data == date(2022, 1, 1)
    assert leitura.ultima_data == date(2026, 10, 16)
    #: O `AEGEBZ29` é o papel do mockup que o arquivo não tem.
    assert leitura.tickers_sem_valor == ["AEGEBZ29"]
    #: O catálogo do arquivo cobre toda coluna de série — nenhum ticker órfão.
    assert leitura.tickers_fora_do_catalogo == []
    #: Nenhuma cotação com os dois lados nulos: seria linha sem informação.
    assert not [
        c for c in leitura.cotacoes if c.preco_pct_par is None and c.yield_pct is None
    ]
    #: E a Aegea tem o 2031 desde 28/09/2023 (o yield vem antes do preço).
    do_2031 = [c for c in leitura.cotacoes if c.ticker == "AEGEBZ31"]
    assert min(c.data for c in do_2031) == date(2023, 9, 28)
