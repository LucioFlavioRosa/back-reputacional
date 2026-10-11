"""A pessoa que aparece em duas instituições não barra mais a planilha.

O CASO REAL, e ele veio de uma subida de verdade. O dono do produto subiu
`modelo-de-agendas-relações-com-investidores.xlsx` na pilha local e recebeu:

    "Artur Santa, Michelle Nasser Vaz Lauande, Nikolau Muller, Paulo Hanaoka
    aparece em agendas de instituições diferentes, e um interlocutor pertence a
    uma só. Cadastre-o pela tela de Administração e aponte para ele na
    conferência."

E a correção que ele deu: "nesse caso a mensagem está correta, mas não era para
impedir subir: era transformar isso em um alerta para revisão na aba correspondente
de cadastro".

MEDIDO NO ARQUIVO DELE: 304 linhas, 321 pessoas distintas, e EXATAMENTE QUATRO em
duas instituições cada. Trezentas agendas barradas por quatro nomes.

    Artur Santa ................. Banco XP (114) e XP Asset (245)
    Nikolau Muller .............. JGP Gestão (56) e Leto Capital (289)
    Michelle Nasser Vaz Lauande   Itaú International PB (252) e Santander (255)
    Paulo Hanaoka ............... Itaú International PB (252) e Santander (255)

OS QUATRO SÃO DADO REAL, não erro de digitação: o primeiro é o mesmo grupo em dois
braços, o segundo parece mudança de casa, e os dois últimos estão nas MESMAS duas
linhas — assinatura de reunião conjunta ou de linha copiada. Nenhum deles é uma
pergunta que a máquina possa responder, e nenhum justifica recusar o arquivo.

DOIS CADASTROS É O QUE O MODELO PREVÊ: a `0002` cria
`unique (nome_normalizado, instituicao_id)` com o comentário "Mesmo nome em
instituições diferentes são pessoas diferentes". A trava nascia de `de_quem` ser um
dicionário nome → UMA instituição: o caso não caBIA na estrutura, e levantar
`RegraViolada` era a saída honesta de quem não podia representá-lo.
"""

from __future__ import annotations

import io
from datetime import date

import pytest
from openpyxl import load_workbook
from sqlalchemy import create_engine, func, select
from sqlalchemy.orm import Session

from app.banco.tabelas_stakeholders import Instituicao, Interlocutor
from app.casos_de_uso import importar_agendas, modelo_de_importacao
from app.dominio.texto import normalizar
from main import app
from tests.test_e2e_postgres import URL
from tests.test_importacoes_api import TIPO_XLSX, _cliente

_engine = create_engine(URL, pool_pre_ping=True)


@pytest.fixture
def sessao():
    conexao = _engine.connect()
    transacao = conexao.begin()
    sessao = Session(bind=conexao, expire_on_commit=False)
    try:
        yield sessao
    finally:
        sessao.close()
        transacao.rollback()
        conexao.close()


@pytest.fixture
def cliente_admin(sessao):
    try:
        yield _cliente(sessao, "plataforma_edicao")
    finally:
        app.dependency_overrides.clear()


def _duas_agendas_com_a_mesma_pessoa(sessao) -> bytes:
    """Duas linhas, duas instituições, a MESMA pessoa nova — o caso do arquivo real."""
    primeira = Instituicao(
        nome="Banco Zz51",
        nome_normalizado=normalizar("Banco Zz51"),
        tipo="investidor",
        uf="SP",
    )
    segunda = Instituicao(
        nome="Asset Zz51",
        nome_normalizado=normalizar("Asset Zz51"),
        tipo="investidor",
        uf="SP",
    )
    sessao.add_all([primeira, segunda])
    sessao.flush()

    modelo = modelo_de_importacao.gerar(importar_agendas.vocabularios(sessao))
    pasta = load_workbook(io.BytesIO(modelo))
    folha = pasta["Agendas"]
    cabecalho = [celula.value for celula in next(folha.iter_rows())]
    for instituicao in (primeira.nome, segunda.nome):
        folha.append(
            [
                {
                    "Data": date(2026, 9, 25),
                    "Instituição": instituicao,
                    "UF": "SP",
                    "Interlocutor 1": "Analista Dos Dois Zz51",
                }.get(coluna)
                for coluna in cabecalho
            ]
        )
    saida = io.BytesIO()
    pasta.save(saida)
    return saida.getvalue()


def test_a_planilha_SOBE_e_a_confirmacao_cria_um_cadastro_por_instituicao(
    cliente_admin, sessao
):
    """O arquivo entra, e a pessoa nasce duas vezes — uma por instituição.

    ERA `RegraViolada` NA CONFIRMAÇÃO: a conferência dizia que estava tudo pronto e
    o último passo recusava o arquivo inteiro. É o pior lugar para uma recusa —
    depois de a pessoa ter conferido tudo.
    """
    criada = cliente_admin.post(
        "/api/importacoes",
        files={"arquivo": ("ri.xlsx", _duas_agendas_com_a_mesma_pessoa(sessao), TIPO_XLSX)},
    ).json()
    assert criada["pendencias"] == 0, criada["grupos"]

    confirmacao = cliente_admin.post(f"/api/importacoes/{criada['id']}/confirmacao")
    assert confirmacao.status_code == 201, confirmacao.text
    assert confirmacao.json()["criadas"] == 2, confirmacao.text

    nascidos = sessao.scalars(
        select(Interlocutor).where(
            Interlocutor.nome_normalizado == normalizar("Analista Dos Dois Zz51")
        )
    ).all()
    #: DOIS CADASTROS, um por instituição — o que a `0002` chama de "pessoas
    #: diferentes".
    assert len(nascidos) == 2, [n.instituicao_id for n in nascidos]
    assert len({n.instituicao_id for n in nascidos}) == 2


def test_os_DOIS_nascem_pendentes_COM_O_MOTIVO(cliente_admin, sessao):
    """O alerta para revisão na aba de cadastro, que é o pedido.

    SEM O MOTIVO a aba lista dois cadastros com o mesmo nome e ninguém sabe por que
    eles estão lá — e a revisão vira adivinhação. A frase diz as três saídas
    possíveis: a mesma pessoa, mudança de casa, ou homônimas.
    """
    criada = cliente_admin.post(
        "/api/importacoes",
        files={"arquivo": ("ri.xlsx", _duas_agendas_com_a_mesma_pessoa(sessao), TIPO_XLSX)},
    ).json()
    cliente_admin.post(f"/api/importacoes/{criada['id']}/confirmacao")

    nascidos = sessao.scalars(
        select(Interlocutor).where(
            Interlocutor.nome_normalizado == normalizar("Analista Dos Dois Zz51")
        )
    ).all()
    for nascido in nascidos:
        assert nascido.pendente_de_revisao is True
        assert nascido.motivo_da_revisao is not None
        assert "2 instituições" in nascido.motivo_da_revisao


def test_a_CONFERENCIA_avisa_antes_de_confirmar(cliente_admin, sessao):
    """A conferência promete o que vai acontecer, e isto é o que ela vai fazer.

    O alerta na aba de cadastro acontece DEPOIS da confirmação. A conferência tem
    de dizer antes: "vou criar dois cadastros para esta pessoa" é exatamente o tipo
    de coisa que ninguém quer descobrir depois.
    """
    criada = cliente_admin.post(
        "/api/importacoes",
        files={"arquivo": ("ri.xlsx", _duas_agendas_com_a_mesma_pessoa(sessao), TIPO_XLSX)},
    ).json()

    mensagens = [
        d["mensagem"]
        for linha in criada["linhas"]
        for d in linha["divergencias"]
    ]
    avisos = [m for m in mensagens if "aparece em 2 instituições" in m]
    #: NAS DUAS LINHAS, porque quem olha uma linha precisa ver o aviso dela.
    assert len(avisos) == 2, mensagens
    #: E DIZ QUAIS SÃO AS INSTITUIÇÕES, senão a pessoa tem de procurar.
    assert "Banco Zz51" in avisos[0] and "Asset Zz51" in avisos[0], avisos[0]
    #: SEM TRAVAR — é o pedido inteiro.
    assert all(
        not d["trava"]
        for linha in criada["linhas"]
        for d in linha["divergencias"]
    ), mensagens


def test_a_pessoa_QUE_JA_EXISTE_em_outra_instituicao_segue_de_fora_da_agenda(
    cliente_admin, sessao
):
    """O outro caso, e ele NÃO mudou: quem já está no cadastro tem instituição.

    A regra 4 do domínio não aceita participante de outra instituição, e a agenda
    entra SEM ele com um aviso — é a mudança de 11/10/2026 mais cedo. Este teste
    está aqui para a correção de hoje não ter apagado aquela: são dois casos
    diferentes, e o nome novo é o único que vira dois cadastros.
    """
    camara = Instituicao(
        nome="Camara Zz52",
        nome_normalizado=normalizar("Camara Zz52"),
        tipo="orgao",
        uf="SP",
    )
    outra = Instituicao(
        nome="Fundo Zz52",
        nome_normalizado=normalizar("Fundo Zz52"),
        tipo="investidor",
        uf="SP",
    )
    sessao.add_all([camara, outra])
    sessao.flush()
    ja = Interlocutor(
        nome="Pessoa Da Camara Zz52",
        nome_normalizado=normalizar("Pessoa Da Camara Zz52"),
        instituicao_id=camara.id,
    )
    sessao.add(ja)
    sessao.flush()

    modelo = modelo_de_importacao.gerar(importar_agendas.vocabularios(sessao))
    pasta = load_workbook(io.BytesIO(modelo))
    folha = pasta["Agendas"]
    cabecalho = [celula.value for celula in next(folha.iter_rows())]
    folha.append(
        [
            {
                "Data": date(2026, 9, 25),
                "Instituição": outra.nome,
                "UF": "SP",
                "Interlocutor 1": ja.nome,
            }.get(coluna)
            for coluna in cabecalho
        ]
    )
    saida = io.BytesIO()
    pasta.save(saida)

    criada = cliente_admin.post(
        "/api/importacoes", files={"arquivo": ("ri.xlsx", saida.getvalue(), TIPO_XLSX)}
    ).json()
    assert criada["pendencias"] == 0, criada["grupos"]
    mensagens = " ".join(
        d["mensagem"] for linha in criada["linhas"] for d in linha["divergencias"]
    )
    assert "SEM esta pessoa" in mensagens, mensagens

    #: E NÃO NASCEU UM SEGUNDO CADASTRO DELA: ela já existe, então não está em
    #: `a_criar` — o recorte por `a_nascer` é quem garante.
    assert (
        sessao.scalar(
            select(func.count())
            .select_from(Interlocutor)
            .where(Interlocutor.nome_normalizado == normalizar(ja.nome))
        )
        == 1
    )


def _duas_grafias_do_mesmo_nome(sessao) -> bytes:
    """Duas linhas, a MESMA instituição, o mesmo nome com e sem acento.

    É o caso real: a planilha de RI traz "Renan Tiburcio" e "Renan Tibúrcio" no BTG
    Pactual, "Nathalia Maximo" e "Nathália Máximo" na XP Asset, e "Barbara
    Halberstadt" e "Bárbara Halberstadt" no Deutsche Bank.
    """
    banco = Instituicao(
        nome="Banco Zz61",
        nome_normalizado=normalizar("Banco Zz61"),
        tipo="investidor",
        uf="SP",
    )
    sessao.add(banco)
    sessao.flush()

    modelo = modelo_de_importacao.gerar(importar_agendas.vocabularios(sessao))
    pasta = load_workbook(io.BytesIO(modelo))
    folha = pasta["Agendas"]
    cabecalho = [celula.value for celula in next(folha.iter_rows())]
    for pessoa in ("Renan Tiburcio Zz61", "Renan Tibúrcio Zz61"):
        folha.append(
            [
                {
                    "Data": date(2026, 9, 25),
                    "Instituição": banco.nome,
                    "UF": "SP",
                    "Interlocutor 1": pessoa,
                }.get(coluna)
                for coluna in cabecalho
            ]
        )
    saida = io.BytesIO()
    pasta.save(saida)
    return saida.getvalue()


def test_duas_grafias_do_mesmo_nome_CRIAM_UM_cadastro(cliente_admin, sessao):
    """O 422 real que isto conserta, e ele veio de uma subida de verdade.

    A conferência dizia zero pendências e a confirmação devolvia: "já existe uma
    pessoa chamada 'Renan Tiburcio' nesta instituição — pode estar desativada".
    E ELA NÃO EXISTIA NO BANCO: medido, zero linhas. A colisão era dentro da própria
    planilha, que traz o nome com e sem acento.

    A CAUSA, E A IRONIA DELA: `a_criar` era indexado pelo valor CRU, e a unicidade
    do banco é `(nome_normalizado, instituicao_id)`. Duas grafias viravam duas
    criações, que colidem. E `normalizar` existe PRECISAMENTE para isso — o
    docstring dela diz que serve "para a importação deduplicar 'Radamés' e
    'Radames' antes de criar as chaves estrangeiras".
    """
    criada = cliente_admin.post(
        "/api/importacoes",
        files={"arquivo": ("ri.xlsx", _duas_grafias_do_mesmo_nome(sessao), TIPO_XLSX)},
    ).json()
    assert criada["pendencias"] == 0, criada["grupos"]

    confirmacao = cliente_admin.post(f"/api/importacoes/{criada['id']}/confirmacao")
    assert confirmacao.status_code == 201, confirmacao.text

    #: UM CADASTRO, com a PRIMEIRA grafia do arquivo.
    nascidos = sessao.scalars(
        select(Interlocutor).where(
            Interlocutor.nome_normalizado == normalizar("Renan Tiburcio Zz61")
        )
    ).all()
    assert len(nascidos) == 1, [n.nome for n in nascidos]
    assert nascidos[0].nome == "Renan Tiburcio Zz61"


def test_a_conferencia_AVISA_das_duas_grafias(cliente_admin, sessao):
    """A deduplicação é silenciosa, e este aviso é o que a torna honesta.

    O cadastro nasce com uma das grafias e ninguém ficaria sabendo que havia outra.
    "Vou cadastrar como 'Renan Tiburcio', não 'Renan Tibúrcio'" é uma escolha que a
    pessoa tem direito de ver — e de corrigir na planilha, se a errada for a que
    ficou.
    """
    criada = cliente_admin.post(
        "/api/importacoes",
        files={"arquivo": ("ri.xlsx", _duas_grafias_do_mesmo_nome(sessao), TIPO_XLSX)},
    ).json()

    avisos = [
        d["mensagem"]
        for linha in criada["linhas"]
        for d in linha["divergencias"]
        if "escrito de" in d["mensagem"]
    ]
    assert avisos, [
        d["mensagem"] for linha in criada["linhas"] for d in linha["divergencias"]
    ]
    #: DIZ AS DUAS FORMAS E QUAL VAI FICAR.
    assert "Renan Tiburcio Zz61" in avisos[0]
    assert "Renan Tibúrcio Zz61" in avisos[0]
    #: E NÃO TRAVA.
    assert criada["pendencias"] == 0
