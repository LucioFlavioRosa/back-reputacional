"""As rotas da importação: baixar o modelo, subir o arquivo, retomar a conferência.

ESCONDER O BOTÃO É CONVENIÊNCIA, NUNCA CONTROLE. A tela só oferece a importação a
quem administra cadastros, e isso não protege nada: um `curl` basta. A guarda é a
dependência do router, e a âncora estrutural deste arquivo é o que garante que
uma rota nova sob este prefixo nasça protegida em vez de depender de alguém
lembrar.
"""

import io

import pytest
from fastapi.testclient import TestClient
from sqlalchemy import create_engine, select
from sqlalchemy.orm import Session

from app.banco.sessao import obter_sessao
from app.banco.tabelas_catalogo import Clima, FormatoInteracao
from app.banco.tabelas_stakeholders import Instituicao, Interlocutor, PessoaAegea
from app.casos_de_uso import importar_agendas, modelo_de_importacao
from app.dominio.importacao_de_agendas import (
    COLUNA_DE_REPETICAO,
    ROTULO_DO_VOCABULARIO,
    VALOR_DA_REPETICAO,
    colunas_do_cadastro,
)
from app.dominio.texto import normalizar
from main import app
from tests.test_e2e_postgres import URL

TIPO_XLSX = "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet"

_engine = create_engine(URL, pool_pre_ping=True)


@pytest.fixture
def sessao():
    conexao = _engine.connect()
    transacao = conexao.begin()
    sessao = Session(bind=conexao, expire_on_commit=False)
    try:
        yield sessao
    finally:
        sessao.close()
        transacao.rollback()
        conexao.close()


def _cliente(sessao, perfil: str) -> TestClient:
    """Um cliente logado com o perfil pedido.

    SOBRESCREVE `obter_configuracao()` VIA `app.dependency_overrides`, e não
    `monkeypatch.setattr`: o `Depends(obter_configuracao)` já foi resolvido, na
    importação do módulo, com uma referência direta à função original — trocar o
    nome no módulo depois não alcança quem já guardou o objeto.
    """
    from app.configuracao import Configuracao, obter_configuracao

    padrao = obter_configuracao()
    como = Configuracao(**{**padrao.model_dump(), "auth_mock": True, "auth_mock_perfil": perfil})
    app.dependency_overrides[obter_sessao] = lambda: sessao
    app.dependency_overrides[obter_configuracao] = lambda: como
    return TestClient(app)


@pytest.fixture
def cliente_admin(sessao):
    """`plataforma_edicao` é o ÚNICO perfil com `administra_dicionarios` — ver
    `app/api/dependencias.py`, `exigir_administracao_de_cadastros`."""
    try:
        yield _cliente(sessao, "plataforma_edicao")
    finally:
        app.dependency_overrides.clear()


@pytest.fixture
def cliente_sem_admin(sessao):
    """`crm_edicao` abre o portal do CRM e NÃO administra cadastros: é o
    contrapeso exato da guarda, e não um perfil sem acesso a nada — que passaria
    o teste de 403 pelo motivo errado."""
    try:
        yield _cliente(sessao, "crm_edicao")
    finally:
        app.dependency_overrides.clear()


@pytest.fixture
def semente(sessao):
    instituicao = Instituicao(
        nome="Valor Econômico",
        nome_normalizado=normalizar("Valor Econômico"),
        tipo="veiculo",
        uf="SP",
    )
    sessao.add(instituicao)
    sessao.flush()
    interlocutor = Interlocutor(
        nome="Ana Prado",
        nome_normalizado=normalizar("Ana Prado"),
        instituicao_id=instituicao.id,
    )
    sessao.add_all(
        [
            interlocutor,
            PessoaAegea(nome="Radamés Casseb", nome_normalizado=normalizar("Radamés Casseb")),
        ]
    )
    sessao.flush()
    return {
        "instituicao": instituicao,
        "interlocutor": interlocutor,
        "formato": sessao.scalars(select(FormatoInteracao).limit(1)).first(),
        "clima": sessao.scalars(select(Clima).limit(1)).first(),
    }


def _declarar(pasta, chave: str, **valores) -> None:
    """Escreve uma linha na aba de cadastro daquele vocabulário, POR NOME DE COLUNA.

    As abas de cadastro passaram a ter os campos do formulário da plataforma, então a
    categoria de público deixou de ser a segunda coluna. Um teste que escreve por
    posição grava a categoria no "Nome completo" e não reclama de nada — e aí falha por
    um motivo que não é o que ele queria provar.
    """
    colunas = [coluna.nome for coluna in colunas_do_cadastro(chave)]
    desconhecidas = set(valores) - set(colunas)
    assert not desconhecidas, f"a aba {chave!r} não tem {desconhecidas}; tem {colunas}"
    pasta[ROTULO_DO_VOCABULARIO[chave]].append([valores.get(coluna) for coluna in colunas])


def _planilha_de_um_dia(sessao, semente, quantas: int = 3) -> bytes:
    """O modelo de verdade com algumas agendas preenchidas."""
    from datetime import date

    from openpyxl import load_workbook

    modelo = modelo_de_importacao.gerar(importar_agendas.vocabularios(sessao))
    pasta = load_workbook(io.BytesIO(modelo))
    folha = pasta["Agendas"]
    cabecalho = [celula.value for celula in next(folha.iter_rows())]
    for i in range(quantas):
        valores = {
            # DATAS DIFERENTES de propósito: com a mesma instituição e a mesma
            # data, as linhas 2, 3 e 4 seriam duplicatas uma da outra, e todo
            # teste que usa este helper passaria a carregar um aviso que não tem
            # nada a ver com o que ele quer provar.
            "Data": date(2026, 9, 21 + i),
            "Instituição": semente["instituicao"].nome,
            "UF": "SP",
        }
        folha.append([valores.get(coluna) for coluna in cabecalho])
    saida = io.BytesIO()
    pasta.save(saida)
    return saida.getvalue()


def _agenda_ja_no_sistema(cliente, sessao, semente, quando, arquivada: bool = False):
    """Uma agenda que JÁ existe, para a detecção de duplicata ter o que achar.

    Monta com os ids que o banco de teste tem, e não com códigos escritos aqui:
    `interacao` exige `frente_id`, `status_id` e `criado_por`, e inventar
    qualquer um deles daria erro de chave estrangeira em vez de testar o aviso.

    O `cliente` NÃO É DECORAÇÃO: o banco de teste é recriado a cada execução com
    só o que as migrations inserem, e `usuario` nasce VAZIA. O usuário do
    `auth_mock` só passa a existir na primeira requisição autenticada — então
    esta chamada barata é o que dá a `criado_por` alguém a quem apontar.
    """
    from app.banco.tabelas_acesso import Usuario
    from app.banco.tabelas_catalogo import Frente, Status
    from app.banco.tabelas_interacoes import InteracaoRegistro

    cliente.get("/api/importacoes/modelo")

    registro = InteracaoRegistro(
        data_interacao=quando,
        instituicao_id=semente["instituicao"].id,
        uf="SP",
        frente_id=sessao.scalars(select(Frente).limit(1)).first().id,
        status_id=sessao.scalars(select(Status).limit(1)).first().id,
        criado_por=sessao.scalars(select(Usuario).limit(1)).first().id,
        arquivado_em=quando if arquivada else None,
    )
    sessao.add(registro)
    sessao.flush()
    return registro


# =============================================================================
# a guarda
# =============================================================================


def test_quem_nao_administra_cadastros_nao_importa(cliente_sem_admin):
    """A decisão de produto: só a coordenação importa. E esconder o botão é
    conveniência — o servidor recusa por conta própria."""
    resposta = cliente_sem_admin.post(
        "/api/importacoes", files={"arquivo": ("a.xlsx", b"PK\x03\x04", TIPO_XLSX)}
    )

    assert resposta.status_code == 403


def test_quem_nao_administra_cadastros_nao_baixa_o_modelo(cliente_sem_admin):
    """O modelo lista o cadastro inteiro — instituições, interlocutores, pessoas
    da Aegea. Deixá-lo aberto entregaria o diretório a quem não o vê na tela."""
    assert cliente_sem_admin.get("/api/importacoes/modelo").status_code == 403


def test_toda_rota_sob_importacoes_exige_administrar_cadastros():
    """A ÂNCORA ESTRUTURAL, no mesmo molde da do portal: uma rota nova sob este
    prefixo nasce protegida, ou este teste quebra.

    Vale mais que os dois testes de 403 acima: eles provam duas rotas, este
    prova TODAS — inclusive as que a Tarefa 10 e a 11 ainda vão acrescentar.
    """
    from fastapi.routing import APIRoute

    from app.api.dependencias import exigir_administracao_de_cadastros
    from tests.test_portal_no_backend import _rotas_montadas

    desprotegidas = [
        f"{sorted(rota.methods)} {rota.path}"
        for rota in _rotas_montadas(app)
        if isinstance(rota, APIRoute)
        and rota.path.startswith("/api/importacoes")
        and exigir_administracao_de_cadastros not in {d.call for d in rota.dependant.dependencies}
    ]

    assert desprotegidas == []


def test_a_varredura_enxerga_as_rotas_de_importacao():
    """Contrapeso da âncora: uma varredura que não encontra nada passaria sem
    provar nada."""
    from fastapi.routing import APIRoute

    from tests.test_portal_no_backend import _rotas_montadas

    achadas = [
        rota.path
        for rota in _rotas_montadas(app)
        if isinstance(rota, APIRoute) and rota.path.startswith("/api/importacoes")
    ]

    assert len(achadas) >= 3, achadas


# =============================================================================
# o modelo
# =============================================================================


def test_o_modelo_baixa_um_xlsx_que_abre(cliente_admin):
    from openpyxl import load_workbook

    resposta = cliente_admin.get("/api/importacoes/modelo")

    assert resposta.status_code == 200
    assert resposta.content[:4] == b"PK\x03\x04"
    assert "Agendas" in load_workbook(io.BytesIO(resposta.content)).sheetnames


def test_o_modelo_vem_como_anexo_com_nome(cliente_admin):
    resposta = cliente_admin.get("/api/importacoes/modelo")

    assert "attachment" in resposta.headers["content-disposition"]
    # O NOME DIZ O RECORTE: dois arquivos de mesmo nome na pasta de downloads viram
    # "modelo (1).xlsx", e a pessoa abre o errado — descobrindo só ao procurar uma
    # coluna que aquele recorte não tem.
    assert "modelo-de-agendas-completo.xlsx" in resposta.headers["content-disposition"]


# =============================================================================
# subir
# =============================================================================


def test_subir_grava_a_importacao_e_devolve_as_propostas(cliente_admin, sessao, semente):
    resposta = cliente_admin.post(
        "/api/importacoes",
        files={"arquivo": ("dia.xlsx", _planilha_de_um_dia(sessao, semente), TIPO_XLSX)},
    )

    assert resposta.status_code == 201, resposta.text
    corpo = resposta.json()
    assert corpo["situacao"] == "aguardando_conferencia"
    assert len(corpo["linhas"]) == 3


def test_subir_guarda_de_que_linha_do_arquivo_cada_uma_veio(cliente_admin, sessao, semente):
    """`linha_origem` é o que deixa a pessoa voltar à planilha e conferir."""
    corpo = cliente_admin.post(
        "/api/importacoes",
        files={"arquivo": ("dia.xlsx", _planilha_de_um_dia(sessao, semente), TIPO_XLSX)},
    ).json()

    assert [linha["linha_origem"] for linha in corpo["linhas"]] == [2, 3, 4]


def test_subir_preserva_os_dados_brutos(cliente_admin, sessao, semente):
    """O bruto fica mesmo depois de aceito: é o que permite reprocessar quando a
    regra de leitura mudar, e responder de onde veio o registro."""
    corpo = cliente_admin.post(
        "/api/importacoes",
        files={"arquivo": ("dia.xlsx", _planilha_de_um_dia(sessao, semente), TIPO_XLSX)},
    ).json()

    assert corpo["linhas"][0]["dados_brutos"]["Instituição"] == semente["instituicao"].nome


def test_subir_grava_o_nome_do_arquivo(cliente_admin, sessao, semente):
    corpo = cliente_admin.post(
        "/api/importacoes",
        files={
            "arquivo": (
                "agendas de quinta.xlsx",
                _planilha_de_um_dia(sessao, semente),
                TIPO_XLSX,
            )
        },
    ).json()

    assert corpo["arquivo_nome"] == "agendas de quinta.xlsx"


def test_subir_um_arquivo_que_nao_e_planilha_devolve_422_e_nao_500(cliente_admin):
    """`RegraViolada` é traduzida centralmente. Sem isso a pessoa leria "erro
    interno" para um arquivo que ela mesma pode trocar."""
    resposta = cliente_admin.post(
        "/api/importacoes", files={"arquivo": ("a.xlsx", b"isto nao e planilha", TIPO_XLSX)}
    )

    assert resposta.status_code == 422
    assert "xlsx" in resposta.text


def test_subir_uma_planilha_sem_as_abas_devolve_422(cliente_admin):
    from openpyxl import Workbook

    livro = Workbook()
    saida = io.BytesIO()
    livro.save(saida)

    resposta = cliente_admin.post(
        "/api/importacoes", files={"arquivo": ("a.xlsx", saida.getvalue(), TIPO_XLSX)}
    )

    assert resposta.status_code == 422


# =============================================================================
# retomar
# =============================================================================


def test_a_importacao_sobrevive_a_fechar_o_navegador(cliente_admin, sessao, semente):
    """É para isto que a 0008 criou tabela em vez de resolver em memória."""
    criada = cliente_admin.post(
        "/api/importacoes",
        files={"arquivo": ("dia.xlsx", _planilha_de_um_dia(sessao, semente), TIPO_XLSX)},
    ).json()

    retomada = cliente_admin.get(f"/api/importacoes/{criada['id']}")

    assert retomada.status_code == 200
    assert retomada.json()["linhas"] == criada["linhas"]


def test_uma_importacao_que_nao_existe_devolve_404(cliente_admin):
    resposta = cliente_admin.get("/api/importacoes/00000000-0000-0000-0000-000000000000")

    assert resposta.status_code == 404


# =============================================================================
# a duplicata possível
# =============================================================================


def test_agenda_que_ja_existe_no_banco_AVISA_e_nao_trava(cliente_admin, sessao, semente):
    """Duas reuniões com o mesmo órgão no mesmo dia acontecem. Travar por isso
    ensinaria a pessoa a ignorar o aviso — e é o aviso que a protege do caso que
    importa: ela subiu o mesmo arquivo duas vezes."""
    from datetime import date

    _agenda_ja_no_sistema(cliente_admin, sessao, semente, date(2026, 9, 21))

    corpo = cliente_admin.post(
        "/api/importacoes",
        files={"arquivo": ("dia.xlsx", _planilha_de_um_dia(sessao, semente, quantas=1), TIPO_XLSX)},
    ).json()

    duplicatas = [grupo for grupo in corpo["grupos"] if grupo["campo"] == "duplicata"]
    assert duplicatas, corpo["grupos"]
    assert duplicatas[0]["trava"] is False
    assert corpo["pendencias"] == 0


def test_a_mesma_agenda_repetida_DENTRO_do_arquivo_tambem_avisa(cliente_admin, sessao, semente):
    """Colar a mesma agenda duas vezes na planilha é o mesmo erro visto de outro
    ângulo, e quem confere precisa vê-lo antes de criar as duas."""
    from datetime import date

    from openpyxl import load_workbook

    modelo = modelo_de_importacao.gerar(importar_agendas.vocabularios(sessao))
    pasta = load_workbook(io.BytesIO(modelo))
    folha = pasta["Agendas"]
    cabecalho = [celula.value for celula in next(folha.iter_rows())]
    for _ in range(2):
        valores = {
            "Data": date(2026, 9, 25),
            "Instituição": semente["instituicao"].nome,
            "UF": "SP",
        }
        folha.append([valores.get(coluna) for coluna in cabecalho])
    saida = io.BytesIO()
    pasta.save(saida)

    corpo = cliente_admin.post(
        "/api/importacoes", files={"arquivo": ("dia.xlsx", saida.getvalue(), TIPO_XLSX)}
    ).json()

    duplicatas = [grupo for grupo in corpo["grupos"] if grupo["campo"] == "duplicata"]
    assert duplicatas, corpo["grupos"]
    # A PRIMEIRA não é duplicata de nada; a segunda é.
    assert duplicatas[0]["linhas"] == [3]


def test_a_agenda_arquivada_nao_conta_como_duplicata(cliente_admin, sessao, semente):
    """Uma agenda arquivada não é a mesma reunião acontecendo de novo, e avisar
    sobre ela ensinaria a ignorar o aviso."""
    from datetime import date

    _agenda_ja_no_sistema(cliente_admin, sessao, semente, date(2026, 9, 21), arquivada=True)

    corpo = cliente_admin.post(
        "/api/importacoes",
        files={"arquivo": ("dia.xlsx", _planilha_de_um_dia(sessao, semente, quantas=1), TIPO_XLSX)},
    ).json()

    assert [grupo for grupo in corpo["grupos"] if grupo["campo"] == "duplicata"] == []


# =============================================================================
# os grupos na resposta
# =============================================================================


def test_a_mesma_instituicao_desconhecida_vira_UM_grupo_com_as_linhas(
    cliente_admin, sessao, semente
):
    """A vista principal da conferência: uma decisão, várias linhas."""
    from datetime import date

    from openpyxl import load_workbook

    modelo = modelo_de_importacao.gerar(importar_agendas.vocabularios(sessao))
    pasta = load_workbook(io.BytesIO(modelo))
    folha = pasta["Agendas"]
    cabecalho = [celula.value for celula in next(folha.iter_rows())]
    for _ in range(4):
        valores = {
            "Data": date(2026, 9, 25),
            "Instituição": "Prefeitura de Campinas",
            "UF": "SP",
        }
        folha.append([valores.get(coluna) for coluna in cabecalho])
    saida = io.BytesIO()
    pasta.save(saida)

    corpo = cliente_admin.post(
        "/api/importacoes", files={"arquivo": ("dia.xlsx", saida.getvalue(), TIPO_XLSX)}
    ).json()

    (grupo,) = [g for g in corpo["grupos"] if g["campo"] == "instituicao_id"]
    assert grupo["valor"] == "Prefeitura de Campinas"
    assert grupo["linhas"] == [2, 3, 4, 5]
    assert grupo["trava"] is True
    # QUATRO pendências e UMA decisão. Este teste dizia `pendencias == 1`, e com
    # isso codificava o defeito que a revisão do Codex achou: o campo promete
    # "quantas linhas seguram a confirmação" e contava grupos, fazendo o cabeçalho
    # mentir sobre o tamanho do trabalho.
    assert corpo["pendencias"] == 4
    assert corpo["decisoes_pendentes"] == 1


def test_o_grupo_oferece_o_nome_parecido_que_JA_existe(cliente_admin, sessao, semente):
    """É o que transforma "não existe" em um clique: a pessoa reconhece
    "Prefeitura Municipal de Campinas" e aponta para ela."""
    from datetime import date

    from openpyxl import load_workbook

    from app.banco.tabelas_stakeholders import Instituicao as Inst

    sessao.add(
        Inst(
            nome="Prefeitura Municipal de Campinas",
            nome_normalizado=normalizar("Prefeitura Municipal de Campinas"),
            tipo="orgao",
            uf="SP",
        )
    )
    sessao.flush()

    modelo = modelo_de_importacao.gerar(importar_agendas.vocabularios(sessao))
    pasta = load_workbook(io.BytesIO(modelo))
    folha = pasta["Agendas"]
    cabecalho = [celula.value for celula in next(folha.iter_rows())]
    valores = {
        "Data": date(2026, 9, 25),
        "Instituição": "Prefeitura de Campinas",
        "UF": "SP",
    }
    folha.append([valores.get(coluna) for coluna in cabecalho])
    saida = io.BytesIO()
    pasta.save(saida)

    corpo = cliente_admin.post(
        "/api/importacoes", files={"arquivo": ("dia.xlsx", saida.getvalue(), TIPO_XLSX)}
    ).json()

    (grupo,) = [g for g in corpo["grupos"] if g["campo"] == "instituicao_id"]
    # A sugestão carrega nome E alvo: o nome que a pessoa reconhece e o id que
    # o servidor aceita em `apontar`.
    assert "Prefeitura Municipal de Campinas" in [s["nome"] for s in grupo["sugestoes"]]


def test_o_que_trava_vem_ANTES_do_que_so_avisa(cliente_admin, sessao, semente):
    """A tela lista em ordem de urgência: o que segura a confirmação primeiro."""
    from datetime import date

    from openpyxl import load_workbook

    _agenda_ja_no_sistema(cliente_admin, sessao, semente, date(2026, 9, 25))

    modelo = modelo_de_importacao.gerar(importar_agendas.vocabularios(sessao))
    pasta = load_workbook(io.BytesIO(modelo))
    folha = pasta["Agendas"]
    cabecalho = [celula.value for celula in next(folha.iter_rows())]
    # A primeira resolve inteira e é duplicata (avisa); a segunda tem
    # instituição desconhecida (trava).
    for instituicao in (semente["instituicao"].nome, "Prefeitura de Campinas"):
        valores = {
            "Data": date(2026, 9, 25),
            "Instituição": instituicao,
            "UF": "SP",
        }
        folha.append([valores.get(coluna) for coluna in cabecalho])
    saida = io.BytesIO()
    pasta.save(saida)

    corpo = cliente_admin.post(
        "/api/importacoes", files={"arquivo": ("dia.xlsx", saida.getvalue(), TIPO_XLSX)}
    ).json()

    assert [g["trava"] for g in corpo["grupos"]] == sorted(
        [g["trava"] for g in corpo["grupos"]], reverse=True
    )
    assert corpo["grupos"][0]["trava"] is True


def test_a_linha_limpa_nao_gera_grupo_nenhum(cliente_admin, sessao, semente):
    corpo = cliente_admin.post(
        "/api/importacoes",
        files={"arquivo": ("dia.xlsx", _planilha_de_um_dia(sessao, semente), TIPO_XLSX)},
    ).json()

    assert corpo["grupos"] == []
    assert corpo["pendencias"] == 0


# =============================================================================
# os defeitos que a revisão do Codex achou
# =============================================================================


def test_pendencias_conta_LINHAS_e_nao_grupos(cliente_admin, sessao, semente):
    """O DEFEITO 5: o campo diz "quantas linhas seguram a confirmação" e contava
    grupos. Quatro linhas travadas por um grupo devolviam 1, e o cabeçalho da
    tela mentia sobre o tamanho do trabalho."""
    from datetime import date

    from openpyxl import load_workbook

    modelo = modelo_de_importacao.gerar(importar_agendas.vocabularios(sessao))
    pasta = load_workbook(io.BytesIO(modelo))
    folha = pasta["Agendas"]
    cabecalho = [celula.value for celula in next(folha.iter_rows())]
    for i in range(4):
        valores = {
            "Data": date(2026, 9, 21 + i),
            "Instituição": "Prefeitura de Campinas",
            "UF": "SP",
        }
        folha.append([valores.get(coluna) for coluna in cabecalho])
    saida = io.BytesIO()
    pasta.save(saida)

    corpo = cliente_admin.post(
        "/api/importacoes", files={"arquivo": ("dia.xlsx", saida.getvalue(), TIPO_XLSX)}
    ).json()

    assert corpo["pendencias"] == 4
    assert corpo["decisoes_pendentes"] == 1

def test_celula_de_HORA_nao_estoura_no_commit(cliente_admin, sessao, semente):
    """O DEFEITO 6: `para_json` cobria date, datetime e UUID, e uma célula
    formatada como hora chega `time` — estourando no flush do JSONB, longe de
    quem montou o dado. É a classe de erro que a função existe para evitar."""
    from datetime import date, time

    from openpyxl import load_workbook

    modelo = modelo_de_importacao.gerar(importar_agendas.vocabularios(sessao))
    pasta = load_workbook(io.BytesIO(modelo))
    folha = pasta["Agendas"]
    cabecalho = [celula.value for celula in next(folha.iter_rows())]
    valores = {
        "Data": date(2026, 9, 25),
        "Instituição": semente["instituicao"].nome,
        "UF": "SP",
        "Local": time(14, 30),
    }
    folha.append([valores.get(coluna) for coluna in cabecalho])
    saida = io.BytesIO()
    pasta.save(saida)

    resposta = cliente_admin.post(
        "/api/importacoes", files={"arquivo": ("dia.xlsx", saida.getvalue(), TIPO_XLSX)}
    )

    assert resposta.status_code == 201, resposta.text
    assert resposta.json()["linhas"][0]["dados_brutos"]["Local"] == "14:30:00"


# =============================================================================
# resolver um grupo resolve todas as linhas dele
# =============================================================================


@pytest.fixture
def importacao_com_quatro(cliente_admin, sessao, semente):
    """Quatro agendas presas pela MESMA instituição desconhecida.

    É a forma do problema real: um dia inteiro com o mesmo órgão que ninguém
    cadastrou ainda. Uma decisão tem de destravar as quatro.
    """
    from datetime import date

    from openpyxl import load_workbook

    modelo = modelo_de_importacao.gerar(importar_agendas.vocabularios(sessao))
    pasta = load_workbook(io.BytesIO(modelo))
    folha = pasta["Agendas"]
    cabecalho = [celula.value for celula in next(folha.iter_rows())]
    for i in range(4):
        valores = {
            "Data": date(2026, 9, 21 + i),
            "Instituição": "Prefeitura de Campinas",
            "UF": "SP",
        }
        folha.append([valores.get(coluna) for coluna in cabecalho])
    saida = io.BytesIO()
    pasta.save(saida)

    resposta = cliente_admin.post(
        "/api/importacoes", files={"arquivo": ("dia.xlsx", saida.getvalue(), TIPO_XLSX)}
    )
    assert resposta.status_code == 201, resposta.text
    return resposta.json()


def test_apontar_para_uma_existente_resolve_as_quatro_linhas(
    cliente_admin, sessao, semente, importacao_com_quatro
):
    """UMA DECISÃO, QUATRO LINHAS — é o que faz a conferência escalar."""
    parecida = Instituicao(
        nome="Prefeitura Municipal de Campinas",
        nome_normalizado=normalizar("Prefeitura Municipal de Campinas"),
        tipo="orgao",
        uf="SP",
    )
    sessao.add(parecida)
    sessao.flush()

    resposta = cliente_admin.patch(
        f"/api/importacoes/{importacao_com_quatro['id']}/resolucoes",
        json={
            "campo": "instituicao_id",
            "valor": "Prefeitura de Campinas",
            "decisao": "apontar",
            "alvo": str(parecida.id),
        },
    )

    assert resposta.status_code == 200, resposta.text
    corpo = resposta.json()
    assert corpo["pendencias"] == 0
    assert corpo["decisoes_pendentes"] == 0


def test_apontar_registra_para_qual_cadastro(cliente_admin, sessao, semente, importacao_com_quatro):
    """Sem guardar o alvo, a confirmação não saberia para onde apontar — e a
    decisão da pessoa teria sido só um número que baixou na tela."""
    parecida = Instituicao(
        nome="Prefeitura Municipal de Campinas",
        nome_normalizado=normalizar("Prefeitura Municipal de Campinas"),
        tipo="orgao",
        uf="SP",
    )
    sessao.add(parecida)
    sessao.flush()

    cliente_admin.patch(
        f"/api/importacoes/{importacao_com_quatro['id']}/resolucoes",
        json={
            "campo": "instituicao_id",
            "valor": "Prefeitura de Campinas",
            "decisao": "apontar",
            "alvo": str(parecida.id),
        },
    )

    estado = cliente_admin.get(f"/api/importacoes/{importacao_com_quatro['id']}").json()
    resolvidas = [
        divergencia
        for linha in estado["linhas"]
        for divergencia in linha["divergencias"]
        if divergencia.get("acao") == "apontar"
    ]
    assert len(resolvidas) == 4
    assert all(d["alvo"] == str(parecida.id) for d in resolvidas)


def test_criar_transforma_o_que_travava_em_cadastro_a_criar(
    cliente_admin, importacao_com_quatro
):
    """A pessoa confirma que é cadastro novo mesmo — digitou na célula sem
    declarar na aba, e agora decide na tela. Deixa de travar."""
    resposta = cliente_admin.patch(
        f"/api/importacoes/{importacao_com_quatro['id']}/resolucoes",
        json={
            "campo": "instituicao_id",
            "valor": "Prefeitura de Campinas",
            "decisao": "criar",
        },
    )

    assert resposta.status_code == 200, resposta.text
    assert resposta.json()["pendencias"] == 0


def test_descartar_tira_as_linhas_sem_apagar_o_bruto(cliente_admin, importacao_com_quatro):
    """`dados_brutos` fica: é o que permite reprocessar quando a regra de leitura
    mudar, e responder de onde veio o registro."""
    cliente_admin.patch(
        f"/api/importacoes/{importacao_com_quatro['id']}/resolucoes",
        json={
            "campo": "instituicao_id",
            "valor": "Prefeitura de Campinas",
            "decisao": "descartar",
        },
    )

    estado = cliente_admin.get(f"/api/importacoes/{importacao_com_quatro['id']}").json()
    descartadas = [linha for linha in estado["linhas"] if linha["decisao"] == "descartada"]
    assert len(descartadas) == 4
    assert all(linha["dados_brutos"] for linha in descartadas)


def test_a_linha_descartada_nao_conta_mais_como_pendencia(cliente_admin, importacao_com_quatro):
    """Descartar é uma resolução: a linha sai do caminho da confirmação."""
    resposta = cliente_admin.patch(
        f"/api/importacoes/{importacao_com_quatro['id']}/resolucoes",
        json={
            "campo": "instituicao_id",
            "valor": "Prefeitura de Campinas",
            "decisao": "descartar",
        },
    )

    assert resposta.json()["pendencias"] == 0


def test_resolver_um_valor_que_nao_esta_pendente_recusa(cliente_admin, importacao_com_quatro):
    """Recusar em vez de não fazer nada: um 200 silencioso faria a tela mostrar
    "resolvido" para uma decisão que não alcançou linha nenhuma."""
    resposta = cliente_admin.patch(
        f"/api/importacoes/{importacao_com_quatro['id']}/resolucoes",
        json={"campo": "instituicao_id", "valor": "Não existe", "decisao": "criar"},
    )

    assert resposta.status_code == 422


def test_apontar_sem_alvo_recusa(cliente_admin, importacao_com_quatro):
    resposta = cliente_admin.patch(
        f"/api/importacoes/{importacao_com_quatro['id']}/resolucoes",
        json={
            "campo": "instituicao_id",
            "valor": "Prefeitura de Campinas",
            "decisao": "apontar",
        },
    )

    assert resposta.status_code == 422


def test_apontar_para_um_cadastro_que_nao_existe_recusa(cliente_admin, importacao_com_quatro):
    """Sem esta guarda, a decisão gravaria um id que a confirmação não encontra —
    e o erro apareceria lá, longe de quem escolheu."""
    resposta = cliente_admin.patch(
        f"/api/importacoes/{importacao_com_quatro['id']}/resolucoes",
        json={
            "campo": "instituicao_id",
            "valor": "Prefeitura de Campinas",
            "decisao": "apontar",
            "alvo": "00000000-0000-0000-0000-000000000000",
        },
    )

    assert resposta.status_code == 422


def test_decisao_desconhecida_recusa(cliente_admin, importacao_com_quatro):
    resposta = cliente_admin.patch(
        f"/api/importacoes/{importacao_com_quatro['id']}/resolucoes",
        json={
            "campo": "instituicao_id",
            "valor": "Prefeitura de Campinas",
            "decisao": "inventada",
        },
    )

    assert resposta.status_code == 422


def test_resolver_nao_mexe_em_linha_de_outro_valor(cliente_admin, sessao, semente):
    """Duas divergências diferentes, uma decisão: só as linhas daquele valor."""
    from datetime import date

    from openpyxl import load_workbook

    modelo = modelo_de_importacao.gerar(importar_agendas.vocabularios(sessao))
    pasta = load_workbook(io.BytesIO(modelo))
    folha = pasta["Agendas"]
    cabecalho = [celula.value for celula in next(folha.iter_rows())]
    for instituicao in ("Prefeitura de Campinas", "Prefeitura de Campos"):
        valores = {
            "Data": date(2026, 9, 25),
            "Instituição": instituicao,
            "UF": "SP",
        }
        folha.append([valores.get(coluna) for coluna in cabecalho])
    saida = io.BytesIO()
    pasta.save(saida)

    criada = cliente_admin.post(
        "/api/importacoes", files={"arquivo": ("dia.xlsx", saida.getvalue(), TIPO_XLSX)}
    ).json()
    assert criada["pendencias"] == 2

    resposta = cliente_admin.patch(
        f"/api/importacoes/{criada['id']}/resolucoes",
        json={
            "campo": "instituicao_id",
            "valor": "Prefeitura de Campinas",
            "decisao": "descartar",
        },
    )

    assert resposta.json()["pendencias"] == 1


def test_resolver_uma_importacao_que_nao_existe_devolve_404(cliente_admin):
    resposta = cliente_admin.patch(
        "/api/importacoes/00000000-0000-0000-0000-000000000000/resolucoes",
        json={"campo": "instituicao_id", "valor": "X", "decisao": "criar"},
    )

    assert resposta.status_code == 404


# =============================================================================
# a confirmação: uma transação só, reconferida contra o banco
# =============================================================================


def _quantas(sessao, tabela) -> int:
    from sqlalchemy import func
    from sqlalchemy import select as sel

    return sessao.scalar(sel(func.count()).select_from(tabela))


@pytest.fixture
def limpa(cliente_admin, sessao, semente):
    """Três agendas que resolvem inteiras — nada a decidir."""
    corpo = cliente_admin.post(
        "/api/importacoes",
        files={"arquivo": ("dia.xlsx", _planilha_de_um_dia(sessao, semente), TIPO_XLSX)},
    ).json()
    assert corpo["pendencias"] == 0, corpo["grupos"]
    return corpo


@pytest.fixture
def com_cadastro_novo(cliente_admin, sessao, semente):
    """Uma agenda cuja instituição foi DECLARADA na aba editável — a criar."""
    from datetime import date

    from openpyxl import load_workbook


    modelo = modelo_de_importacao.gerar(importar_agendas.vocabularios(sessao))
    pasta = load_workbook(io.BytesIO(modelo))
    folha = pasta["Agendas"]
    cabecalho = [celula.value for celula in next(folha.iter_rows())]
    valores = {
        "Data": date(2026, 9, 25),
        "Instituição": "Prefeitura de Campinas",
        "UF": "SP",
    }
    folha.append([valores.get(coluna) for coluna in cabecalho])
    # COM O TIPO: ele deriva a frente da agenda, e a importação recusa criar
    # instituição sem ele — chutar erraria a frente de toda agenda dela.
    _declarar(
        pasta,
        "instituicoes",
        **{
            "Instituição": "Prefeitura de Campinas",
            "Categoria de público": "Poder Executivo",
        },
    )
    saida = io.BytesIO()
    pasta.save(saida)

    corpo = cliente_admin.post(
        "/api/importacoes", files={"arquivo": ("dia.xlsx", saida.getvalue(), TIPO_XLSX)}
    ).json()
    assert corpo["pendencias"] == 0, corpo["grupos"]
    return corpo


def test_confirmar_cria_as_agendas(cliente_admin, sessao, limpa):
    from app.banco.tabelas_interacoes import InteracaoRegistro

    antes = _quantas(sessao, InteracaoRegistro)

    resposta = cliente_admin.post(f"/api/importacoes/{limpa['id']}/confirmacao")

    assert resposta.status_code == 201, resposta.text
    assert _quantas(sessao, InteracaoRegistro) == antes + 3
    assert resposta.json()["criadas"] == 3


def test_confirmar_cria_o_cadastro_declarado_junto(cliente_admin, sessao, com_cadastro_novo):
    """Os cadastros novos e as agendas nascem no MESMO commit."""
    from app.banco.tabelas_interacoes import InteracaoRegistro

    antes_inst = _quantas(sessao, Instituicao)
    antes_ag = _quantas(sessao, InteracaoRegistro)

    resposta = cliente_admin.post(f"/api/importacoes/{com_cadastro_novo['id']}/confirmacao")

    assert resposta.status_code == 201, resposta.text
    assert _quantas(sessao, Instituicao) == antes_inst + 1
    assert _quantas(sessao, InteracaoRegistro) == antes_ag + 1
    assert resposta.json()["cadastros"] == 1


def test_a_agenda_criada_guarda_de_onde_veio(cliente_admin, sessao, limpa):
    """A procedência é o que permite responder "de onde veio este registro" e
    reprocessar quando a regra de leitura mudar."""
    from sqlalchemy import select as sel

    from app.banco.tabelas_interacoes import InteracaoRegistro

    cliente_admin.post(f"/api/importacoes/{limpa['id']}/confirmacao")

    criadas = sessao.scalars(
        sel(InteracaoRegistro).where(InteracaoRegistro.fonte == "importacao_planilha")
    ).all()
    assert len(criadas) == 3
    assert {linha.origem_aba for linha in criadas} == {"Agendas"}
    assert sorted(linha.origem_linha for linha in criadas) == [2, 3, 4]


def test_confirmar_liga_cada_linha_a_sua_agenda(cliente_admin, limpa):
    """`interacao_id` é o que fecha o laço entre a planilha e o registro."""
    cliente_admin.post(f"/api/importacoes/{limpa['id']}/confirmacao")

    estado = cliente_admin.get(f"/api/importacoes/{limpa['id']}").json()
    das_agendas = [linha for linha in estado["linhas"] if linha["aba"] == "Agendas"]
    assert all(linha["interacao_id"] for linha in das_agendas)


def test_confirmar_marca_a_importacao_como_confirmada(cliente_admin, limpa):
    cliente_admin.post(f"/api/importacoes/{limpa['id']}/confirmacao")

    estado = cliente_admin.get(f"/api/importacoes/{limpa['id']}").json()
    assert estado["situacao"] == "confirmada"
    assert estado["confirmado_em"]


def test_o_cadastro_que_NASCEU_entre_subir_e_confirmar_para_a_confirmacao(
    cliente_admin, sessao, com_cadastro_novo
):
    """O PONTO MAIS DELICADO DA FUNCIONALIDADE.

    Entre subir e confirmar, alguém cadastrou a mesma instituição pela tela de
    Administração. Confirmar cego criaria a duplicata que a conferência existia
    para evitar — e duplicata de instituição é o defeito mais caro aqui, porque
    espalha por toda leitura que agrupa por órgão.
    """
    sessao.add(
        Instituicao(
            nome="Prefeitura de Campinas",
            nome_normalizado=normalizar("Prefeitura de Campinas"),
            tipo="orgao",
            uf="SP",
        )
    )
    sessao.flush()

    resposta = cliente_admin.post(f"/api/importacoes/{com_cadastro_novo['id']}/confirmacao")

    assert resposta.status_code == 409, resposta.text
    assert "Prefeitura de Campinas" in resposta.text


def test_confirmar_com_pendencia_que_trava_recusa(cliente_admin, sessao, semente):
    """O botão fica apagado na tela, e o servidor recusa por conta própria."""
    from datetime import date

    from openpyxl import load_workbook

    modelo = modelo_de_importacao.gerar(importar_agendas.vocabularios(sessao))
    pasta = load_workbook(io.BytesIO(modelo))
    folha = pasta["Agendas"]
    cabecalho = [celula.value for celula in next(folha.iter_rows())]
    valores = {
        "Data": date(2026, 9, 25),
        "Instituição": "Orgao Que Ninguem Cadastrou",
        "UF": "SP",
    }
    folha.append([valores.get(coluna) for coluna in cabecalho])
    saida = io.BytesIO()
    pasta.save(saida)

    criada = cliente_admin.post(
        "/api/importacoes", files={"arquivo": ("dia.xlsx", saida.getvalue(), TIPO_XLSX)}
    ).json()
    assert criada["pendencias"] == 1

    resposta = cliente_admin.post(f"/api/importacoes/{criada['id']}/confirmacao")

    assert resposta.status_code == 422


def test_confirmar_duas_vezes_recusa(cliente_admin, limpa):
    """Sem esta guarda, dois cliques no botão criariam as agendas duas vezes."""
    primeira = cliente_admin.post(f"/api/importacoes/{limpa['id']}/confirmacao")
    assert primeira.status_code == 201

    segunda = cliente_admin.post(f"/api/importacoes/{limpa['id']}/confirmacao")

    assert segunda.status_code == 422


def test_a_linha_descartada_nao_vira_agenda(cliente_admin, sessao, importacao_com_quatro):
    """Descartar é decisão, e ela tem de valer na confirmação."""
    from app.banco.tabelas_interacoes import InteracaoRegistro

    cliente_admin.patch(
        f"/api/importacoes/{importacao_com_quatro['id']}/resolucoes",
        json={
            "campo": "instituicao_id",
            "valor": "Prefeitura de Campinas",
            "decisao": "descartar",
        },
    )
    antes = _quantas(sessao, InteracaoRegistro)

    resposta = cliente_admin.post(f"/api/importacoes/{importacao_com_quatro['id']}/confirmacao")

    assert resposta.status_code == 201, resposta.text
    assert _quantas(sessao, InteracaoRegistro) == antes
    assert resposta.json()["criadas"] == 0


def test_apontar_e_confirmar_usa_o_cadastro_escolhido(
    cliente_admin, sessao, semente, importacao_com_quatro
):
    """O caminho completo da conferência: a pessoa aponta e a agenda nasce no
    cadastro que ela escolheu, não num criado às pressas."""
    from sqlalchemy import select as sel

    from app.banco.tabelas_interacoes import InteracaoRegistro

    escolhida = Instituicao(
        nome="Prefeitura Municipal de Campinas",
        nome_normalizado=normalizar("Prefeitura Municipal de Campinas"),
        tipo="orgao",
        uf="SP",
    )
    sessao.add(escolhida)
    sessao.flush()

    cliente_admin.patch(
        f"/api/importacoes/{importacao_com_quatro['id']}/resolucoes",
        json={
            "campo": "instituicao_id",
            "valor": "Prefeitura de Campinas",
            "decisao": "apontar",
            "alvo": str(escolhida.id),
        },
    )

    resposta = cliente_admin.post(f"/api/importacoes/{importacao_com_quatro['id']}/confirmacao")

    assert resposta.status_code == 201, resposta.text
    criadas = sessao.scalars(
        sel(InteracaoRegistro).where(InteracaoRegistro.instituicao_id == escolhida.id)
    ).all()
    assert len(criadas) == 4


def test_cancelar_nao_deixa_cadastro_orfao(cliente_admin, sessao, com_cadastro_novo):
    """Se os cadastros nascessem no upload, cancelar deixaria instituições
    criadas sem nenhuma agenda a que servissem. É por isso que nada nasce lá."""
    antes = _quantas(sessao, Instituicao)

    resposta = cliente_admin.post(f"/api/importacoes/{com_cadastro_novo['id']}/cancelamento")

    assert resposta.status_code == 200, resposta.text
    assert _quantas(sessao, Instituicao) == antes
    assert resposta.json()["situacao"] == "cancelada"


def test_cancelar_preserva_o_bruto(cliente_admin, com_cadastro_novo):
    """Cancelar não apaga o que a pessoa preencheu: ela pode querer entender o
    que deu errado, e o bruto é a única cópia daquilo no sistema."""
    cliente_admin.post(f"/api/importacoes/{com_cadastro_novo['id']}/cancelamento")

    estado = cliente_admin.get(f"/api/importacoes/{com_cadastro_novo['id']}").json()
    assert all(linha["dados_brutos"] for linha in estado["linhas"])


def test_confirmar_uma_cancelada_recusa(cliente_admin, limpa):
    cliente_admin.post(f"/api/importacoes/{limpa['id']}/cancelamento")

    resposta = cliente_admin.post(f"/api/importacoes/{limpa['id']}/confirmacao")

    assert resposta.status_code == 422


def test_confirmar_uma_importacao_que_nao_existe_devolve_404(cliente_admin):
    resposta = cliente_admin.post(
        "/api/importacoes/00000000-0000-0000-0000-000000000000/confirmacao"
    )

    assert resposta.status_code == 404


# =============================================================================
# os defeitos que a revisão do Codex achou na Tarefa 10
# =============================================================================


def test_a_divergencia_resolvida_SAI_das_pendencias_e_vira_a_criar(
    cliente_admin, importacao_com_quatro
):
    """DEFEITO 1. A divergência resolvida ficava na lista com `trava=False`, e o
    agrupamento a devolvia como se fosse aviso comum — escondendo a decisão que a
    pessoa tomou e fazendo a tela parecer ter uma pendência branda que não existe.
    Resolvido é resolvido: sai dos grupos e entra no bloco "o que vou criar"."""
    resposta = cliente_admin.patch(
        f"/api/importacoes/{importacao_com_quatro['id']}/resolucoes",
        json={
            "campo": "instituicao_id",
            "valor": "Prefeitura de Campinas",
            "decisao": "criar",
        },
    )

    corpo = resposta.json()
    assert corpo["grupos"] == [], corpo["grupos"]
    assert ["Prefeitura de Campinas"] == [item["valor"] for item in corpo["a_criar"]]


def test_reaplicar_a_mesma_decisao_recusa(cliente_admin, importacao_com_quatro):
    """DEFEITO 1, a outra metade. `linhas_com` continuava encontrando a
    divergência já resolvida, então a guarda de "zero linhas alcançadas" não
    protegia contra sobrescrever uma decisão tomada — um segundo clique podia
    trocar "apontar para X" por "criar" sem ninguém perceber."""
    primeira = cliente_admin.patch(
        f"/api/importacoes/{importacao_com_quatro['id']}/resolucoes",
        json={
            "campo": "instituicao_id",
            "valor": "Prefeitura de Campinas",
            "decisao": "criar",
        },
    )
    assert primeira.status_code == 200

    segunda = cliente_admin.patch(
        f"/api/importacoes/{importacao_com_quatro['id']}/resolucoes",
        json={
            "campo": "instituicao_id",
            "valor": "Prefeitura de Campinas",
            "decisao": "criar",
        },
    )

    assert segunda.status_code == 422


def test_o_declarado_na_aba_ja_aparece_como_a_criar(cliente_admin, sessao, semente):
    """A declaração na aba editável e a decisão "criar" na tela passam a ter a
    MESMA forma, e as duas aparecem no mesmo bloco — que é o que a spec descreve."""
    from datetime import date

    from openpyxl import load_workbook


    modelo = modelo_de_importacao.gerar(importar_agendas.vocabularios(sessao))
    pasta = load_workbook(io.BytesIO(modelo))
    folha = pasta["Agendas"]
    cabecalho = [celula.value for celula in next(folha.iter_rows())]
    valores = {
        "Data": date(2026, 9, 25),
        "Instituição": "Prefeitura de Campinas",
        "UF": "SP",
    }
    folha.append([valores.get(coluna) for coluna in cabecalho])
    # COM O TIPO: ele deriva a frente da agenda, e a importação recusa criar
    # instituição sem ele — chutar erraria a frente de toda agenda dela.
    _declarar(
        pasta,
        "instituicoes",
        **{
            "Instituição": "Prefeitura de Campinas",
            "Categoria de público": "Poder Executivo",
        },
    )
    saida = io.BytesIO()
    pasta.save(saida)

    corpo = cliente_admin.post(
        "/api/importacoes", files={"arquivo": ("dia.xlsx", saida.getvalue(), TIPO_XLSX)}
    ).json()

    assert corpo["grupos"] == []
    assert [item["valor"] for item in corpo["a_criar"]] == ["Prefeitura de Campinas"]


def test_apontar_para_um_vocabulario_DE_CODIGO_funciona(cliente_admin, sessao, semente):
    """DEFEITO 2. `_cadastro_existe` só olhava `NO_BANCO`, então uma Modalidade
    escrita errada não podia ser resolvida apontando para um código válido — a
    pessoa ficava sem saída numa pendência que travava."""
    from datetime import date

    from openpyxl import load_workbook

    modelo = modelo_de_importacao.gerar(importar_agendas.vocabularios(sessao))
    pasta = load_workbook(io.BytesIO(modelo))
    folha = pasta["Agendas"]
    cabecalho = [celula.value for celula in next(folha.iter_rows())]
    valores = {
        "Data": date(2026, 9, 25),
        "Instituição": semente["instituicao"].nome,
        "UF": "SP",
        "Modalidade": "pessoalmente",
    }
    folha.append([valores.get(coluna) for coluna in cabecalho])
    saida = io.BytesIO()
    pasta.save(saida)

    criada = cliente_admin.post(
        "/api/importacoes", files={"arquivo": ("dia.xlsx", saida.getvalue(), TIPO_XLSX)}
    ).json()
    assert criada["pendencias"] == 1

    resposta = cliente_admin.patch(
        f"/api/importacoes/{criada['id']}/resolucoes",
        json={
            "campo": "modalidade",
            "valor": "pessoalmente",
            "decisao": "apontar",
            "alvo": "presencial",
        },
    )

    assert resposta.status_code == 200, resposta.text
    assert resposta.json()["pendencias"] == 0


def test_criar_num_vocabulario_FECHADO_recusa(cliente_admin, sessao, semente):
    """DEFEITO 2, a metade pior. `criar` não era validado contra vocabulário
    fechado, então a resolução contornava a regra que `classificar` enforça no
    upload: mudar a lista de modalidades é mudança de regra, não de cadastro."""
    from datetime import date

    from openpyxl import load_workbook

    modelo = modelo_de_importacao.gerar(importar_agendas.vocabularios(sessao))
    pasta = load_workbook(io.BytesIO(modelo))
    folha = pasta["Agendas"]
    cabecalho = [celula.value for celula in next(folha.iter_rows())]
    valores = {
        "Data": date(2026, 9, 25),
        "Instituição": semente["instituicao"].nome,
        "UF": "SP",
        "Modalidade": "pessoalmente",
    }
    folha.append([valores.get(coluna) for coluna in cabecalho])
    saida = io.BytesIO()
    pasta.save(saida)

    criada = cliente_admin.post(
        "/api/importacoes", files={"arquivo": ("dia.xlsx", saida.getvalue(), TIPO_XLSX)}
    ).json()

    resposta = cliente_admin.patch(
        f"/api/importacoes/{criada['id']}/resolucoes",
        json={"campo": "modalidade", "valor": "pessoalmente", "decisao": "criar"},
    )

    assert resposta.status_code == 422
    assert "fechad" in resposta.text.lower()


def test_alvo_malformado_recusa_com_422_e_nao_estoura(cliente_admin, importacao_com_quatro):
    """DEFEITO 3. `alvo` chega como texto e ia direto comparar com uma coluna
    UUID. "abc" viraria erro de banco — 500 e "erro interno" para quem só digitou
    algo inesperado num campo da API."""
    resposta = cliente_admin.patch(
        f"/api/importacoes/{importacao_com_quatro['id']}/resolucoes",
        json={
            "campo": "instituicao_id",
            "valor": "Prefeitura de Campinas",
            "decisao": "apontar",
            "alvo": "abc",
        },
    )

    assert resposta.status_code == 422, resposta.text


def test_alvo_nao_numerico_num_campo_de_id_inteiro_recusa(cliente_admin, sessao, semente):
    """O mesmo, do lado dos dicionários com id inteiro."""
    from datetime import date

    from openpyxl import load_workbook

    modelo = modelo_de_importacao.gerar(importar_agendas.vocabularios(sessao))
    pasta = load_workbook(io.BytesIO(modelo))
    folha = pasta["Agendas"]
    cabecalho = [celula.value for celula in next(folha.iter_rows())]
    valores = {
        "Data": date(2026, 9, 25),
        "Instituição": semente["instituicao"].nome,
        "UF": "SP",
        "Unidade de negócio": "Unidade Que Nao Existe",
    }
    folha.append([valores.get(coluna) for coluna in cabecalho])
    saida = io.BytesIO()
    pasta.save(saida)

    criada = cliente_admin.post(
        "/api/importacoes", files={"arquivo": ("dia.xlsx", saida.getvalue(), TIPO_XLSX)}
    ).json()

    resposta = cliente_admin.patch(
        f"/api/importacoes/{criada['id']}/resolucoes",
        json={
            "campo": "unidade_negocio_id",
            "valor": "Unidade Que Nao Existe",
            "decisao": "apontar",
            "alvo": "nao-e-numero",
        },
    )

    assert resposta.status_code == 422, resposta.text


def test_a_confirmacao_nao_faz_commit_por_conta_propria(cliente_admin, sessao, limpa):
    """OU TUDO ENTRA, OU NADA ENTRA — e é o que garante isso.

    A atomicidade não vem de um `try/except` dentro do caso de uso: vem de a
    transação ser a da requisição, comitada no teardown da dependência (ver
    `app/banco/sessao.py`). Um `sessao.commit()` no meio de `confirmar` criaria
    exatamente o estado que ela existe para impedir — instituições criadas e
    agendas não —, e nenhum teste de caminho feliz notaria.

    Este teste vigia a ausência de um commit. É a única forma de prender a
    invariante sem simular a falha, porque dentro da transação do próprio teste
    um rollback parcial não se distingue de um total.
    """
    chamadas = []
    original = sessao.commit
    sessao.commit = lambda *a, **k: chamadas.append(1) or original(*a, **k)
    try:
        resposta = cliente_admin.post(f"/api/importacoes/{limpa['id']}/confirmacao")
    finally:
        sessao.commit = original

    assert resposta.status_code == 201, resposta.text
    assert chamadas == [], "confirmar comitou por conta própria"


# =============================================================================
# os defeitos que a revisão do Codex achou na Tarefa 11
# =============================================================================


def _com_declaracao(sessao, semente, nome: str, tipo: str | None = None):
    """Uma agenda com instituição nova, declarada na aba com (ou sem) o tipo."""
    from datetime import date

    from openpyxl import load_workbook


    modelo = modelo_de_importacao.gerar(importar_agendas.vocabularios(sessao))
    pasta = load_workbook(io.BytesIO(modelo))
    folha = pasta["Agendas"]
    cabecalho = [celula.value for celula in next(folha.iter_rows())]
    valores = {"Data": date(2026, 9, 25), "Instituição": nome, "UF": "SP"}
    folha.append([valores.get(coluna) for coluna in cabecalho])
    _declarar(
        pasta,
        "instituicoes",
        **({"Instituição": nome, "Categoria de público": tipo} if tipo else {"Instituição": nome}),
    )
    saida = io.BytesIO()
    pasta.save(saida)
    return saida.getvalue()


def test_a_aba_de_instituicoes_pede_a_CATEGORIA(sessao):
    """DEFEITO 1. O tipo da instituição DERIVA A FRENTE da agenda — o próprio
    `derivar_frente` diz que "o tipo já basta, sozinho, para todos os tipos menos
    dois". Criar com um `orgao` adivinhado dava frente errada em toda agenda
    daquela instituição, contaminando toda leitura agrupada por frente. A planilha
    tem de perguntar."""
    from openpyxl import load_workbook

    pasta = load_workbook(
        io.BytesIO(modelo_de_importacao.gerar(importar_agendas.vocabularios(sessao)))
    )
    from app.dominio.importacao_de_agendas import ROTULO_DO_VOCABULARIO

    folha = pasta[ROTULO_DO_VOCABULARIO["instituicoes"]]
    cabecalho = [celula.value for celula in next(folha.iter_rows())]

    # A CATEGORIA ESTÁ NA ABA, e não numa posição fixa: a aba ganhou os campos do
    # formulário da plataforma, e fixar a posição faria este teste falhar por uma
    # coluna nova em vez de por um defeito.
    assert cabecalho[0] == "Instituição"
    assert "Categoria de público" in cabecalho


def test_a_instituicao_criada_DERIVA_o_tipo_da_categoria(cliente_admin, sessao, semente):
    """O TIPO NASCE DA CATEGORIA, como na tela de cadastro.

    `api/stakeholders.py` diz que "a tela de cadastro nao pergunta mais o tipo;
    ausente, ele vem da categoria de publico". Pedir o tipo direto na planilha
    divergiria disso e, pior, deixaria `categoria_publico_id` NULO — e é essa
    coluna que a taxonomia de públicos do Score usa para agrupar. A instituição
    importada ficaria invisível para uma área inteira do produto.
    """
    from sqlalchemy import select as sel

    conteudo = _com_declaracao(
        sessao, semente, "Valor Novo", tipo="Imprensa e Formadores de Opinião"
    )
    criada = cliente_admin.post(
        "/api/importacoes", files={"arquivo": ("a.xlsx", conteudo, TIPO_XLSX)}
    ).json()
    assert criada["pendencias"] == 0, criada["grupos"]

    cliente_admin.post(f"/api/importacoes/{criada['id']}/confirmacao")

    nova = sessao.scalars(
        sel(Instituicao).where(Instituicao.nome_normalizado == normalizar("Valor Novo"))
    ).first()
    assert nova is not None
    assert nova.tipo == "veiculo"
    assert nova.categoria_publico_id is not None


def test_instituicao_declarada_SEM_categoria_trava(cliente_admin, sessao, semente):
    """Sem o tipo não há como criar sem adivinhar, e adivinhar erra a frente.
    Travar devolve a decisão a quem sabe."""
    conteudo = _com_declaracao(sessao, semente, "Orgao Sem Tipo")

    criada = cliente_admin.post(
        "/api/importacoes", files={"arquivo": ("a.xlsx", conteudo, TIPO_XLSX)}
    ).json()

    assert criada["pendencias"] == 1
    mensagens = " ".join(
        d["mensagem"] for linha in criada["linhas"] for d in linha["divergencias"]
    ).lower()
    assert "categoria" in mensagens, mensagens


def test_instituicao_declarada_com_categoria_INVALIDA_trava(cliente_admin, sessao, semente):
    conteudo = _com_declaracao(sessao, semente, "Coisa", tipo="inventado")

    criada = cliente_admin.post(
        "/api/importacoes", files={"arquivo": ("a.xlsx", conteudo, TIPO_XLSX)}
    ).json()

    assert criada["pendencias"] == 1


def test_cadastro_INATIVO_nao_resolve(cliente_admin, sessao, semente):
    """DEFEITO 2. `_indice` não filtrava `ativo`, então uma linha limpa podia
    confirmar apontando para instituição desativada — divergindo das APIs
    normais, que filtram ativos, e ressuscitando na Base um cadastro que alguém
    tirou de circulação de propósito."""
    from datetime import date

    from openpyxl import load_workbook

    aposentada = Instituicao(
        nome="Jornal Extinto",
        nome_normalizado=normalizar("Jornal Extinto"),
        tipo="veiculo",
        uf="SP",
        ativo=False,
    )
    sessao.add(aposentada)
    sessao.flush()

    modelo = modelo_de_importacao.gerar(importar_agendas.vocabularios(sessao))
    pasta = load_workbook(io.BytesIO(modelo))
    folha = pasta["Agendas"]
    cabecalho = [celula.value for celula in next(folha.iter_rows())]
    valores = {
        "Data": date(2026, 9, 25),
        "Instituição": "Jornal Extinto",
        "UF": "SP",
    }
    folha.append([valores.get(coluna) for coluna in cabecalho])
    saida = io.BytesIO()
    pasta.save(saida)

    criada = cliente_admin.post(
        "/api/importacoes", files={"arquivo": ("a.xlsx", saida.getvalue(), TIPO_XLSX)}
    ).json()

    assert criada["pendencias"] == 1, criada["grupos"]


def test_o_modelo_nao_oferece_cadastro_inativo_na_lista(cliente_admin, sessao, semente):
    """A outra ponta: se o inativo não resolve, oferecê-lo na lista suspensa seria
    convidar a pessoa a escolher o que vai ser recusado."""
    sessao.add(
        Instituicao(
            nome="Jornal Extinto",
            nome_normalizado=normalizar("Jornal Extinto"),
            tipo="veiculo",
            uf="SP",
            ativo=False,
        )
    )
    sessao.flush()

    listas = importar_agendas.vocabularios(sessao)

    assert "Jornal Extinto" not in listas["instituicoes"]


def test_apontar_para_um_cadastro_INATIVO_recusa(cliente_admin, sessao, importacao_com_quatro):
    """E a reconferência também: um alvo desativado depois da decisão passava."""
    aposentada = Instituicao(
        nome="Prefeitura Extinta",
        nome_normalizado=normalizar("Prefeitura Extinta"),
        tipo="orgao",
        uf="SP",
        ativo=False,
    )
    sessao.add(aposentada)
    sessao.flush()

    resposta = cliente_admin.patch(
        f"/api/importacoes/{importacao_com_quatro['id']}/resolucoes",
        json={
            "campo": "instituicao_id",
            "valor": "Prefeitura de Campinas",
            "decisao": "apontar",
            "alvo": str(aposentada.id),
        },
    )

    assert resposta.status_code == 422


def test_dicionario_administrado_nao_promete_criar_e_trava_na_hora(
    cliente_admin, sessao, semente
):
    """DEFEITO 3. `unidades_negocio`, `formatos_interacao` e `areas_pessoa` são
    editáveis, então `_resolver` prometia "vou criar" — mas `_criar_cadastros` não
    sabe criá-los, e a recusa só aparecia no ÚLTIMO passo, depois de a pessoa ter
    conferido tudo. Eles são dicionário administrado, não cadastro livre: a
    divergência trava desde o upload, dizendo para cadastrar na Administração."""
    from datetime import date

    from openpyxl import load_workbook

    from app.dominio.importacao_de_agendas import ROTULO_DO_VOCABULARIO

    modelo = modelo_de_importacao.gerar(importar_agendas.vocabularios(sessao))
    pasta = load_workbook(io.BytesIO(modelo))
    folha = pasta["Agendas"]
    cabecalho = [celula.value for celula in next(folha.iter_rows())]
    valores = {
        "Data": date(2026, 9, 25),
        "Instituição": semente["instituicao"].nome,
        "UF": "SP",
        "Unidade de negócio": "Unidade Inventada",
    }
    folha.append([valores.get(coluna) for coluna in cabecalho])
    pasta[ROTULO_DO_VOCABULARIO["unidades_negocio"]].append(["Unidade Inventada"])
    saida = io.BytesIO()
    pasta.save(saida)

    criada = cliente_admin.post(
        "/api/importacoes", files={"arquivo": ("a.xlsx", saida.getvalue(), TIPO_XLSX)}
    ).json()

    assert criada["pendencias"] == 1, criada["grupos"]
    assert criada["a_criar"] == []
    mensagens = " ".join(
        d["mensagem"] for linha in criada["linhas"] for d in linha["divergencias"]
    )
    assert "Administração" in mensagens


# =============================================================================
# os achados Important da revisão final do Codex
# =============================================================================


def test_a_sugestao_traz_o_ALVO_que_o_servidor_aceita(cliente_admin, sessao, importacao_com_quatro):
    """DEFEITO 1 DA REVISÃO FINAL, e o pior de todos: a sugestão era só um NOME, a
    tela mandava o nome como `alvo`, e `_cadastro_existe` valida id ou código —
    então o atalho PRINCIPAL da conferência devolvia 422. A ação central da tela
    não funcionava, e nenhum teste cruzava os dois lados."""
    parecida = Instituicao(
        nome="Prefeitura Municipal de Campinas",
        nome_normalizado=normalizar("Prefeitura Municipal de Campinas"),
        tipo="orgao",
        uf="SP",
    )
    sessao.add(parecida)
    sessao.flush()

    estado = cliente_admin.get(f"/api/importacoes/{importacao_com_quatro['id']}").json()
    (grupo,) = [g for g in estado["grupos"] if g["campo"] == "instituicao_id"]

    assert grupo["sugestoes"], "sem sugestão não há atalho"
    sugestao = grupo["sugestoes"][0]
    assert sugestao["nome"] == "Prefeitura Municipal de Campinas"
    assert sugestao["alvo"] == str(parecida.id)

    # E o alvo que ela oferece é aceito de verdade — é o cruzamento que faltava.
    resposta = cliente_admin.patch(
        f"/api/importacoes/{importacao_com_quatro['id']}/resolucoes",
        json={
            "campo": grupo["campo"],
            "valor": grupo["valor"],
            "decisao": "apontar",
            "alvo": sugestao["alvo"],
        },
    )

    assert resposta.status_code == 200, resposta.text
    assert resposta.json()["pendencias"] == 0


def test_o_grupo_diz_se_da_para_CRIAR(cliente_admin, sessao, semente):
    """DEFEITO 3. A tela oferecia "Cadastrar como novo" em TODO grupo, e o servidor
    só barrava vocabulário fechado — então um dicionário administrado saía da
    pendência como `criar` e falhava na confirmação, depois de a pessoa ter
    conferido tudo. Agora o servidor diz, por grupo, se criar é possível."""
    from datetime import date

    from openpyxl import load_workbook

    modelo = modelo_de_importacao.gerar(importar_agendas.vocabularios(sessao))
    pasta = load_workbook(io.BytesIO(modelo))
    folha = pasta["Agendas"]
    cabecalho = [celula.value for celula in next(folha.iter_rows())]
    valores = {
        "Data": date(2026, 9, 25),
        "Instituição": "Orgao Novo",
        "UF": "SP",
        "Unidade de negócio": "Unidade Inventada",
    }
    folha.append([valores.get(coluna) for coluna in cabecalho])
    saida = io.BytesIO()
    pasta.save(saida)

    corpo = cliente_admin.post(
        "/api/importacoes", files={"arquivo": ("a.xlsx", saida.getvalue(), TIPO_XLSX)}
    ).json()

    por_campo = {g["campo"]: g for g in corpo["grupos"]}
    # Instituição é cadastro livre: a importação cria.
    assert por_campo["instituicao_id"]["pode_criar"] is True
    # Unidade de negócio é dicionário administrado: cadastra-se na Administração.
    assert por_campo["unidade_negocio_id"]["pode_criar"] is False


def test_criar_num_campo_SEM_vocabulario_recusa(cliente_admin, sessao, semente):
    """DEFEITO 3, a outra metade: uma data ilegível não tem vocabulário nenhum, e
    "criar" não significa nada ali. Sem a recusa, a decisão apagava a pendência e
    ela voltava como conflito na confirmação."""

    from openpyxl import load_workbook

    modelo = modelo_de_importacao.gerar(importar_agendas.vocabularios(sessao))
    pasta = load_workbook(io.BytesIO(modelo))
    folha = pasta["Agendas"]
    cabecalho = [celula.value for celula in next(folha.iter_rows())]
    valores = {
        "Data": "25 de setembro",
        "Instituição": semente["instituicao"].nome,
        "UF": "SP",
    }
    folha.append([valores.get(coluna) for coluna in cabecalho])
    saida = io.BytesIO()
    pasta.save(saida)

    criada = cliente_admin.post(
        "/api/importacoes", files={"arquivo": ("a.xlsx", saida.getvalue(), TIPO_XLSX)}
    ).json()
    (grupo,) = [g for g in criada["grupos"] if g["campo"] == "data_interacao"]
    assert grupo["pode_criar"] is False

    resposta = cliente_admin.patch(
        f"/api/importacoes/{criada['id']}/resolucoes",
        json={"campo": "data_interacao", "valor": grupo["valor"], "decisao": "criar"},
    )

    assert resposta.status_code == 422


def test_a_conferencia_mostra_o_que_foi_HERDADO(cliente_admin, sessao, semente):
    """O PROBLEMA QUE O DONO ACHOU AO TESTAR.

    A herança por `idem` é a única parte desta funcionalidade cujo resultado é
    INVISÍVEL antes de criar as agendas: a célula continua visualmente vazia na
    planilha, e a conferência mostrava só `dados_brutos` — o que foi digitado. A
    pessoa não tinha como saber, antes de confirmar, o que de fato seria gravado
    naquelas células.

    Cada linha passa a dizer também o que foi herdado, por coluna. É o que permite
    conferir uma linha de `idem` sem confiar na memória do que havia acima.
    """
    from datetime import date

    from openpyxl import load_workbook

    modelo = modelo_de_importacao.gerar(importar_agendas.vocabularios(sessao))
    pasta = load_workbook(io.BytesIO(modelo))
    folha = pasta["Agendas"]
    cabecalho = [celula.value for celula in next(folha.iter_rows())]
    primeira = {
        "Data": date(2026, 9, 25),
        "Instituição": semente["instituicao"].nome,
        "UF": "SP",
        "Local": "Sede, sala 3",
    }
    segunda = {COLUNA_DE_REPETICAO: VALOR_DA_REPETICAO}
    for valores in (primeira, segunda):
        folha.append([valores.get(coluna) for coluna in cabecalho])
    saida = io.BytesIO()
    pasta.save(saida)

    corpo = cliente_admin.post(
        "/api/importacoes", files={"arquivo": ("a.xlsx", saida.getvalue(), TIPO_XLSX)}
    ).json()

    por_linha = {linha["linha_origem"]: linha for linha in corpo["linhas"]}

    # A primeira não herdou nada — ela é a origem.
    assert por_linha[2]["herdado"] == {}

    # A segunda herdou o que deixou em branco, e a tela pode mostrar cada valor.
    herdado = por_linha[3]["herdado"]
    assert herdado["Data"] == "2026-09-25"
    assert herdado["Instituição"] == semente["instituicao"].nome
    assert herdado["UF"] == "SP"
    assert herdado["Local"] == "Sede, sala 3"


def test_o_que_a_pessoa_digitou_nao_entra_no_herdado(cliente_admin, sessao, semente):
    """O contrapeso: a tela precisa distinguir o que ela escreveu do que veio de
    cima. Sem isso, "herdado" viraria um despejo da linha inteira e não diria
    nada."""
    from datetime import date

    from openpyxl import load_workbook

    modelo = modelo_de_importacao.gerar(importar_agendas.vocabularios(sessao))
    pasta = load_workbook(io.BytesIO(modelo))
    folha = pasta["Agendas"]
    cabecalho = [celula.value for celula in next(folha.iter_rows())]
    for valores in (
        {
            "Data": date(2026, 9, 25),
            "Instituição": semente["instituicao"].nome,
            "UF": "SP",
        },
        {COLUNA_DE_REPETICAO: VALOR_DA_REPETICAO, "UF": "RJ"},
    ):
        folha.append([valores.get(coluna) for coluna in cabecalho])
    saida = io.BytesIO()
    pasta.save(saida)

    corpo = cliente_admin.post(
        "/api/importacoes", files={"arquivo": ("a.xlsx", saida.getvalue(), TIPO_XLSX)}
    ).json()

    herdado = {linha["linha_origem"]: linha["herdado"] for linha in corpo["linhas"]}[3]
    assert "UF" not in herdado, "UF foi digitada, não herdada"
    assert herdado["Instituição"] == semente["instituicao"].nome


def test_interlocutor_novo_declarado_VIRA_participante_na_confirmacao(
    cliente_admin, sessao, semente
):
    """ACHADO DA REVISÃO FINAL, e a aba única simplificou a correção.

    O interlocutor nascia só com nome, sem `instituicao_id`, e a recusa 4 — o
    participante tem de pertencer à instituição da agenda — o rejeitava NA
    CONFIRMAÇÃO: o upload dizia "vou criar", a pessoa conferia tudo, e o último
    passo devolvia conflito por uma pessoa que ela mesma declarou.

    COM TUDO NUMA ABA, a resposta está na própria linha: a instituição da agenda e
    o nome do interlocutor são vizinhos de coluna. Não há nada a inferir entre
    abas, nem uma coluna nova a pedir.
    """
    from datetime import date

    from openpyxl import load_workbook

    from app.banco.tabelas_interacoes import InteracaoInterlocutor

    modelo = modelo_de_importacao.gerar(importar_agendas.vocabularios(sessao))
    pasta = load_workbook(io.BytesIO(modelo))
    folha = pasta["Agendas"]
    cabecalho = [celula.value for celula in next(folha.iter_rows())]
    folha.append(
        [
            {
                "Data": date(2026, 9, 25),
                "Instituição": semente["instituicao"].nome,
                "UF": "SP",
                "Interlocutor 1": "Carla Nova",
                "Presença 1": "presente",
            }.get(coluna)
            for coluna in cabecalho
        ]
    )
    _declarar(pasta, "interlocutores", **{"Interlocutor": "Carla Nova"})
    saida = io.BytesIO()
    pasta.save(saida)

    criada = cliente_admin.post(
        "/api/importacoes", files={"arquivo": ("a.xlsx", saida.getvalue(), TIPO_XLSX)}
    ).json()
    assert criada["pendencias"] == 0, criada["grupos"]

    resposta = cliente_admin.post(f"/api/importacoes/{criada['id']}/confirmacao")

    assert resposta.status_code == 201, resposta.text
    assert resposta.json()["criadas"] == 1

    nova = sessao.scalars(
        select(Interlocutor).where(Interlocutor.nome_normalizado == normalizar("Carla Nova"))
    ).first()
    assert nova is not None
    assert nova.instituicao_id == semente["instituicao"].id

    ligacoes = sessao.scalars(
        select(InteracaoInterlocutor).where(InteracaoInterlocutor.interlocutor_id == nova.id)
    ).all()
    assert len(ligacoes) == 1


# =============================================================================
# o achado Alto da revisão da aba única
# =============================================================================


def test_a_instituicao_APONTADA_e_a_que_o_interlocutor_novo_recebe(
    cliente_admin, sessao, semente
):
    """ACHADO ALTO DO CODEX, e o defeito mais feio que a revisão achou.

    Com DUAS instituições de mesmo nome normalizado, o upload trava a linha por
    ambiguidade e a pessoa escolhe qual na conferência. A criação do interlocutor
    novo, porém, relia `dados_brutos["Instituição"]` contra o índice CRU — onde
    aquele nome vale `AMBIGUO`, um sentinela, não um id. O interlocutor nascia com
    o sentinela no lugar da chave estrangeira.

    O caminho todo já tinha corrido: a pessoa subiu, conferiu, escolheu a
    instituição certa e mandou confirmar. A decisão dela existia e estava gravada
    — só não era consultada nesse ponto."""
    from datetime import date

    from openpyxl import load_workbook

    from app.banco.tabelas_stakeholders import Instituicao

    # A HOMÔNIMA: mesmo nome normalizado, outra UF, outro cadastro. O banco
    # permite — `instituicao` é única por `(nome_normalizado, tipo)`.
    gemea = Instituicao(
        nome=semente["instituicao"].nome,
        nome_normalizado=semente["instituicao"].nome_normalizado,
        tipo="orgao",
        uf="RJ",
    )
    sessao.add(gemea)
    sessao.flush()

    modelo = modelo_de_importacao.gerar(importar_agendas.vocabularios(sessao))
    pasta = load_workbook(io.BytesIO(modelo))
    folha = pasta["Agendas"]
    cabecalho = [celula.value for celula in next(folha.iter_rows())]
    folha.append(
        [
            {
                "Data": date(2026, 9, 25),
                "Instituição": semente["instituicao"].nome,
                "UF": "SP",
                "Interlocutor 1": "Bruno Novo",
                "Presença 1": "presente",
            }.get(coluna)
            for coluna in cabecalho
        ]
    )
    _declarar(pasta, "interlocutores", **{"Interlocutor": "Bruno Novo"})
    saida = io.BytesIO()
    pasta.save(saida)

    criada = cliente_admin.post(
        "/api/importacoes", files={"arquivo": ("a.xlsx", saida.getvalue(), TIPO_XLSX)}
    ).json()
    # A linha trava: o nome da instituição não decide qual das duas é.
    assert criada["pendencias"] >= 1, criada["grupos"]

    # A pessoa escolhe a gêmea do Rio na conferência.
    resolucao = cliente_admin.patch(
        f"/api/importacoes/{criada['id']}/resolucoes",
        json={
            "campo": "instituicao_id",
            "valor": semente["instituicao"].nome,
            "decisao": "apontar",
            "alvo": str(gemea.id),
        },
    )
    assert resolucao.status_code == 200, resolucao.text

    confirmacao = cliente_admin.post(f"/api/importacoes/{criada['id']}/confirmacao")

    assert confirmacao.status_code == 201, confirmacao.text
    nova = sessao.scalars(
        select(Interlocutor).where(Interlocutor.nome_normalizado == normalizar("Bruno Novo"))
    ).first()
    assert nova is not None
    # O ID QUE A PESSOA ESCOLHEU, e não o sentinela nem a outra homônima.
    assert nova.instituicao_id == gemea.id


# =============================================================================
# a relação instituição -> interlocutor, no servidor
# =============================================================================


def _com_interlocutor(sessao, semente, instituicao_nome, pessoa, declarar=None):
    """Uma agenda naquela instituição, com aquela pessoa no `Interlocutor 1`.

    `declarar` é `(nome, instituicao)` escrito na aba de interlocutores — a
    declaração de intenção, agora com a instituição ao lado.
    """
    from datetime import date

    from openpyxl import load_workbook


    modelo = modelo_de_importacao.gerar(
        importar_agendas.vocabularios(sessao),
        importar_agendas.pares_de_vocabulario(sessao),
    )
    pasta = load_workbook(io.BytesIO(modelo))
    folha = pasta["Agendas"]
    cabecalho = [celula.value for celula in next(folha.iter_rows())]
    folha.append(
        [
            {
                "Data": date(2026, 9, 25),
                "Instituição": instituicao_nome,
                "UF": "SP",
                "Interlocutor 1": pessoa,
                "Presença 1": "presente",
            }.get(coluna)
            for coluna in cabecalho
        ]
    )
    if declarar:
        _declarar(
            pasta,
            "interlocutores",
            **{"Interlocutor": declarar[0], "Instituição": declarar[1]},
        )
    saida = io.BytesIO()
    pasta.save(saida)
    return saida.getvalue()


def _outra_instituicao(sessao, nome="Câmara Municipal", tipo="orgao"):
    from app.banco.tabelas_stakeholders import Instituicao

    instituicao = Instituicao(
        nome=nome, nome_normalizado=normalizar(nome), tipo=tipo, uf="SP"
    )
    sessao.add(instituicao)
    sessao.flush()
    return instituicao


def test_o_HOMONIMO_e_desfeito_pela_instituicao_da_linha(cliente_admin, sessao, semente):
    """O GANHO DE GRAÇA DA RELAÇÃO. "Assessoria da liderança" existe em vários
    órgãos — é cargo, não nome próprio —, e `interlocutor` é único por
    `(nome_normalizado, instituicao_id)`, então o banco permite o homônimo.

    Pelo nome sozinho, a importação travava: "existe mais de um cadastro com este
    nome, escolha qual". Com a instituição da linha, não há o que escolher: é a
    pessoa daquele órgão. Uma pendência que a pessoa não precisava resolver."""
    from app.banco.tabelas_interacoes import InteracaoInterlocutor
    from app.banco.tabelas_stakeholders import Interlocutor as Tabela

    outra = _outra_instituicao(sessao)
    for instituicao in (semente["instituicao"], outra):
        sessao.add(
            Tabela(
                nome="Assessoria da liderança",
                nome_normalizado=normalizar("Assessoria da liderança"),
                instituicao_id=instituicao.id,
            )
        )
    sessao.flush()

    conteudo = _com_interlocutor(
        sessao, semente, outra.nome, "Assessoria da liderança"
    )
    criada = cliente_admin.post(
        "/api/importacoes", files={"arquivo": ("a.xlsx", conteudo, TIPO_XLSX)}
    ).json()

    assert criada["pendencias"] == 0, criada["grupos"]
    confirmacao = cliente_admin.post(f"/api/importacoes/{criada['id']}/confirmacao")
    assert confirmacao.status_code == 201, confirmacao.text
    # A pessoa DAQUELE órgão, e não a homônima do outro.
    escolhida = sessao.scalars(
        select(Tabela).where(
            Tabela.nome_normalizado == normalizar("Assessoria da liderança"),
            Tabela.instituicao_id == outra.id,
        )
    ).first()
    ligacoes = sessao.scalars(
        select(InteracaoInterlocutor).where(
            InteracaoInterlocutor.interlocutor_id == escolhida.id
        )
    ).all()
    assert len(ligacoes) == 1


def test_interlocutor_de_OUTRA_instituicao_trava_e_diz_de_quem_ele_e(
    cliente_admin, sessao, semente
):
    """O ERRO QUE A SUSPENSA DEPENDENTE PREVINE, travado também no servidor — a
    planilha convertida para o Google Sheets pode perder a validação, e o arquivo
    pode nem ter saído do nosso modelo.

    "Ana Prado" é do Valor Econômico. Usada numa agenda da Câmara, ela violaria a
    regra de que o participante pertence à instituição da agenda — e essa regra
    recusava na CONFIRMAÇÃO, depois de a pessoa ter conferido tudo. Agora recusa no
    upload, junto das outras pendências, dizendo de quem a pessoa é."""
    outra = _outra_instituicao(sessao)

    conteudo = _com_interlocutor(sessao, semente, outra.nome, "Ana Prado")
    criada = cliente_admin.post(
        "/api/importacoes", files={"arquivo": ("a.xlsx", conteudo, TIPO_XLSX)}
    ).json()

    assert criada["pendencias"] >= 1
    # ONDE A PESSOA LÊ: a tabela de linhas da conferência mostra a mensagem de cada
    # divergência (`ConferirImportacao.tsx`); o bloco agrupado monta o próprio
    # cabeçalho a partir de campo e valor.
    mensagens = " ".join(
        divergencia["mensagem"]
        for linha in criada["linhas"]
        for divergencia in linha["divergencias"]
    )
    assert "Ana Prado" in mensagens
    # A mensagem diz de QUEM ela é, que é o que permite consertar sem procurar.
    assert semente["instituicao"].nome in mensagens


def test_a_instituicao_DECLARADA_ao_lado_precisa_bater_com_a_da_agenda(
    cliente_admin, sessao, semente
):
    """A contradição declarada, pega no upload.

    Declarar "Carla Nova — Câmara Municipal" e usá-la numa agenda do Valor
    Econômico é pedir duas coisas incompatíveis: a regra 4 diz que o participante
    pertence à instituição da agenda. Criar a pessoa em um dos dois órgãos e
    esperar que dê certo é o tipo de chute que erra calado."""
    _outra_instituicao(sessao)

    conteudo = _com_interlocutor(
        sessao,
        semente,
        semente["instituicao"].nome,
        "Carla Nova",
        declarar=("Carla Nova", "Câmara Municipal"),
    )
    criada = cliente_admin.post(
        "/api/importacoes", files={"arquivo": ("a.xlsx", conteudo, TIPO_XLSX)}
    ).json()

    assert criada["pendencias"] >= 1, criada["grupos"]
    mensagens = " ".join(
        divergencia["mensagem"]
        for linha in criada["linhas"]
        for divergencia in linha["divergencias"]
    )
    assert "Carla Nova" in mensagens
    assert "uma instituição só" in mensagens


def test_a_instituicao_declarada_IGUAL_a_da_agenda_cria_a_pessoa_nela(
    cliente_admin, sessao, semente
):
    """O caminho feliz da declaração: o órgão escrito ao lado é o da agenda, e a
    pessoa nasce nele. Escrever o órgão é o que a suspensa da coluna B oferece, e
    confirmá-lo aqui é o que torna a coluna útil em vez de decorativa."""
    conteudo = _com_interlocutor(
        sessao,
        semente,
        semente["instituicao"].nome,
        "Carla Nova",
        declarar=("Carla Nova", semente["instituicao"].nome),
    )
    criada = cliente_admin.post(
        "/api/importacoes", files={"arquivo": ("a.xlsx", conteudo, TIPO_XLSX)}
    ).json()
    assert criada["pendencias"] == 0, criada["grupos"]

    confirmacao = cliente_admin.post(f"/api/importacoes/{criada['id']}/confirmacao")

    assert confirmacao.status_code == 201, confirmacao.text
    nova = sessao.scalars(
        select(Interlocutor).where(Interlocutor.nome_normalizado == normalizar("Carla Nova"))
    ).first()
    assert nova.instituicao_id == semente["instituicao"].id


# =============================================================================
# completar na tela a informação que falta na linha
# =============================================================================


def _sem_data(sessao, semente):
    """Uma agenda completa e uma SEM A DATA — a falta que nenhuma decisão resolve.

    O grupo de divergência de uma AUSÊNCIA não tem valor: não há nome errado para
    apontar nem cadastro para criar. Antes desta rota, a única saída era corrigir a
    planilha e subir de novo, ou descartar a linha.
    """
    from datetime import date

    from openpyxl import load_workbook

    modelo = modelo_de_importacao.gerar(
        importar_agendas.vocabularios(sessao),
        importar_agendas.pares_de_vocabulario(sessao),
    )
    pasta = load_workbook(io.BytesIO(modelo))
    folha = pasta["Agendas"]
    cabecalho = [celula.value for celula in next(folha.iter_rows())]
    for valores in (
        {
            "Data": date(2026, 9, 25),
            "Instituição": semente["instituicao"].nome,
            "UF": "SP",
        },
        {"Instituição": semente["instituicao"].nome, "UF": "SP"},
    ):
        folha.append([valores.get(coluna) for coluna in cabecalho])
    saida = io.BytesIO()
    pasta.save(saida)
    return saida.getvalue()


def test_a_divergencia_diz_QUAL_COLUNA_a_pessoa_tem_de_preencher():
    """A tela precisa saber onde oferecer o campo, e `campo` não responde isso:
    `data_interacao` é o nome interno, e a pessoa procura "Data" na planilha.

    Sem a coluna na divergência, a tela teria de adivinhar a partir do texto da
    mensagem — e `outra_parte.interlocutor_id` corresponde a QUATRO colunas."""
    from app.dominio.importacao_de_agendas import Divergencia

    assert Divergencia(campo="x", valor="y", mensagem="z", trava=True).coluna == ""


def test_a_mensagem_da_falta_usa_o_ROTULO_da_coluna(cliente_admin, sessao, semente):
    """"Falta data interacao" é o nome do campo no código. A pessoa lê "Data" no
    cabeçalho da planilha, e é esse nome que ela procura para consertar."""
    criada = cliente_admin.post(
        "/api/importacoes",
        files={"arquivo": ("a.xlsx", _sem_data(sessao, semente), TIPO_XLSX)},
    ).json()

    (faltando,) = [
        divergencia
        for linha in criada["linhas"]
        for divergencia in linha["divergencias"]
        if divergencia["trava"]
    ]

    assert "Data" in faltando["mensagem"]
    assert "data interacao" not in faltando["mensagem"]
    # E a coluna vem separada, para a tela saber onde pôr o campo.
    assert faltando["coluna"] == "Data"


def test_preencher_a_celula_na_tela_RESOLVE_a_linha(cliente_admin, sessao, semente):
    """O PEDIDO: completar a informação sem voltar à planilha.

    A linha é reproposta com o valor novo — pela MESMA máquina do upload, não por
    uma segunda implementação —, e a pendência desaparece."""
    criada = cliente_admin.post(
        "/api/importacoes",
        files={"arquivo": ("a.xlsx", _sem_data(sessao, semente), TIPO_XLSX)},
    ).json()
    assert criada["pendencias"] == 1
    presa = next(
        linha for linha in criada["linhas"] if any(d["trava"] for d in linha["divergencias"])
    )

    resposta = cliente_admin.patch(
        f"/api/importacoes/{criada['id']}/linhas/{presa['id']}",
        json={"celulas": {"Data": "26/09/2026"}},
    )

    assert resposta.status_code == 200, resposta.text
    depois = resposta.json()
    assert depois["pendencias"] == 0, [
        d["mensagem"] for linha in depois["linhas"] for d in linha["divergencias"]
    ]
    # E confirma: a agenda nasce com a data que a pessoa digitou na tela.
    confirmacao = cliente_admin.post(f"/api/importacoes/{criada['id']}/confirmacao")
    assert confirmacao.status_code == 201, confirmacao.text
    assert confirmacao.json()["criadas"] == 2


def test_a_celula_corrigida_FICA_MARCADA_como_editada(cliente_admin, sessao, semente):
    """A TELA TEM DE DIZER, e é a contrapartida honesta de editar aqui: o registro
    passa a divergir da planilha que a pessoa guardou. Sem a marca, ela abriria o
    arquivo meses depois para entender uma agenda e encontraria a célula vazia."""
    criada = cliente_admin.post(
        "/api/importacoes",
        files={"arquivo": ("a.xlsx", _sem_data(sessao, semente), TIPO_XLSX)},
    ).json()
    presa = next(
        linha for linha in criada["linhas"] if any(d["trava"] for d in linha["divergencias"])
    )

    depois = cliente_admin.patch(
        f"/api/importacoes/{criada['id']}/linhas/{presa['id']}",
        json={"celulas": {"Data": "26/09/2026"}},
    ).json()

    editada = next(linha for linha in depois["linhas"] if linha["id"] == presa["id"])
    assert editada["corrigido"] == {"Data": "26/09/2026"}
    # A linha que ninguém tocou não ganha marca nenhuma.
    outra = next(linha for linha in depois["linhas"] if linha["id"] != presa["id"])
    assert outra["corrigido"] == {}


def test_coluna_que_nao_existe_RECUSA_em_vez_de_nao_fazer_nada(
    cliente_admin, sessao, semente
):
    """Uma coluna escrita errado gravaria uma chave que ninguém lê: a pessoa
    clicaria em salvar, veria a pendência continuar e não teria como saber por quê.
    """
    criada = cliente_admin.post(
        "/api/importacoes",
        files={"arquivo": ("a.xlsx", _sem_data(sessao, semente), TIPO_XLSX)},
    ).json()
    presa = next(
        linha for linha in criada["linhas"] if any(d["trava"] for d in linha["divergencias"])
    )

    resposta = cliente_admin.patch(
        f"/api/importacoes/{criada['id']}/linhas/{presa['id']}",
        json={"celulas": {"Datta": "26/09/2026"}},
    )

    assert resposta.status_code == 422
    assert "Datta" in resposta.json()["detalhe"]


def test_nao_se_edita_linha_de_importacao_ja_CONFIRMADA(cliente_admin, sessao, semente):
    """Depois de confirmada, a linha já virou agenda: editá-la aqui não mudaria a
    agenda nenhuma, e daria à pessoa a impressão de ter corrigido algo."""
    criada = cliente_admin.post(
        "/api/importacoes",
        files={"arquivo": ("a.xlsx", _sem_data(sessao, semente), TIPO_XLSX)},
    ).json()
    presa = next(
        linha for linha in criada["linhas"] if any(d["trava"] for d in linha["divergencias"])
    )
    cliente_admin.patch(
        f"/api/importacoes/{criada['id']}/linhas/{presa['id']}",
        json={"celulas": {"Data": "26/09/2026"}},
    )
    cliente_admin.post(f"/api/importacoes/{criada['id']}/confirmacao")

    resposta = cliente_admin.patch(
        f"/api/importacoes/{criada['id']}/linhas/{presa['id']}",
        json={"celulas": {"Data": "27/09/2026"}},
    )

    assert resposta.status_code == 422


def test_TODO_campo_da_divergencia_atravessa_a_gravacao():
    """A GRAVAÇÃO É ESCRITA À MÃO, campo por campo, e é onde um campo novo se perde.

    Já aconteceu: `acao` e `alvo` faltavam, e a declaração feita na aba editável
    reaparecia como pendência em vez de ir para "o que vou criar". O defeito é
    silencioso — nada quebra, o valor simplesmente não chega à tela.

    Esta guarda compara a lista gravada com os campos VIVOS da dataclass, então um
    campo novo em `Divergencia` falha aqui em vez de desaparecer em produção."""
    import re
    from dataclasses import fields
    from pathlib import Path

    from app.dominio.importacao_de_agendas import Divergencia

    fonte = Path("app/api/importacoes.py").read_text(encoding="utf-8")
    trecho = fonte[fonte.index("divergencias=[") : fonte.index("for divergencia in proposta")]
    # `sugestoes` é gravado como `list(divergencia.sugestoes)`: o padrão aceita
    # qualquer envelope em volta, e não só a referência nua.
    gravados = set(re.findall(r'"(\w+)":[^,\n]*divergencia\.', trecho))

    assert {campo.name for campo in fields(Divergencia)} == gravados


# =============================================================================
# escolher o modelo no download, e as colunas na saída
# =============================================================================


def test_o_download_entrega_o_COMPLETO_por_padrao(cliente_admin, sessao):
    from openpyxl import load_workbook

    from app.dominio.importacao_de_agendas import MODELOS

    resposta = cliente_admin.get("/api/importacoes/modelo")

    assert resposta.status_code == 200
    folha = load_workbook(io.BytesIO(resposta.content))["Agendas"]
    cabecalho = [celula.value for celula in next(folha.iter_rows())]
    assert cabecalho == list(MODELOS["completo"])


def test_o_download_entrega_o_SIMPLIFICADO_quando_pedido(cliente_admin, sessao):
    """O modelo do evento: 54 agendas no mesmo dia, 22 colunas em vez de 58."""
    from openpyxl import load_workbook

    from app.dominio.importacao_de_agendas import MODELOS

    resposta = cliente_admin.get("/api/importacoes/modelo?modelo=simplificado")

    assert resposta.status_code == 200
    folha = load_workbook(io.BytesIO(resposta.content))["Agendas"]
    cabecalho = [celula.value for celula in next(folha.iter_rows())]
    assert cabecalho == list(MODELOS["simplificado"])


def test_o_nome_do_arquivo_baixado_diz_QUAL_modelo(cliente_admin, sessao):
    """Duas planilhas na pasta de downloads com o mesmo nome viram
    "modelo (1).xlsx", e a pessoa abre a errada."""
    resposta = cliente_admin.get("/api/importacoes/modelo?modelo=simplificado")

    assert "simplificado" in resposta.headers["content-disposition"]


def test_modelo_invalido_no_download_recusa_dizendo_os_validos(cliente_admin, sessao):
    resposta = cliente_admin.get("/api/importacoes/modelo?modelo=resumido")

    assert resposta.status_code == 422
    assert "simplificado" in resposta.json()["detalhe"]


def test_a_conferencia_diz_as_COLUNAS_daquele_arquivo(cliente_admin, sessao, semente):
    """A GRADE PRECISA DA ORDEM. `dados_brutos` é um objeto, e a tela não pode
    depender da ordem de um objeto JSON para montar o cabeçalho de uma tabela.

    E as colunas são as DAQUELE arquivo: quem subiu o simplificado não deve ver 58
    colunas vazias na conferência."""
    from app.dominio.importacao_de_agendas import MODELOS

    conteudo = _planilha_de_um_dia(sessao, semente, quantas=1)
    resposta = cliente_admin.post(
        "/api/importacoes", files={"arquivo": ("a.xlsx", conteudo, TIPO_XLSX)}
    )
    assert resposta.status_code == 201, resposta.text
    criada = resposta.json()

    assert [coluna["nome"] for coluna in criada["colunas"]] == list(MODELOS["completo"])


def test_as_colunas_da_conferencia_seguem_o_arquivo_SIMPLIFICADO(
    cliente_admin, sessao, semente
):
    from datetime import date

    from openpyxl import load_workbook

    from app.dominio.importacao_de_agendas import MODELOS

    modelo = modelo_de_importacao.gerar(
        importar_agendas.vocabularios(sessao),
        importar_agendas.pares_de_vocabulario(sessao),
        modelo="simplificado",
    )
    pasta = load_workbook(io.BytesIO(modelo))
    folha = pasta["Agendas"]
    cabecalho = [celula.value for celula in next(folha.iter_rows())]
    folha.append(
        [
            {
                "Data": date(2026, 9, 25),
                "Instituição": semente["instituicao"].nome,
                "UF": "SP",
            }.get(coluna)
            for coluna in cabecalho
        ]
    )
    saida = io.BytesIO()
    pasta.save(saida)

    criada = cliente_admin.post(
        "/api/importacoes", files={"arquivo": ("s.xlsx", saida.getvalue(), TIPO_XLSX)}
    ).json()

    assert [coluna["nome"] for coluna in criada["colunas"]] == list(
        MODELOS["simplificado"]
    )
    assert criada["pendencias"] == 0, criada["grupos"]


# =============================================================================
# a coluna na divergência: é o que pinta a célula na grade
# =============================================================================


def test_a_divergencia_de_VOCABULARIO_tambem_diz_a_coluna(cliente_admin, sessao, semente):
    """O CASO MAIS COMUM DE TODOS, e era o que ficava sem cor.

    "Este órgão não existe no cadastro" é a divergência que a conferência mais
    mostra, e ela nasce de UMA célula. Sem a coluna, a grade não tem onde pintar —
    e a pessoa recebe "2 pendências" olhando uma tabela sem nada destacado."""
    conteudo = _com_declaracao(sessao, semente, "Órgão Inventado")
    criada = cliente_admin.post(
        "/api/importacoes", files={"arquivo": ("a.xlsx", conteudo, TIPO_XLSX)}
    ).json()

    de_instituicao = [
        divergencia
        for linha in criada["linhas"]
        for divergencia in linha["divergencias"]
        if divergencia["campo"] == "instituicao_id"
    ]

    assert de_instituicao, criada["linhas"]
    assert all(d["coluna"] == "Instituição" for d in de_instituicao), de_instituicao


def test_TODA_divergencia_de_celula_diz_a_coluna(cliente_admin, sessao, semente):
    """A GUARDA GERAL. Um arquivo com problemas de tipos diferentes numa linha só:
    cada divergência que veio de uma célula tem de saber qual é.

    As que NÃO têm coluna são as da linha inteira — a duplicata de agenda —, e
    essas continuam sem, de propósito: pintar uma célula arbitrária mandaria a
    pessoa consertar uma coluna que não tem nada de errado.
    """
    from datetime import date

    from openpyxl import load_workbook

    modelo = modelo_de_importacao.gerar(
        importar_agendas.vocabularios(sessao),
        importar_agendas.pares_de_vocabulario(sessao),
    )
    pasta = load_workbook(io.BytesIO(modelo))
    folha = pasta["Agendas"]
    cabecalho = [celula.value for celula in next(folha.iter_rows())]
    folha.append(
        [
            {
                "Data": date(2026, 9, 25),
                "Instituição": "Órgão Que Ninguém Cadastrou",
                "UF": "SP",
                "Clima": "eufórico",
                "Interlocutor 1": "Pessoa Que Ninguém Cadastrou",
                "Tema 1": "Tema Inexistente",
            }.get(coluna)
            for coluna in cabecalho
        ]
    )
    saida = io.BytesIO()
    pasta.save(saida)

    criada = cliente_admin.post(
        "/api/importacoes", files={"arquivo": ("a.xlsx", saida.getvalue(), TIPO_XLSX)}
    ).json()

    sem_coluna = [
        (d["campo"], d["mensagem"][:60])
        for linha in criada["linhas"]
        for d in linha["divergencias"]
        # A duplicata é da linha inteira e não tem coluna: ela é a única exceção,
        # e aqui não há duplicata nenhuma.
        if not d["coluna"]
    ]

    assert sem_coluna == []


# =============================================================================
# os dois achados da revisão dos dois modelos
# =============================================================================


# =============================================================================
# excluir e restaurar uma linha
# =============================================================================


def test_excluir_UMA_linha_tira_ela_da_confirmacao(cliente_admin, sessao, semente):
    """O PEDIDO: dois botões por linha, editar e excluir.

    Excluir é a saída para a linha que não deveria estar ali — a duplicata colada
    por engano, o rascunho no fim do arquivo. Antes só existia "descartar" no nível
    do GRUPO: descartava todas as linhas que compartilhavam um valor, o que é outra
    coisa e às vezes é demais.
    """
    criada = cliente_admin.post(
        "/api/importacoes",
        files={"arquivo": ("a.xlsx", _planilha_de_um_dia(sessao, semente, quantas=3), TIPO_XLSX)},
    ).json()
    (primeira, *_) = criada["linhas"]

    resposta = cliente_admin.patch(
        f"/api/importacoes/{criada['id']}/linhas/{primeira['id']}",
        json={"descartada": True},
    )

    assert resposta.status_code == 200, resposta.text
    depois = resposta.json()
    excluida = next(linha for linha in depois["linhas"] if linha["id"] == primeira["id"])
    assert excluida["decisao"] == "descartada"

    confirmacao = cliente_admin.post(f"/api/importacoes/{criada['id']}/confirmacao")
    assert confirmacao.status_code == 201, confirmacao.text
    # Três linhas no arquivo, uma excluída: duas agendas.
    assert confirmacao.json()["criadas"] == 2


def test_a_linha_excluida_pode_ser_RESTAURADA(cliente_admin, sessao, semente):
    """NADA SE PERDE ANTES DE CONFIRMAR. Excluir por engano numa tela de 54 linhas é
    fácil, e a planilha não é a fonte de volta — o arquivo não fica guardado.

    A linha excluída continua na grade, marcada, e volta com um clique."""
    criada = cliente_admin.post(
        "/api/importacoes",
        files={"arquivo": ("a.xlsx", _planilha_de_um_dia(sessao, semente, quantas=2), TIPO_XLSX)},
    ).json()
    (primeira, *_) = criada["linhas"]
    cliente_admin.patch(
        f"/api/importacoes/{criada['id']}/linhas/{primeira['id']}",
        json={"descartada": True},
    )

    depois = cliente_admin.patch(
        f"/api/importacoes/{criada['id']}/linhas/{primeira['id']}",
        json={"descartada": False},
    ).json()

    restaurada = next(linha for linha in depois["linhas"] if linha["id"] == primeira["id"])
    assert restaurada["decisao"] == "pendente"
    confirmacao = cliente_admin.post(f"/api/importacoes/{criada['id']}/confirmacao")
    assert confirmacao.json()["criadas"] == 2


def test_excluir_a_linha_TIRA_a_pendencia_dela(cliente_admin, sessao, semente):
    """É a segunda razão de excluir existir por linha: a linha presa que a pessoa
    decide não importar não pode continuar segurando a confirmação das outras 53."""
    criada = cliente_admin.post(
        "/api/importacoes",
        files={
            "arquivo": (
                "a.xlsx",
                _com_declaracao(sessao, semente, "Órgão Inventado"),
                TIPO_XLSX,
            )
        },
    ).json()
    assert criada["pendencias"] == 1
    (presa,) = [
        linha for linha in criada["linhas"] if any(d["trava"] for d in linha["divergencias"])
    ]

    depois = cliente_admin.patch(
        f"/api/importacoes/{criada['id']}/linhas/{presa['id']}",
        json={"descartada": True},
    ).json()

    assert depois["pendencias"] == 0, depois["grupos"]


def test_editar_e_excluir_na_MESMA_chamada_recusa(cliente_admin, sessao, semente):
    """As duas coisas juntas são contraditórias: preencher uma célula de uma linha
    que não vai entrar. Recusar é melhor que escolher uma das duas em silêncio."""
    criada = cliente_admin.post(
        "/api/importacoes",
        files={"arquivo": ("a.xlsx", _planilha_de_um_dia(sessao, semente, quantas=1), TIPO_XLSX)},
    ).json()
    (primeira, *_) = criada["linhas"]

    resposta = cliente_admin.patch(
        f"/api/importacoes/{criada['id']}/linhas/{primeira['id']}",
        json={"descartada": True, "celulas": {"UF": "RJ"}},
    )

    assert resposta.status_code == 422


def test_a_conferencia_diz_o_TIPO_de_cada_coluna(cliente_admin, sessao, semente):
    """A TELA NÃO PODE CHUTAR LARGURA. "Relato" guarda um parágrafo e "UF" guarda
    duas letras: a mesma largura nas duas desperdiça a tela numa e trunca a outra.

    O tipo vem do FORMATO, que é quem sabe — e não de uma lista de 59 nomes no
    front, que envelheceria na primeira coluna nova sem ninguém perceber."""
    criada = cliente_admin.post(
        "/api/importacoes",
        files={
            "arquivo": ("a.xlsx", _planilha_de_um_dia(sessao, semente, quantas=1), TIPO_XLSX)
        },
    ).json()

    tipos = {coluna["nome"]: coluna["tipo"] for coluna in criada["colunas"]}

    assert tipos["Repetir a linha de cima"] == "marca"
    assert tipos["Data"] == "data"
    assert tipos["UF"] == "sigla"
    assert tipos["Instituição"] == "lista"
    assert tipos["Relato"] == "prosa"
    assert tipos["Local"] == "texto"


def test_instituicao_E_interlocutor_NOVOS_no_mesmo_arquivo(cliente_admin, sessao, semente):
    """O QUE O DONO DO PRODUTO QUER FAZER: declarar a instituição uma vez, e o
    interlocutor dela ao lado, sem preencher o mesmo dado duas vezes.

    É O CASO MAIS DIFÍCIL DA CRIAÇÃO EM CASCATA: o interlocutor precisa do id de uma
    instituição que ainda não existe quando a linha é lida. A ordem em `_criar_cadastros`
    existe para isto — instituições primeiro, um `flush`, e o índice relido — e este
    teste é o que prova que ela funciona de ponta a ponta."""
    from datetime import date

    from openpyxl import load_workbook

    from app.banco.tabelas_stakeholders import Instituicao

    modelo = modelo_de_importacao.gerar(
        importar_agendas.vocabularios(sessao),
        importar_agendas.pares_de_vocabulario(sessao),
    )
    pasta = load_workbook(io.BytesIO(modelo))
    folha = pasta["Agendas"]
    cabecalho = [celula.value for celula in next(folha.iter_rows())]
    folha.append(
        [
            {
                "Data": date(2026, 9, 27),
                "Instituição": "Instituto Novo",
                "UF": "SP",
                "Interlocutor 1": "Pessoa Nova",
            }.get(coluna)
            for coluna in cabecalho
        ]
    )
    # Declarados nas abas: a instituição com a categoria, e a pessoa com a instituição.
    _declarar(
        pasta,
        "instituicoes",
        **{"Instituição": "Instituto Novo", "Categoria de público": "Poder Executivo"},
    )
    _declarar(
        pasta,
        "interlocutores",
        **{"Interlocutor": "Pessoa Nova", "Instituição": "Instituto Novo"},
    )
    saida = io.BytesIO()
    pasta.save(saida)

    criada = cliente_admin.post(
        "/api/importacoes", files={"arquivo": ("a.xlsx", saida.getvalue(), TIPO_XLSX)}
    ).json()
    assert criada["pendencias"] == 0, [
        d["mensagem"] for linha in criada["linhas"] for d in linha["divergencias"]
    ]

    confirmacao = cliente_admin.post(f"/api/importacoes/{criada['id']}/confirmacao")

    assert confirmacao.status_code == 201, confirmacao.text
    instituicao = sessao.scalars(
        select(Instituicao).where(Instituicao.nome_normalizado == normalizar("Instituto Novo"))
    ).first()
    assert instituicao is not None
    pessoa = sessao.scalars(
        select(Interlocutor).where(Interlocutor.nome_normalizado == normalizar("Pessoa Nova"))
    ).first()
    # A PESSOA NASCE NA INSTITUIÇÃO QUE NASCEU AGORA — é o vínculo que o dono não quer
    # ter de digitar duas vezes.
    assert pessoa is not None
    assert pessoa.instituicao_id == instituicao.id


def test_o_cadastro_nasce_com_TODOS_os_campos_que_ela_preencheu(
    cliente_admin, sessao, semente
):
    """O PENTE FINO, de ponta a ponta. As abas ganharam os campos do formulário da
    plataforma, e o que importa é que eles CHEGUEM ao banco — uma aba com colunas
    bonitas que o servidor ignora é pior que não ter as colunas, porque a pessoa
    preenche e não desconfia.

    A ABRANGÊNCIA É O CASO QUE O DONO DO PRODUTO LEVANTOU: ela não é a UF da agenda. A
    reunião é em SP e o órgão é de MG — antes disto a instituição nascia com `NA` fixo,
    e alguém teria de completar na Administração depois.
    """
    from datetime import date

    from openpyxl import load_workbook

    from app.banco.tabelas_catalogo import SubcategoriaPublico
    from app.banco.tabelas_stakeholders import Instituicao

    modelo = modelo_de_importacao.gerar(
        importar_agendas.vocabularios(sessao),
        importar_agendas.pares_de_vocabulario(sessao),
    )
    pasta = load_workbook(io.BytesIO(modelo))
    folha = pasta["Agendas"]
    cabecalho = [celula.value for celula in next(folha.iter_rows())]
    folha.append(
        [
            {
                "Data": date(2026, 9, 29),
                "Instituição": "Secretaria Nova",
                # A UF DA AGENDA é SP: a reunião aconteceu aqui.
                "UF": "SP",
                "Interlocutor 1": "Pessoa Nova",
            }.get(coluna)
            for coluna in cabecalho
        ]
    )
    _declarar(
        pasta,
        "instituicoes",
        **{
            "Instituição": "Secretaria Nova",
            "Nome completo": "Secretaria Nova de Meio Ambiente",
            # A ABRANGÊNCIA DO ÓRGÃO é MG: ele é de Minas.
            "Abrangência": "MG",
            "Categoria de público": "Poder Executivo",
        },
    )
    _declarar(
        pasta,
        "interlocutores",
        **{
            "Interlocutor": "Pessoa Nova",
            "Instituição": "Secretaria Nova",
            "Cargo": "Secretária adjunta",
            "E-mail": "pessoa.nova@mg.gov.br",
        },
    )
    saida = io.BytesIO()
    pasta.save(saida)

    criada = cliente_admin.post(
        "/api/importacoes", files={"arquivo": ("a.xlsx", saida.getvalue(), TIPO_XLSX)}
    ).json()
    assert criada["pendencias"] == 0, [
        d["mensagem"] for linha in criada["linhas"] for d in linha["divergencias"]
    ]

    confirmacao = cliente_admin.post(f"/api/importacoes/{criada['id']}/confirmacao")
    assert confirmacao.status_code == 201, confirmacao.text

    instituicao = sessao.scalars(
        select(Instituicao).where(Instituicao.nome_normalizado == normalizar("Secretaria Nova"))
    ).first()
    assert instituicao is not None
    assert instituicao.nome_completo == "Secretaria Nova de Meio Ambiente"
    # A ABRANGÊNCIA DO ÓRGÃO, e não a UF da agenda.
    assert instituicao.uf == "MG"
    assert instituicao.tipo == "orgao"

    pessoa = sessao.scalars(
        select(Interlocutor).where(Interlocutor.nome_normalizado == normalizar("Pessoa Nova"))
    ).first()
    assert pessoa is not None
    assert pessoa.cargo == "Secretária adjunta"
    assert pessoa.email == "pessoa.nova@mg.gov.br"
    assert pessoa.instituicao_id == instituicao.id
    assert SubcategoriaPublico is not None  # o modelo existe; a subcategoria é opcional


def test_a_SUBCATEGORIA_declarada_chega_ao_banco(cliente_admin, sessao, semente):
    """A subcategoria é o campo que exige tradução: a pessoa escreve o NOME e a coluna
    guarda o id. Sem a tradução, o cadastro nasceria sem ela e ninguém veria."""
    from datetime import date

    from openpyxl import load_workbook

    from app.banco.tabelas_catalogo import CategoriaPublico, SubcategoriaPublico
    from app.banco.tabelas_stakeholders import Instituicao

    sub = sessao.scalars(
        select(SubcategoriaPublico)
        .join(CategoriaPublico, CategoriaPublico.id == SubcategoriaPublico.categoria_publico_id)
        .where(SubcategoriaPublico.ativo.is_(True))
        .limit(1)
    ).first()
    if sub is None:
        import pytest

        pytest.skip("a base de teste não tem subcategoria cadastrada")
    categoria = sessao.get(CategoriaPublico, sub.categoria_publico_id)

    modelo = modelo_de_importacao.gerar(
        importar_agendas.vocabularios(sessao),
        importar_agendas.pares_de_vocabulario(sessao),
    )
    pasta = load_workbook(io.BytesIO(modelo))
    folha = pasta["Agendas"]
    cabecalho = [celula.value for celula in next(folha.iter_rows())]
    folha.append(
        [
            {
                "Data": date(2026, 9, 29),
                "Instituição": "Órgão Com Subcategoria",
                "UF": "SP",
            }.get(coluna)
            for coluna in cabecalho
        ]
    )
    _declarar(
        pasta,
        "instituicoes",
        **{
            "Instituição": "Órgão Com Subcategoria",
            "Categoria de público": categoria.nome,
            "Subcategoria": sub.nome,
        },
    )
    saida = io.BytesIO()
    pasta.save(saida)

    criada = cliente_admin.post(
        "/api/importacoes", files={"arquivo": ("a.xlsx", saida.getvalue(), TIPO_XLSX)}
    ).json()
    assert criada["pendencias"] == 0, [
        d["mensagem"] for linha in criada["linhas"] for d in linha["divergencias"]
    ]
    cliente_admin.post(f"/api/importacoes/{criada['id']}/confirmacao")

    criada_no_banco = sessao.scalars(
        select(Instituicao).where(
            Instituicao.nome_normalizado == normalizar("Órgão Com Subcategoria")
        )
    ).first()
    assert criada_no_banco is not None
    assert criada_no_banco.subcategoria_publico_id == sub.id


# =============================================================================
# os achados da revisão do pente fino
# =============================================================================


def _com_cadastro(sessao, semente, instituicao="Órgão Declarado", **colunas):
    """Uma agenda com uma instituição nova, declarada com as colunas que o teste quiser."""
    from datetime import date

    from openpyxl import load_workbook

    modelo = modelo_de_importacao.gerar(
        importar_agendas.vocabularios(sessao),
        importar_agendas.pares_de_vocabulario(sessao),
    )
    pasta = load_workbook(io.BytesIO(modelo))
    folha = pasta["Agendas"]
    cabecalho = [celula.value for celula in next(folha.iter_rows())]
    folha.append(
        [
            {"Data": date(2026, 9, 30), "Instituição": instituicao, "UF": "SP"}.get(coluna)
            for coluna in cabecalho
        ]
    )
    _declarar(
        pasta,
        "instituicoes",
        **{
            "Instituição": instituicao,
            "Categoria de público": "Poder Executivo",
            **colunas,
        },
    )
    saida = io.BytesIO()
    pasta.save(saida)
    return saida.getvalue()


def test_abrangencia_INVALIDA_trava_no_upload_e_nao_estoura_na_confirmacao(
    cliente_admin, sessao, semente
):
    """ACHADO CRÍTICO DA REVISÃO. "Minas Gerais" no lugar de "MG": o upload dizia que
    estava tudo pronto, e a confirmação estourava — `instituicao.uf` tem domínio no
    banco, e o erro vinha como falha de integridade DEPOIS de a pessoa conferir tudo.

    O PIOR TIPO DE ERRO: ela lê "nada pendente", clica em subir, e recebe erro interno
    sobre uma célula que ninguém apontou. A resposta certa é pendência no upload, junto
    das outras, dizendo a coluna e o valor."""
    conteudo = _com_cadastro(sessao, semente, **{"Abrangência": "Minas Gerais"})

    criada = cliente_admin.post(
        "/api/importacoes", files={"arquivo": ("a.xlsx", conteudo, TIPO_XLSX)}
    ).json()

    assert criada["pendencias"] >= 1, criada["a_criar"]
    mensagens = " ".join(
        d["mensagem"] for linha in criada["linhas"] for d in linha["divergencias"]
    )
    assert "Minas Gerais" in mensagens
    assert "Abrangência" in mensagens


def test_editar_uma_celula_NAO_perde_o_cadastro_declarado(cliente_admin, sessao, semente):
    """ACHADO CRÍTICO DA REVISÃO, e é o mais traiçoeiro: a pessoa declara a instituição,
    o upload promete criá-la, ela corrige QUALQUER célula daquela linha na grade — e a
    promessa desaparece, virando pendência travada.

    A CAUSA: a reproposição de uma linha só não tinha como saber o que foi declarado,
    porque o arquivo não é guardado. O que ela tem é a própria divergência gravada, que
    já carrega a declaração — e é de lá que ela precisa ler."""
    conteudo = _com_cadastro(sessao, semente)
    criada = cliente_admin.post(
        "/api/importacoes", files={"arquivo": ("a.xlsx", conteudo, TIPO_XLSX)}
    ).json()
    assert criada["pendencias"] == 0, criada["grupos"]
    assert criada["a_criar"], "o upload tinha de prometer criar a instituição"
    (linha,) = [linha for linha in criada["linhas"] if linha["aba"] == "Agendas"]

    depois = cliente_admin.patch(
        f"/api/importacoes/{criada['id']}/linhas/{linha['id']}",
        json={"celulas": {"UF": "RJ"}},
    ).json()

    assert depois["pendencias"] == 0, [
        d["mensagem"] for linha in depois["linhas"] for d in linha["divergencias"]
    ]
    assert depois["a_criar"], "a promessa de criar a instituição tinha de continuar"
    confirmacao = cliente_admin.post(f"/api/importacoes/{criada['id']}/confirmacao")
    assert confirmacao.status_code == 201, confirmacao.text


def test_a_importacao_ANTIGA_com_o_campo_velho_ainda_confirma(cliente_admin, sessao, semente):
    """ACHADO ALTO DA REVISÃO. Importações criadas antes de `categoria_declarada` virar
    `declarado` estão no banco com o campo velho, e a conferência delas continua aberta.
    Sem ler o campo antigo, a confirmação recusa a instituição como "sem categoria" —
    uma importação que a pessoa já conferiu deixa de poder ser confirmada."""
    from uuid import UUID

    from app.banco.tabelas_importacao import ImportacaoLinha

    conteudo = _com_cadastro(sessao, semente, instituicao="Órgão Do Passado")
    criada = cliente_admin.post(
        "/api/importacoes", files={"arquivo": ("a.xlsx", conteudo, TIPO_XLSX)}
    ).json()

    # O QUE O BANCO ANTIGO TEM: o campo velho, e nada do novo.
    linha = sessao.scalars(
        select(ImportacaoLinha).where(ImportacaoLinha.importacao_id == UUID(criada["id"]))
    ).first()
    linha.divergencias = [
        {
            **bruta,
            "categoria_declarada": (bruta.get("declarado") or {}).get(
                "categoria_publico_id"
            ),
            "declarado": None,
        }
        for bruta in linha.divergencias
    ]
    sessao.flush()

    confirmacao = cliente_admin.post(f"/api/importacoes/{criada['id']}/confirmacao")

    assert confirmacao.status_code == 201, confirmacao.text
    assert confirmacao.json()["cadastros"] == 1


def test_subcategoria_de_OUTRA_categoria_trava_em_vez_de_ser_ignorada(
    cliente_admin, sessao, semente
):
    """ACHADO ALTO DA REVISÃO. A subcategoria de outra categoria era DESCARTADA em
    silêncio: a pessoa preenchia, a confirmação criava o cadastro sem ela, e ninguém
    ficava sabendo. Dado fechado incompatível é pendência — o valor que ela escolheu não
    pertence à categoria que ela escolheu, e só ela sabe qual dos dois está errado."""
    from app.banco.tabelas_catalogo import CategoriaPublico, SubcategoriaPublico

    executivo = sessao.scalars(
        select(CategoriaPublico).where(CategoriaPublico.nome == "Poder Executivo")
    ).first()
    # O NOME NÃO PODE EXISTIR na categoria declarada: "Federal" se repete de propósito
    # em Poder Executivo, Legislativo e Reguladores, e escolher um repetido faria o teste
    # passar por ele ser válido ali — provando o contrário do que ele quer.
    do_executivo = {
        normalizar(nome)
        for (nome,) in sessao.execute(
            select(SubcategoriaPublico.nome).where(
                SubcategoriaPublico.categoria_publico_id == executivo.id
            )
        ).all()
    }
    de_outra = next(
        (
            sub
            for sub in sessao.scalars(
                select(SubcategoriaPublico).where(
                    SubcategoriaPublico.categoria_publico_id != executivo.id,
                    SubcategoriaPublico.ativo.is_(True),
                )
            ).all()
            if normalizar(sub.nome) not in do_executivo
        ),
        None,
    )
    if de_outra is None:
        import pytest

        pytest.skip("a base não tem subcategoria de outra categoria")

    conteudo = _com_cadastro(sessao, semente, **{"Subcategoria": de_outra.nome})
    criada = cliente_admin.post(
        "/api/importacoes", files={"arquivo": ("a.xlsx", conteudo, TIPO_XLSX)}
    ).json()

    assert criada["pendencias"] >= 1, criada["a_criar"]
    mensagens = " ".join(
        d["mensagem"] for linha in criada["linhas"] for d in linha["divergencias"]
    )
    assert de_outra.nome in mensagens


def test_a_RELEVANCIA_invalida_trava_no_upload(cliente_admin, sessao, semente):
    """ACHADO MÉDIO DA REVISÃO: `tier = "999"` era gravado e só falhava por chave
    estrangeira no fim. Vocabulário fechado na aba de cadastro tem de ser conferido onde
    todos os outros são — no upload."""
    conteudo = _com_cadastro(sessao, semente, **{"Relevância": "999"})

    criada = cliente_admin.post(
        "/api/importacoes", files={"arquivo": ("a.xlsx", conteudo, TIPO_XLSX)}
    ).json()

    assert criada["pendencias"] >= 1, criada["a_criar"]
    mensagens = " ".join(
        d["mensagem"] for linha in criada["linhas"] for d in linha["divergencias"]
    )
    assert "999" in mensagens


def test_o_PORTA_VOZ_com_resposta_que_nao_e_sim_nem_nao_trava(cliente_admin, sessao, semente):
    """ACHADO MÉDIO DA REVISÃO: "talvez" virava `False` em silêncio. É a mesma regra que
    a coluna de sim/não da agenda já tem — o "talvez" lá vira pendência, e aqui não
    virava."""
    from datetime import date

    from openpyxl import load_workbook

    modelo = modelo_de_importacao.gerar(
        importar_agendas.vocabularios(sessao),
        importar_agendas.pares_de_vocabulario(sessao),
    )
    pasta = load_workbook(io.BytesIO(modelo))
    folha = pasta["Agendas"]
    cabecalho = [celula.value for celula in next(folha.iter_rows())]
    folha.append(
        [
            {
                "Data": date(2026, 9, 30),
                "Instituição": semente["instituicao"].nome,
                "UF": "SP",
                "Pessoa da Aegea 1": "Alguém Novo",
            }.get(coluna)
            for coluna in cabecalho
        ]
    )
    _declarar(
        pasta,
        "pessoas_aegea",
        **{"Pessoa da Aegea": "Alguém Novo", "É porta-voz?": "talvez"},
    )
    saida = io.BytesIO()
    pasta.save(saida)

    criada = cliente_admin.post(
        "/api/importacoes", files={"arquivo": ("a.xlsx", saida.getvalue(), TIPO_XLSX)}
    ).json()

    assert criada["pendencias"] >= 1, criada["a_criar"]
    mensagens = " ".join(
        d["mensagem"] for linha in criada["linhas"] for d in linha["divergencias"]
    )
    assert "talvez" in mensagens
