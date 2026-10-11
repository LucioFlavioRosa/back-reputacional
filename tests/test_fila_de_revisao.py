"""O cadastro que nasce de planilha espera revisão, e a fila nasce vazia.

PEDIDO DO DONO DO PRODUTO (11/10/2026):

  "Quando subirmos uma base que tenha algum ponto de conferência de cadastro, pode
  subir sem problemas. Porém, dentro de cada uma das 3 abas de Cadastro
  compartilhado, deve existir uma aba específica para cadastros que precisam ser
  revisados para serem consolidados na base de instituições/contatos, temas e
  representantes da Aegea. Assim o profissional consegue subir a planilha sem
  problemas e pode revisar depois — e inclusive pode ser trabalho de outra pessoa.
  Queremos dar padronização, mas com menor impacto."

QUATRO NASCIMENTOS PARA TRÊS ABAS, e o mapa é exato: "Instituições e Contatos" é
`instituicao` + `interlocutor`, "Temas" é `tema`, "Representantes Aegea" é
`pessoa_aegea`. São também os quatro cadastros que a importação de agendas sabe
criar (`VOCABULARIOS_QUE_A_IMPORTACAO_CRIA`), mais a instituição que a ingestão de
menções cria.

POR QUE ESTE ARQUIVO EXISTE. A marca é uma linha por nascimento, e uma linha
esquecida não quebra nada: o cadastro entra, a tela funciona, e ele simplesmente
nunca aparece na fila. Uma fila que não tem tudo é pior que fila nenhuma, porque
ensina a confiar nela.

E A SEGUNDA METADE, de 11/10/2026: a ROTA que lê a fila e o ato que tira dela.
Durante um dia a coluna existiu sem ninguém a ler — o pior estado possível,
porque a importação PAROU DE TRAVAR contando com uma revisão que não acontecia.
O dono do produto cobrou: "e pedi também por uma aba extra somente para
conferência de novos registros e não vi".
"""

from __future__ import annotations

import io
import uuid
from datetime import date

import pytest
from openpyxl import load_workbook
from sqlalchemy import create_engine, func, select
from sqlalchemy.orm import Session

from app.banco.tabelas_catalogo import (
    CategoriaPublico,
    SubcategoriaPublico,
    Tema,
)
from app.banco.tabelas_stakeholders import Instituicao, Interlocutor, PessoaAegea
from app.casos_de_uso import importar_agendas, modelo_de_importacao
from app.casos_de_uso.veiculos_da_imprensa import criar, reconhecer
from app.dominio.importacao_de_agendas import (
    COLUNA_DA_CATEGORIA_DE_INSTITUICAO,
    ROTULO_DO_VOCABULARIO,
    colunas_do_cadastro,
)
from app.dominio.texto import normalizar

#: O AJUDANTE DE ONDE ELE JÁ EXISTE, e não uma cópia: `_cliente` monta o
#: `TestClient` sobrescrevendo `obter_configuracao` por `dependency_overrides` — o
#: `Depends` já foi resolvido na importação do módulo, então `monkeypatch.setattr`
#: não alcança quem guardou a função. Uma segunda cópia disso divergiria no dia em
#: que o login mudasse.
#:
#: A FUNÇÃO E NÃO A FIXTURE: importar a fixture e usá-la como parâmetro redefine o
#: nome (F811), e este projeto não compartilha fixture entre arquivos de teste —
#: cada um monta as suas.
from main import app
from tests.test_e2e_postgres import URL
from tests.test_importacoes_api import TIPO_XLSX, _cliente

_engine = create_engine(URL, pool_pre_ping=True)


@pytest.fixture
def sessao():
    conexao = _engine.connect()
    transacao = conexao.begin()
    sessao = Session(bind=conexao, expire_on_commit=False)
    try:
        yield sessao
    finally:
        sessao.close()
        transacao.rollback()
        conexao.close()


@pytest.fixture
def cliente_admin(sessao):
    """`plataforma_edicao` é o único perfil com `administra_dicionarios`, e é o
    que a importação exige — ver `exigir_administracao_de_cadastros`.

    O `finally` LIMPA OS OVERRIDES: `_cliente` os instala em `app`, que é global, e
    um teste que os deixasse montados contaminaria o próximo arquivo da suíte.
    """
    try:
        yield _cliente(sessao, "plataforma_edicao")
    finally:
        app.dependency_overrides.clear()


#: As quatro tabelas da fila, com o nome da aba de Cadastro compartilhado a que
#: cada uma pertence — é o que faz a falha dizer ONDE o registro deixaria de
#: aparecer.
DA_FILA = (
    (Instituicao, "Instituições e Contatos"),
    (Interlocutor, "Instituições e Contatos"),
    (Tema, "Temas"),
    (PessoaAegea, "Representantes Aegea"),
)


def test_as_quatro_tabelas_da_fila_TEM_a_coluna(sessao):
    """A fila das três abas sai destas quatro, e de nenhuma outra.

    SE UMA DELAS NÃO TIVER A COLUNA, a aba correspondente não tem o que listar — e
    a descoberta seria pela tela vazia, não por aqui.
    """
    for tabela, aba in DA_FILA:
        quantos = sessao.scalar(
            select(func.count()).select_from(tabela).where(tabela.pendente_de_revisao)
        )
        assert quantos is not None, f"{tabela.__tablename__} não tem a coluna (aba {aba})"


def test_a_fila_NASCE_VAZIA(sessao):
    """O "menor impacto" do pedido, afirmado.

    `default false` É A DECISÃO: os 3.785 cadastros de instituição, 129 de
    interlocutor, 149 de tema e 12 de pessoa da Aegea que já existem foram feitos
    por gente, ou já passaram pela conferência da ingestão. Se a coluna nascesse
    `true`, a fila abriria pedindo revisão de tudo — e uma fila que nasce cheia de
    registro antigo é ruído, que é o que faz ninguém olhar (o argumento da `0078`).
    """
    for tabela, aba in DA_FILA:
        pendentes = sessao.scalars(
            select(tabela.nome).where(tabela.pendente_de_revisao).limit(5)
        ).all()
        assert pendentes == [], f"a aba {aba} abriria com {pendentes} na fila"


def _agenda_com(sessao, **valores) -> bytes:
    """Uma planilha de agendas com uma linha, pelos nomes de coluna."""
    modelo = modelo_de_importacao.gerar(importar_agendas.vocabularios(sessao))
    pasta = load_workbook(io.BytesIO(modelo))
    folha = pasta["Agendas"]
    cabecalho = [celula.value for celula in next(folha.iter_rows())]
    folha.append([valores.get(coluna) for coluna in cabecalho])
    saida = io.BytesIO()
    pasta.save(saida)
    return saida.getvalue()


def test_os_QUATRO_cadastros_da_importacao_de_agendas_nascem_pendentes(
    cliente_admin, sessao
):
    """Instituição, interlocutor, pessoa da Aegea e tema, de uma planilha só.

    OS QUATRO NUMA LINHA, de propósito: é o caso que o dono do produto descreveu —
    "quando temos várias reuniões em um dia o registro completo pode ser uma
    dificuldade". A linha entra com nomes que não existem em cadastro nenhum, e os
    quatro nascem.

    PASSA PELA CONFIRMAÇÃO, e isto não é detalhe: os cadastros nascem em
    `_criar_cadastros`, que só roda ali. Uma primeira versão deste teste afirmava
    `proposta.a_criar` — o que prova que a importação SE PROPÕE a criar, e nada
    sobre a marca da fila. Passaria com a marca removida dos quatro.
    """
    conteudo = _agenda_com(
        sessao,
        **{
            "Data": date(2026, 9, 25),
            "Instituição": "Órgão Novo Da Fila ZZ",
            "UF": "SP",
            "Interlocutor 1": "Pessoa Nova Da Fila ZZ",
            "Pessoa da Aegea 1": "Representante Novo Da Fila ZZ",
            "Tema 1": "Assunto Novo Da Fila ZZ",
        },
    )

    criada = cliente_admin.post(
        "/api/importacoes", files={"arquivo": ("dia.xlsx", conteudo, TIPO_XLSX)}
    ).json()
    #: NÃO BARRA: é a premissa do pedido, e sem ela o resto não acontece.
    assert criada["pendencias"] == 0, criada["grupos"]

    confirmacao = cliente_admin.post(f"/api/importacoes/{criada['id']}/confirmacao")
    assert confirmacao.status_code == 201, confirmacao.text

    #: OS QUATRO NASCERAM, E OS QUATRO ESTÃO NA FILA.
    nascidos = (
        (Instituicao, "Órgão Novo Da Fila ZZ", "Instituições e Contatos"),
        (Interlocutor, "Pessoa Nova Da Fila ZZ", "Instituições e Contatos"),
        (PessoaAegea, "Representante Novo Da Fila ZZ", "Representantes Aegea"),
        (Tema, "Assunto Novo Da Fila ZZ", "Temas"),
    )
    for tabela, nome, aba in nascidos:
        registro = sessao.scalar(select(tabela).where(tabela.nome == nome))
        assert registro is not None, f"{nome} não foi criado"
        assert registro.pendente_de_revisao is True, (
            f"{nome} nasceu fora da fila — não apareceria na aba {aba}"
        )


def test_o_veiculo_da_ingestao_NASCE_pendente(sessao):
    """A ingestão de menções também, e o motivo de ela entrar na fila importa.

    ELA JÁ TEM CONFERÊNCIA, e ainda assim entra — porque as duas perguntam coisas
    diferentes. Na conferência da ingestão a pessoa autoriza a CRIAÇÃO olhando
    nome, praça, alcance e cargo: ela diz "este veículo existe e deve entrar". O
    ENQUADRAMENTO é outra pergunta — a subcategoria é lógica editorial que nenhuma
    coluna da Clipei informa, e é dela que a lente Mercado depende. Quem autoriza
    2.631 criações numa tela não decide 2.631 enquadramentos na mesma tela.
    """
    nome = "Veículo Da Fila ZZ"
    reco = reconhecer(
        sessao, {nome: {"uf": "SP", "abrangencia": "Local", "mencoes": 3, "tier": ""}}
    )
    criar(sessao, reco.novos)
    sessao.flush()

    novo = sessao.scalar(
        select(Instituicao).where(Instituicao.nome_normalizado == normalizar(nome))
    )
    assert novo is not None, "o veículo não foi criado"
    assert novo.pendente_de_revisao is True
    #: E O MOTIVO DIZ QUAL ENQUADRAMENTO FALTA — ver os três casos abaixo.
    assert novo.motivo_da_revisao is not None


def test_quem_a_PESSOA_cadastra_pela_tela_NAO_entra_na_fila(sessao):
    """O outro lado da regra, e é ele que mantém a fila útil.

    O `default` da coluna é `false`, então um cadastro criado por qualquer caminho
    que não seja planilha nasce fora da fila. Sem esta asserção, alguém poderia
    trocar o default por `true` para "garantir que nada escape" — e a fila passaria
    a conter tudo que a plataforma cria, incluindo o que acabou de ser conferido
    por gente na própria tela de cadastro.
    """
    nome = "Entidade Feita A Mao ZZ"
    sessao.add(
        Instituicao(nome=nome, nome_normalizado=normalizar(nome), tipo="entidade")
    )
    sessao.flush()

    feita = sessao.scalar(
        select(Instituicao).where(Instituicao.nome_normalizado == normalizar(nome))
    )
    assert feita.pendente_de_revisao is False


# -- A ABA: a rota que lê a fila, e o ato que tira dela ------------------------


@pytest.fixture
def cliente_sem_admin(sessao):
    """`crm_edicao` abre o portal do CRM e NÃO administra cadastros.

    É o contrapeso EXATO da guarda, e não um perfil sem acesso a nada — que
    passaria com qualquer dependência, inclusive com a errada.
    """
    try:
        yield _cliente(sessao, "crm_edicao")
    finally:
        app.dependency_overrides.clear()


def test_a_rota_LISTA_as_tres_abas_sempre(cliente_admin):
    """As três, inclusive as vazias, e na ordem da tela.

    UMA ABA QUE DESAPARECE QUANDO ESTÁ LIMPA é uma aba que ninguém aprende a
    olhar: quem abre o Cadastro compartilhado precisa ver que a revisão existe e
    que não há nada pendente — não um espaço em branco indistinguível de uma
    tela que não carregou.
    """
    resposta = cliente_admin.get("/api/revisao-de-cadastros")
    assert resposta.status_code == 200, resposta.text

    abas = resposta.json()
    assert [aba["id"] for aba in abas] == ["instituicoes", "assuntos", "porta_vozes"]
    assert [aba["rotulo"] for aba in abas] == [
        "Instituições e Contatos",
        "Temas",
        "Representantes Aegea",
    ]


def test_quem_NAO_administra_cadastros_nao_le_a_fila(cliente_sem_admin):
    """Mais estrito que as listagens de cadastro, que pedem `ve_diretorio`.

    A FILA NÃO É A BASE: um registro aqui diz que a plataforma não confia naquele
    enquadramento ainda. Quem não pode consertar não ganha nada em ver a lista do
    que está por consertar — e a lista carrega o motivo, que é uma dúvida
    interna sobre o dado.
    """
    assert cliente_sem_admin.get("/api/revisao-de-cadastros").status_code == 403


def test_os_quatro_nascimentos_APARECEM_cada_um_na_sua_aba(cliente_admin, sessao):
    """O mapa de `DA_FILA`, afirmado pela ROTA e não pela coluna.

    `test_os_QUATRO_cadastros_da_importacao_de_agendas_nascem_pendentes` prova que
    a marca foi gravada. Este prova que ela CHEGA NA TELA, e na aba certa — que é
    o que o dono do produto pediu e o que faltava.
    """
    conteudo = _agenda_com(
        sessao,
        **{
            "Data": date(2026, 9, 25),
            "Instituição": "Órgão Da Aba Zz73",
            "UF": "SP",
            "Interlocutor 1": "Pessoa Da Aba Zz73",
            "Pessoa da Aegea 1": "Representante Da Aba Zz73",
            "Tema 1": "Assunto Da Aba Zz73",
        },
    )
    criada = cliente_admin.post(
        "/api/importacoes", files={"arquivo": ("dia.xlsx", conteudo, TIPO_XLSX)}
    ).json()
    assert criada["pendencias"] == 0, criada["grupos"]
    assert (
        cliente_admin.post(f"/api/importacoes/{criada['id']}/confirmacao").status_code
        == 201
    )

    abas = {
        aba["id"]: {p["nome"]: p for p in aba["pendentes"]}
        for aba in cliente_admin.get("/api/revisao-de-cadastros").json()
    }

    esperado = (
        ("instituicoes", "Órgão Da Aba Zz73", "instituicao"),
        ("instituicoes", "Pessoa Da Aba Zz73", "interlocutor"),
        ("assuntos", "Assunto Da Aba Zz73", "tema"),
        ("porta_vozes", "Representante Da Aba Zz73", "pessoa_aegea"),
    )
    for aba, nome, tipo in esperado:
        assert nome in abas[aba], f"{nome} não apareceu na aba {aba}: {list(abas[aba])}"
        assert abas[aba][nome]["tipo"] == tipo


def test_a_fila_DIZ_a_instituicao_de_cada_pessoa_homonima(cliente_admin, sessao):
    """O caso real, e é a razão de `contexto` existir.

    A PLANILHA DE RI tem quatro nomes em duas instituições cada, e desde
    11/10/2026 cada um nasce DUAS VEZES, de propósito — a `0002` diz que "mesmo
    nome em instituições diferentes são pessoas diferentes".

    SEM A INSTITUIÇÃO AO LADO a aba mostraria "Artur Santa" duas vezes, com o
    mesmo motivo, e quem revisa não teria como saber qual é o do Banco XP e qual
    é o da XP Asset. A fila ficaria ilegível justamente no caso que a fez
    existir — e pior que ilegível: convidaria a fundir os dois.
    """
    modelo = modelo_de_importacao.gerar(importar_agendas.vocabularios(sessao))
    pasta = load_workbook(io.BytesIO(modelo))
    folha = pasta["Agendas"]
    cabecalho = [celula.value for celula in next(folha.iter_rows())]
    for casa in ("Banco Zz74 Um", "Banco Zz74 Dois"):
        folha.append(
            [
                {
                    "Data": date(2026, 9, 25),
                    "Instituição": casa,
                    "UF": "SP",
                    "Interlocutor 1": "Pessoa Em Duas Zz74",
                }.get(coluna)
                for coluna in cabecalho
            ]
        )
    saida = io.BytesIO()
    pasta.save(saida)

    criada = cliente_admin.post(
        "/api/importacoes", files={"arquivo": ("ri.xlsx", saida.getvalue(), TIPO_XLSX)}
    ).json()
    assert criada["pendencias"] == 0, criada["grupos"]
    assert (
        cliente_admin.post(f"/api/importacoes/{criada['id']}/confirmacao").status_code
        == 201
    )

    aba = next(
        a
        for a in cliente_admin.get("/api/revisao-de-cadastros").json()
        if a["id"] == "instituicoes"
    )
    as_duas = [
        p for p in aba["pendentes"] if p["nome"] == "Pessoa Em Duas Zz74"
    ]
    assert len(as_duas) == 2, as_duas

    #: AS DUAS SÃO DISTINGUÍVEIS, e é só o `contexto` que as distingue.
    assert {p["contexto"] for p in as_duas} == {"Banco Zz74 Um", "Banco Zz74 Dois"}
    #: E AS DUAS DIZEM POR QUÊ.
    for pendente in as_duas:
        assert "2 instituições" in pendente["motivo"]


def test_consolidar_TIRA_da_fila_e_APAGA_o_motivo(cliente_admin, sessao):
    """O ato de revisar, e os dois efeitos que ele tem.

    O MOTIVO SAI COM A PENDÊNCIA porque ele é a PERGUNTA que a fila fazia.
    Mantê-lo deixaria "apareceu em 2 instituições — confirme se são a mesma
    pessoa" grudado num cadastro sobre o qual alguém JÁ confirmou, pronto para
    reaparecer na próxima tela que o leia.
    """
    nome = "Entidade Consolidada Zz75"
    registro = Instituicao(
        nome=nome,
        nome_normalizado=normalizar(nome),
        tipo="entidade",
        uf="SP",
        pendente_de_revisao=True,
        motivo_da_revisao="Nasceu da planilha de teste.",
    )
    sessao.add(registro)
    sessao.flush()

    na_fila = cliente_admin.get("/api/revisao-de-cadastros").json()[0]["pendentes"]
    assert nome in [p["nome"] for p in na_fila]

    resposta = cliente_admin.post(
        f"/api/revisao-de-cadastros/instituicao/{registro.id}/consolidacao"
    )
    assert resposta.status_code == 200, resposta.text
    assert resposta.json()["ja_estava_consolidado"] is False
    assert resposta.json()["pendente"]["motivo"] is None
    #: O CONTEXTO CONTINUA, porque ele descreve o cadastro e não a pendência.
    assert resposta.json()["pendente"]["contexto"] == "Entidade · SP"

    sessao.refresh(registro)
    assert registro.pendente_de_revisao is False
    assert registro.motivo_da_revisao is None

    #: E SAIU DA LISTA.
    depois = cliente_admin.get("/api/revisao-de-cadastros").json()[0]["pendentes"]
    assert nome not in [p["nome"] for p in depois]


def test_consolidar_DUAS_VEZES_diz_que_ja_estava(cliente_admin, sessao):
    """Duas pessoas na mesma aba é o caso NORMAL, não a exceção.

    O PEDIDO É EXPLÍCITO: "inclusive pode ser trabalho de outra pessoa". Recusar
    o segundo clique seria inventar um conflito onde não há desacordo — o estado
    final é o que as duas pessoas queriam. O booleano existe para a tela dizer
    "outra pessoa já consolidou este" em vez de piscar uma confirmação falsa.
    """
    nome = "Entidade Duas Vezes Zz76"
    registro = Instituicao(
        nome=nome,
        nome_normalizado=normalizar(nome),
        tipo="entidade",
        pendente_de_revisao=True,
    )
    sessao.add(registro)
    sessao.flush()

    rota = f"/api/revisao-de-cadastros/instituicao/{registro.id}/consolidacao"
    assert cliente_admin.post(rota).json()["ja_estava_consolidado"] is False
    segunda = cliente_admin.post(rota)
    assert segunda.status_code == 200, segunda.text
    assert segunda.json()["ja_estava_consolidado"] is True


def test_o_id_que_nao_e_do_tipo_da_404_e_nao_500(cliente_admin):
    """`sessao.get(Tema, "abc")` estoura `DataError` e aborta a transação.

    UM 500 ONDE O CERTO É 404, e o custo é maior que a resposta errada: o
    `DataError` aborta a transação do pedido INTEIRO no Postgres, então qualquer
    coisa que viesse depois na mesma sessão também falharia.

    E ESTE NÃO É UM CASO HIPOTÉTICO: o id vem da URL, e o front monta a URL com
    o id que a listagem devolveu como TEXTO — um `tema` com id inteiro e uma
    `instituicao` com UUID saem da mesma lista, pela mesma rota.
    """
    assert cliente_admin.post(
        "/api/revisao-de-cadastros/tema/nao-sou-numero/consolidacao"
    ).status_code == 404
    assert cliente_admin.post(
        "/api/revisao-de-cadastros/instituicao/nao-sou-uuid/consolidacao"
    ).status_code == 404
    #: E O TIPO QUE NÃO É DA FILA.
    assert cliente_admin.post(
        "/api/revisao-de-cadastros/mencao/1/consolidacao"
    ).status_code == 404


def test_toda_tabela_da_fila_SAI_pela_rota():
    """Âncora: as quatro que entram na fila são as quatro que a rota consolida.

    SEM ELA, uma tabela acrescentada a `_TABELAS` e esquecida em `POR_TIPO`
    apareceria na aba e recusaria a consolidação — a pessoa veria um registro
    que não sai, e descobriria pelo 404. `POR_TIPO` é derivado de `_TABELAS` por
    essa razão; este teste prende a derivação ao mapa de `DA_FILA`, que é o
    pedido do dono do produto escrito em código.
    """
    from app.casos_de_uso.fila_de_revisao import POR_TIPO

    assert set(POR_TIPO) == {tabela.__tablename__ for tabela, _ in DA_FILA}


# -- o motivo de quem nasce sem categoria de público ---------------------------


def test_a_instituicao_SEM_categoria_diz_na_fila_por_que_ficou_entidade(
    cliente_admin, sessao
):
    """Decisão do dono do produto depois de ver a fila real (11/10/2026).

      "pode registrar como entidade e marcar para revisar"

    O COMPORTAMENTO JÁ ERA ESSE; o que faltava era a marca DIZER o que há para
    fazer. Medido com a planilha de RI: 499 das 502 pendências são instituições e
    TODAS nasceram assim — `3R Investimentos`, `AF Invest`, `ARX`, numa planilha
    de Relações com Investidores, todas como "Entidade".

    SEM ESTE TEXTO a aba mostra 499 linhas repetindo "entrou por planilha e
    ninguém conferiu", que não diz O QUE conferir. Uma fila de 499 linhas
    idênticas não é fila: é parede.
    """
    conteudo = _agenda_com(
        sessao,
        **{
            "Data": date(2026, 9, 25),
            "Instituição": "Gestora Sem Categoria Zz81",
            "UF": "SP",
        },
    )
    criada = cliente_admin.post(
        "/api/importacoes", files={"arquivo": ("ri.xlsx", conteudo, TIPO_XLSX)}
    ).json()
    assert criada["pendencias"] == 0, criada["grupos"]
    assert (
        cliente_admin.post(f"/api/importacoes/{criada['id']}/confirmacao").status_code
        == 201
    )

    nascida = sessao.scalar(
        select(Instituicao).where(
            Instituicao.nome_normalizado == normalizar("Gestora Sem Categoria Zz81")
        )
    )
    #: `entidade` CONTINUA SENDO O TIPO, e é o certo: o único fora de
    #: `FRENTE_UNICA_DO_TIPO`, então ele não crava uma frente errada.
    assert nascida.tipo == "entidade"
    assert nascida.categoria_publico_id is None
    assert nascida.pendente_de_revisao is True
    #: E A MARCA DIZ O QUE FAZER: nomeia o campo que falta e o efeito de faltar.
    assert nascida.motivo_da_revisao is not None
    assert "categoria de público" in nascida.motivo_da_revisao
    assert "Entidade" in nascida.motivo_da_revisao


def test_a_instituicao_COM_categoria_nao_ganha_esse_motivo(cliente_admin, sessao):
    """O outro lado, e é ele que mantém o motivo informativo.

    Se toda instituição nova carregasse a frase, ela deixaria de distinguir
    coisa alguma — e quem lê a aba aprenderia a pular a linha do motivo, que é o
    único lugar onde a fila diz o que fazer.

    ELA SEGUE PENDENTE, porque nasceu de planilha: o `motivo` nulo é "entrou por
    planilha, sem mais nada a dizer", e a tela escreve a frase genérica.
    """
    categoria = sessao.scalar(
        select(CategoriaPublico).where(CategoriaPublico.nome == "Poder Executivo")
    )
    #: A CATEGORIA NÃO É COLUNA DA AGENDA, e esta foi a minha confusão ao escrever
    #: o teste: ela vive na ABA DE CADASTRO da instituição — a aba onde se DECLARA
    #: o que a planilha vai criar. Escrita na linha da agenda, ela é uma coluna que
    #: não existe, e a instituição nasce sem categoria: o teste falhava
    #: `assert None == 1` afirmando o contrário do que queria.
    conteudo = _agenda_declarando(
        sessao,
        agenda={
            "Data": date(2026, 9, 25),
            "Instituição": "Orgao Com Categoria Zz82",
            "UF": "SP",
        },
        instituicao={"Nome": "Orgao Com Categoria Zz82", "categoria": categoria.nome},
    )
    criada = cliente_admin.post(
        "/api/importacoes", files={"arquivo": ("ri.xlsx", conteudo, TIPO_XLSX)}
    ).json()
    assert criada["pendencias"] == 0, criada["grupos"]
    assert (
        cliente_admin.post(f"/api/importacoes/{criada['id']}/confirmacao").status_code
        == 201
    )

    nascida = sessao.scalar(
        select(Instituicao).where(
            Instituicao.nome_normalizado == normalizar("Orgao Com Categoria Zz82")
        )
    )
    assert nascida.categoria_publico_id == categoria.id
    assert nascida.pendente_de_revisao is True
    assert nascida.motivo_da_revisao is None


def _agenda_declarando(sessao, *, agenda: dict, instituicao: dict) -> bytes:
    """Uma agenda com a instituição DECLARADA na aba de cadastro dela.

    `_agenda_com` só escreve a aba Agendas, e é o bastante para quase tudo: a
    importação cria o que encontra nas células. Mas a categoria de público, o
    tier e os outros campos do formulário vêm da aba de CADASTRO — é lá que a
    planilha diz "crie esta instituição ASSIM", em vez de só mencioná-la.

    POR NOME DE COLUNA e não por posição, como o helper de
    `test_importar_agendas.py`: as abas de cadastro ganharam os campos do
    formulário, então a categoria deixou de ser a segunda coluna — e escrever por
    posição gravaria a categoria no "Nome completo" sem reclamar de nada.
    """
    modelo = modelo_de_importacao.gerar(importar_agendas.vocabularios(sessao))
    pasta = load_workbook(io.BytesIO(modelo))

    folha = pasta["Agendas"]
    cabecalho = [celula.value for celula in next(folha.iter_rows())]
    desconhecidas = set(agenda) - set(cabecalho)
    assert not desconhecidas, f"coluna que a aba Agendas não tem: {desconhecidas}"
    folha.append([agenda.get(coluna) for coluna in cabecalho])

    colunas = [coluna.nome for coluna in colunas_do_cadastro("instituicoes")]
    valores = {
        colunas[0]: instituicao["Nome"],
        COLUNA_DA_CATEGORIA_DE_INSTITUICAO: instituicao["categoria"],
    }
    pasta[ROTULO_DO_VOCABULARIO["instituicoes"]].append(
        [valores.get(coluna) for coluna in colunas]
    )

    saida = io.BytesIO()
    pasta.save(saida)
    return saida.getvalue()


# -- os TRÊS motivos do que a ingestão de menções cria -------------------------
#
# ACHADO DE REVISÃO COM SEVERIDADE MÉDIA: a ingestão nascia pendente SEM motivo, e
# ela é o produtor de MAIOR VOLUME da fila — 2.631 veículos de uma subida só. O
# comentário no código já explicava o porquê em prosa; a fila mostrava a frase
# genérica.
#
# E NÃO É O MESMO MOTIVO DA IMPORTAÇÃO DE AGENDAS. Lá falta a CATEGORIA e o tipo
# cai em `entidade`. Aqui a categoria vem sempre — `CATEGORIA_DE_QUEM_FALA` dá
# "Imprensa" ao veículo, e o cargo dá a do perfil. Usar a mesma frase mandaria a
# pessoa procurar um campo que já está preenchido.


def test_o_VEICULO_da_ingestao_diz_que_falta_a_subcategoria(sessao):
    """A categoria é Imprensa; o que falta é a lógica editorial.

    E É DELA QUE A LENTE MERCADO DEPENDE: a lista do Mercado é `tipo='veiculo'`
    com categoria Imprensa e subcategoria "Econômica e de negócios" — 81 de
    2.712. Sem a subcategoria o veículo fica fora da lente, e ninguém saberia
    por quê.
    """
    nome = "Veículo Com Motivo Zz83"
    reco = reconhecer(
        sessao, {nome: {"uf": "SP", "abrangencia": "Local", "mencoes": 3, "tier": ""}}
    )
    criar(sessao, reco.novos)
    sessao.flush()

    novo = sessao.scalar(
        select(Instituicao).where(Instituicao.nome_normalizado == normalizar(nome))
    )
    assert novo.tipo == "veiculo"
    #: A CATEGORIA VEIO: é o que torna errado reusar o motivo da importação.
    assert novo.categoria_publico_id is not None
    assert novo.subcategoria_publico_id is None
    assert "subcategoria" in novo.motivo_da_revisao
    assert "Mercado" in novo.motivo_da_revisao


def test_o_PERFIL_sem_cargo_reconhecido_diz_que_falta_o_cargo(sessao):
    """O caso em que a fila tem mais a ganhar.

    SEM CARGO RECONHECIDO não há categoria nenhuma — é o estado dos 1.108 perfis
    que a primeira carga da Bites criou. A frase nomeia o campo que destrava o
    resto: dado o cargo, a categoria sai dele.
    """
    nome = "perfilsemcargo_zz84"
    reco = reconhecer(
        sessao,
        {nome: {"uf": "SP", "abrangencia": "Local", "mencoes": 2, "cargo": ""}},
        tipo="perfil_rede",
    )
    criar(sessao, reco.novos, tipo="perfil_rede")
    sessao.flush()

    novo = sessao.scalar(
        select(Instituicao).where(Instituicao.nome_normalizado == normalizar(nome))
    )
    assert novo.tipo == "perfil_rede"
    assert novo.categoria_publico_id is None
    assert "cargo" in novo.motivo_da_revisao


def test_o_perfil_COM_cargo_nao_e_chamado_de_SEM_cargo_quando_a_categoria_falta(
    sessao,
):
    """O descolamento que um achado de revisão nomeou, e o estrago é permanente.

    A FRASE DA FILA DECIDIA POR `publico is None` — o EFEITO no banco. O fato que
    a gera é outro: "a fonte mandou um cargo que sabemos traduzir". Os dois
    coincidem no caminho normal e descolam quando a categoria que o cargo aponta
    não está disponível — uma categoria desativada na taxonomia, por exemplo.

    E A FRASE É GRAVADA NA LINHA. Um `id` nulo passageiro viraria uma instrução
    errada que fica no banco pedindo à pessoa um campo que ela já preencheu —
    "diga o cargo" para um perfil que TEM cargo. A fila que mente é a fila que
    ninguém usa.

    O TESTE CHAMA A FUNÇÃO DIRETO, de propósito: reproduzir o descolamento pelo
    caminho de fora exigiria desativar uma categoria da taxonomia no meio do
    teste — mexer no vocabulário fechado do modelo para provar uma frase. A
    função é a unidade onde a decisão mora.
    """
    from app.casos_de_uso.veiculos_da_imprensa import _motivo_da_revisao

    #: CARGO RECONHECIDO: a frase é a de conferência, mesmo que a categoria não
    #: tenha resolvido.
    com_cargo = _motivo_da_revisao("perfil_rede", True)
    assert "confira" in com_cargo
    assert "sem cargo" not in com_cargo

    #: SEM CARGO RECONHECIDO: aí sim a frase pede o cargo.
    sem_cargo = _motivo_da_revisao("perfil_rede", False)
    assert "sem cargo reconhecido" in sem_cargo

    #: E O VEÍCULO NÃO DEPENDE DISSO: a pergunta dele é a subcategoria, sempre.
    assert "subcategoria" in _motivo_da_revisao("veiculo", False)
    assert "subcategoria" in _motivo_da_revisao("veiculo", True)


# -- o enquadramento em LOTE ---------------------------------------------------
#
# POR QUE EXISTE, e o número é o argumento: a planilha real de RI criou 499
# instituições, TODAS como "Entidade", e a resposta para quase todas é a mesma.
# Quatrocentas e noventa e nove idas à lista de cadastro para digitar a mesma
# coisa não é revisão: é a pessoa desistindo na décima, com a base meio
# classificada — pior que não classificada, porque ninguém sabe onde parou.


def _pendente(sessao, nome: str, tipo: str = "entidade") -> Instituicao:
    """Uma instituição como a planilha a cria: tipo genérico e na fila."""
    registro = Instituicao(
        nome=nome,
        nome_normalizado=normalizar(nome),
        tipo=tipo,
        uf="SP",
        pendente_de_revisao=True,
        motivo_da_revisao="Entrou pela planilha sem categoria de público.",
    )
    sessao.add(registro)
    sessao.flush()
    return registro


def _categoria_de(sessao, tipo_desejado: str):
    """Uma categoria ATIVA cujo tipo derivado é o pedido.

    LIDA DO MAPA e não escrita à mão: `TIPO_DA_CATEGORIA_DE_PUBLICO` é quem
    decide, e um teste que fixasse "Poder Executivo" quebraria no dia em que a
    taxonomia mudasse — dizendo que o lote está errado quando mudou foi o mapa.
    """
    from app.dominio.frentes import TIPO_DA_CATEGORIA_DE_PUBLICO

    codigos = [
        codigo
        for codigo, tipo in TIPO_DA_CATEGORIA_DE_PUBLICO.items()
        if tipo == tipo_desejado
    ]
    return sessao.scalar(
        select(CategoriaPublico).where(
            CategoriaPublico.codigo.in_(codigos), CategoriaPublico.ativo.is_(True)
        )
    )


def test_o_LOTE_enquadra_todas_e_deriva_o_tipo_de_uma_vez(cliente_admin, sessao):
    """O gesto inteiro: a categoria vai nas três, e o tipo sai dela.

    O TIPO É O QUE TORNA ISTO MAIS QUE PREENCHER UM CAMPO: `derivar_tipo` traduz
    categoria em tipo, e o tipo da instituição é o que `derivar_frente` usa para
    dizer de que frente é cada agenda dela. Enquadrar o lote reclassifica as
    agendas de todas elas.
    """
    categoria = _categoria_de(sessao, "orgao")
    tres = [_pendente(sessao, f"Entidade Do Lote {n} Zz90") for n in (1, 2, 3)]

    resposta = cliente_admin.post(
        "/api/revisao-de-cadastros/instituicoes/enquadramento",
        json={
            "ids": [str(r.id) for r in tres],
            "categoria_publico_id": categoria.id,
            "consolidar": True,
        },
    )
    assert resposta.status_code == 200, resposta.text
    corpo = resposta.json()
    assert corpo["enquadradas"] == 3
    assert corpo["consolidadas"] == 3
    assert corpo["tipo"] == "orgao"
    assert corpo["recusadas"] == []

    for registro in tres:
        sessao.refresh(registro)
        assert registro.categoria_publico_id == categoria.id
        assert registro.tipo == "orgao"
        #: E SAÍRAM DA FILA, motivo e tudo.
        assert registro.pendente_de_revisao is False
        assert registro.motivo_da_revisao is None


def test_o_LOTE_pode_enquadrar_SEM_consolidar(cliente_admin, sessao):
    """Enquadrar e consolidar são decisões diferentes.

    Dar a categoria resolve o MOTIVO que a fila mostrava; consolidar é afirmar
    que o cadastro está certo. Quem enquadra 499 de uma vez pode querer que outra
    pessoa confira antes — e é o mesmo princípio da consolidação uma a uma: ela
    sai por ação, não por um campo que ficou preenchido.
    """
    categoria = _categoria_de(sessao, "orgao")
    uma = _pendente(sessao, "Entidade Sem Consolidar Zz91")

    corpo = cliente_admin.post(
        "/api/revisao-de-cadastros/instituicoes/enquadramento",
        json={
            "ids": [str(uma.id)],
            "categoria_publico_id": categoria.id,
            "consolidar": False,
        },
    ).json()
    assert corpo["enquadradas"] == 1
    assert corpo["consolidadas"] == 0

    sessao.refresh(uma)
    assert uma.categoria_publico_id == categoria.id
    assert uma.pendente_de_revisao is True


def test_o_LOTE_recusa_a_colisao_de_nome_e_ENQUADRA_o_resto(cliente_admin, sessao):
    """A decisão de desenho mais importante: o lote NÃO é tudo-ou-nada.

    Com 499 linhas, uma colisão de nome faria o lote inteiro falhar e a pessoa não
    teria como saber qual das 499. Enquadrar as que dão e DIZER quais não deram,
    com o motivo de cada uma, é o que permite o gesto acontecer.

    E A COLISÃO É REAL, não zelo: a `0002` cria `unique (nome_normalizado, tipo)`
    e o índice NÃO é parcial em `ativo`. Sem a conferência, mudar o tipo estoura
    em `IntegrityError` no flush — 500 para a pessoa, sem dizer qual nome colidiu.
    """
    categoria = _categoria_de(sessao, "orgao")
    #: QUEM JÁ OCUPA O PAR, do tipo de destino.
    ja_existe = "Nome Que Colide Zz92"
    sessao.add(
        Instituicao(
            nome=ja_existe, nome_normalizado=normalizar(ja_existe), tipo="orgao"
        )
    )
    sessao.flush()

    colidente = _pendente(sessao, ja_existe)
    boa = _pendente(sessao, "Nome Que Passa Zz92")

    corpo = cliente_admin.post(
        "/api/revisao-de-cadastros/instituicoes/enquadramento",
        json={
            "ids": [str(colidente.id), str(boa.id)],
            "categoria_publico_id": categoria.id,
        },
    ).json()

    #: A BOA PASSOU.
    assert corpo["enquadradas"] == 1
    sessao.refresh(boa)
    assert boa.tipo == "orgao"

    #: E A OUTRA FOI RECUSADA COM NOME E MOTIVO ACIONÁVEL.
    assert len(corpo["recusadas"]) == 1
    recusada = corpo["recusadas"][0]
    assert recusada["nome"] == ja_existe
    assert "duplicata" in recusada["motivo"]
    sessao.refresh(colidente)
    assert colidente.tipo == "entidade"
    assert colidente.pendente_de_revisao is True


def test_o_LOTE_recusa_duas_pendencias_de_MESMO_NOME(cliente_admin, sessao):
    """O caso que o dono do produto nomeou, e aqui ele deixa de ser hipótese.

      "se o 'valor' for classificado como financeiro e em outro momento 'o valor'
      for registrado como veículo, devemos apontar essa possível duplicação"

    Duas pendências com o mesmo nome normalizado em tipos diferentes hoje passam a
    querer o MESMO par `(nome_normalizado, tipo)`. É o LOTE que as junta — e sem
    esta conferência a segunda estouraria no flush, derrubando as 499 com ela.
    """
    categoria = _categoria_de(sessao, "orgao")
    nome = "O Valor Zz93"
    uma = _pendente(sessao, nome, tipo="entidade")
    outra = _pendente(sessao, nome, tipo="veiculo")

    corpo = cliente_admin.post(
        "/api/revisao-de-cadastros/instituicoes/enquadramento",
        json={
            "ids": [str(uma.id), str(outra.id)],
            "categoria_publico_id": categoria.id,
        },
    ).json()

    assert len(corpo["recusadas"]) == 2, corpo
    for recusada in corpo["recusadas"]:
        assert "duplicata" in recusada["motivo"]
    #: NENHUMA DAS DUAS MUDOU: a pergunta é qual delas é a de verdade, e o lote
    #: não tem como responder.
    sessao.refresh(uma)
    sessao.refresh(outra)
    assert {uma.tipo, outra.tipo} == {"entidade", "veiculo"}


def test_o_LOTE_nao_alcanca_quem_JA_FOI_consolidado(cliente_admin, sessao):
    """A trava que mantém isto seguro, e ela é a razão de a rota viver sob a fila.

    UM LOTE QUE PUDESSE REESCREVER O TIPO DE QUALQUER INSTITUIÇÃO permitiria a um
    clique na aba de revisão varrer a curadoria de 3.785 cadastros feitos à mão. O
    que está pendente nasceu de planilha e ninguém conferiu — é exatamente o que
    este gesto serve. O resto se edita um a um, onde a tela mostra o cadastro
    inteiro.
    """
    categoria = _categoria_de(sessao, "orgao")
    curada = "Entidade Curada A Mao Zz94"
    registro = Instituicao(
        nome=curada, nome_normalizado=normalizar(curada), tipo="entidade", uf="RS"
    )
    sessao.add(registro)
    sessao.flush()
    assert registro.pendente_de_revisao is False

    corpo = cliente_admin.post(
        "/api/revisao-de-cadastros/instituicoes/enquadramento",
        json={"ids": [str(registro.id)], "categoria_publico_id": categoria.id},
    ).json()

    assert corpo["enquadradas"] == 0
    assert len(corpo["recusadas"]) == 1
    assert "consolidada" in corpo["recusadas"][0]["motivo"]
    sessao.refresh(registro)
    assert registro.tipo == "entidade"
    assert registro.categoria_publico_id is None


def test_o_LOTE_recusa_subcategoria_de_OUTRA_categoria(cliente_admin, sessao):
    """A conferência pesa mais aqui do que no cadastro um a um.

    Uma subcategoria que não pertence à categoria seria gravada em até 500 linhas
    de uma vez, e desfazer isso é outro lote. A FK não expressa o cruzamento —
    ele é regra de domínio.
    """
    categoria = _categoria_de(sessao, "orgao")
    estranha = sessao.scalar(
        select(SubcategoriaPublico).where(
            SubcategoriaPublico.categoria_publico_id != categoria.id,
            SubcategoriaPublico.ativo.is_(True),
        )
    )
    uma = _pendente(sessao, "Entidade Com Sub Errada Zz95")

    resposta = cliente_admin.post(
        "/api/revisao-de-cadastros/instituicoes/enquadramento",
        json={
            "ids": [str(uma.id)],
            "categoria_publico_id": categoria.id,
            "subcategoria_publico_id": estranha.id,
        },
    )
    assert resposta.status_code == 422, resposta.text
    assert "não pertence" in resposta.text

    #: E NADA FOI GRAVADO.
    sessao.refresh(uma)
    assert uma.categoria_publico_id is None


def test_o_LOTE_tem_TETO(cliente_admin, sessao):
    """O limite existe para haver um.

    Sem ele, o tamanho do pedido é decidido por quem o manda. Não é proteção
    contra abuso — a rota exige `administra_dicionarios` —, é a diferença entre
    uma operação cujo custo se conhece e uma que pode levar o tempo que quiser.
    """
    from uuid import uuid4

    from app.casos_de_uso.enquadrar_em_lote import MAXIMO_DO_LOTE

    categoria = _categoria_de(sessao, "orgao")
    resposta = cliente_admin.post(
        "/api/revisao-de-cadastros/instituicoes/enquadramento",
        json={
            "ids": [str(uuid4()) for _ in range(MAXIMO_DO_LOTE + 1)],
            "categoria_publico_id": categoria.id,
        },
    )
    assert resposta.status_code == 422, resposta.text
    assert str(MAXIMO_DO_LOTE) in resposta.text


def test_quem_NAO_administra_cadastros_nao_enquadra_em_lote(cliente_sem_admin, sessao):
    """A mesma guarda da leitura, e aqui ela protege uma escrita de 500 linhas."""
    assert (
        cliente_sem_admin.post(
            "/api/revisao-de-cadastros/instituicoes/enquadramento",
            json={"ids": [str(uuid.uuid4())], "categoria_publico_id": 1},
        ).status_code
        == 403
    )


# -- os três achados da revisão do lote ---------------------------------------


def test_o_PERFIL_DE_REDE_e_recusado_e_NAO_derruba_o_lote(cliente_admin, sessao):
    """ACHADO DE REVISÃO COM SEVERIDADE ALTA, e era alcançável de verdade.

    O perfil de rede carrega `cargo` (o que o fornecedor informou) e
    `interlocutor_id` (a pessoa de quem ele é), e o banco tem CHECK de que os dois
    só existem em `tipo = 'perfil_rede'` — migrations `0073` e `0076`.

    E A INGESTÃO DE MENÇÕES CRIA PERFIL DE REDE PENDENTE, então eles estão na
    fila: selecionar um derrubava o lote INTEIRO em `IntegrityError` no flush. Não
    era hipótese — era a fila real.

    RECUSAR, E NÃO LIMPAR OS CAMPOS: apagar o cargo e o vínculo em silêncio, num
    lote de 178, destruiria informação que ninguém pediu para destruir.

    E O RESTO DO LOTE PASSA, que é a promessa de "não é tudo-ou-nada".
    """
    categoria = _categoria_de(sessao, "orgao")
    perfil = Instituicao(
        nome="perfilnalote_zz96",
        nome_normalizado=normalizar("perfilnalote_zz96"),
        tipo="perfil_rede",
        cargo="Deputado estadual",
        pendente_de_revisao=True,
    )
    sessao.add(perfil)
    sessao.flush()
    boa = _pendente(sessao, "Entidade Que Passa Zz96")

    resposta = cliente_admin.post(
        "/api/revisao-de-cadastros/instituicoes/enquadramento",
        json={
            "ids": [str(perfil.id), str(boa.id)],
            "categoria_publico_id": categoria.id,
        },
    )
    #: 200, E NÃO 500: o lote não estoura.
    assert resposta.status_code == 200, resposta.text
    corpo = resposta.json()

    assert corpo["enquadradas"] == 1
    assert len(corpo["recusadas"]) == 1
    assert corpo["recusadas"][0]["nome"] == "perfilnalote_zz96"
    assert "perfil de rede" in corpo["recusadas"][0]["motivo"]

    #: O PERFIL FICOU INTACTO, cargo e tipo.
    sessao.refresh(perfil)
    assert perfil.tipo == "perfil_rede"
    assert perfil.cargo == "Deputado estadual"
    #: E A OUTRA PASSOU.
    sessao.refresh(boa)
    assert boa.tipo == "orgao"


def test_o_lote_TRAVA_as_linhas_que_vai_mexer(cliente_admin, sessao):
    """ACHADO DE REVISÃO COM SEVERIDADE ALTA: era um lost update.

    O lote lia as linhas, testava `pendente_de_revisao` EM MEMÓRIA e só gravava no
    flush. Dois lotes sobre ids que se cruzam passavam os dois pelo teste antes de
    qualquer commit, e o segundo reescrevia categoria, subcategoria e TIPO de uma
    linha que o primeiro já havia consolidado — quebrando em silêncio a trava "só
    o que está na fila", com as duas requisições respondendo 200.

    O QUE ESTE TESTE AFIRMA é que a leitura pede o lock. Provar a corrida de
    verdade exigiria duas conexões e um `commit` no meio, e a fixture da suíte
    roda tudo numa transação que termina em rollback — duas sessões não veriam
    uma à outra.

    ENTÃO A ÂNCORA É A FONTE, e ela pega o único modo como este conserto se
    perde: alguém reescrevendo a consulta sem o `with_for_update`. Uma primeira
    versão deste teste também compilava a consulta e procurava "FOR UPDATE" no
    SQL — isso testava o SQLAlchemy, não este código, e não dizia nada sobre o
    lote.
    """
    import inspect

    from app.casos_de_uso import enquadrar_em_lote

    fonte = inspect.getsource(enquadrar_em_lote.enquadrar)
    assert "with_for_update()" in fonte, (
        "a leitura do lote precisa travar as linhas: sem isso, dois lotes "
        "cruzados reescrevem o que o outro consolidou"
    )


def test_a_categoria_que_NAO_define_tipo_e_recusada_com_a_mensagem_certa(
    cliente_admin, sessao
):
    """ACHADO DE REVISÃO: a mensagem de `derivar_tipo` mente neste contexto.

    Ela diz "Informe a categoria de público", que é verdade no cadastro um a um —
    onde a categoria pode vir vazia — e MENTIRA aqui: a categoria foi informada, e
    o que falta é ela mapear para um tipo. Quem lê "informe a categoria" depois de
    ter escolhido uma vai procurar o erro no lugar errado.
    """
    from app.dominio.frentes import TIPO_DA_CATEGORIA_DE_PUBLICO

    sem_tipo = sessao.scalar(
        select(CategoriaPublico).where(
            CategoriaPublico.codigo.notin_(list(TIPO_DA_CATEGORIA_DE_PUBLICO)),
            CategoriaPublico.ativo.is_(True),
        )
    )
    if sem_tipo is None:
        pytest.skip("toda categoria ativa define tipo — nada a afirmar aqui")

    uma = _pendente(sessao, "Entidade Sem Tipo Zz97")
    resposta = cliente_admin.post(
        "/api/revisao-de-cadastros/instituicoes/enquadramento",
        json={"ids": [str(uma.id)], "categoria_publico_id": sem_tipo.id},
    )
    assert resposta.status_code == 422, resposta.text
    assert "não define tipo" in resposta.text
    assert "Informe a categoria" not in resposta.text


def test_o_id_REPETIDO_no_pedido_nao_vira_duas_recusas(cliente_admin, sessao):
    """A deduplicação que substituiu o "cinto" do laço.

    O cinto era `ocupados.add(...)` dentro do laço, e a revisão apontou que ele
    "sugeria uma proteção que não resolve concorrência" — o id repetido era o
    único caso que ele cobria, e a contagem `quantas_vezes` já pegava duas linhas
    DIFERENTES de mesmo nome. Deduplicar o pedido resolve o caso real e não
    promete nada sobre concorrência, que é o que o savepoint e o lock fazem.
    """
    categoria = _categoria_de(sessao, "orgao")
    uma = _pendente(sessao, "Entidade Repetida Zz98")

    corpo = cliente_admin.post(
        "/api/revisao-de-cadastros/instituicoes/enquadramento",
        json={
            "ids": [str(uma.id), str(uma.id), str(uma.id)],
            "categoria_publico_id": categoria.id,
        },
    ).json()

    assert corpo["enquadradas"] == 1, corpo
    assert corpo["recusadas"] == []
