"""O `.xlsx` que a pessoa baixa para revisar a taxonomia de subtemas.

ELE NASCE PREENCHIDO, e essa é a decisão central deste módulo. O modelo de
agendas sai vazio porque uma agenda nova não existe em lugar nenhum; a
taxonomia existe inteira no banco, e o trabalho real não é cadastrar 150
subtemas — é mexer em três. Um modelo vazio obrigaria a pessoa a redigitar 147
linhas certas para corrigir as três, e cada redigitação é uma chance de errar
num subtema que estava bom.

O efeito colateral é a ida e volta: baixar o modelo e subi-lo sem tocar em nada
tem de dar "tudo igual, nada a fazer". Isso é teste e é garantia — se a volta
não for idêntica, o formato está perdendo informação em algum lugar.

O FORMATO MORA EM `app/dominio/importacao_de_subtemas.py`, e este módulo apenas
PERCORRE `FORMATO`. Mesmo raciocínio de `modelo_de_importacao.py`: uma coluna
nova entra na descrição e aparece aqui e no leitor, em vez de entrar numa cópia
e não na outra.

`vocabularios` E `linhas` CHEGAM PRONTOS, sem sessão de banco aqui — quem
consulta é a rota. É o que mantém `gerar` testável com dicionários comuns.
"""

from __future__ import annotations

import io
from collections.abc import Mapping, Sequence

from app.dominio.erros import RegraViolada
from app.dominio.importacao_de_subtemas import (
    ABA_PRINCIPAL,
    FORMATO,
    VOCABULARIOS,
)

#: O nome que a pessoa vê na pasta de downloads.
NOME_DO_MODELO = "taxonomia-de-subtemas.xlsx"

#: O rótulo de cada aba de vocabulário. Sem isto a aba se chamaria
#: `camadas_lso`, e ninguém reconhece a própria taxonomia em snake_case.
ROTULO_DA_ABA = {
    "blocos_tema": "Pilares",
    "macro_temas": "Temas estratégicos",
    "camadas_lso": "LSO",
    "sim_nao": "Sim ou Não",
    "riscos": "Matriz de risco",
}

#: Quantas linhas a mais ganham lista suspensa, abaixo da última preenchida.
#:
#: As suspensas têm de alcançar as linhas que a pessoa VAI escrever, não só as
#: que já existem. Cem é folga confortável contra o teto de 400 do leitor.
FOLGA_DE_LINHAS = 100


def _valor_e_nota(item: str | Sequence[str]) -> tuple[str, str]:
    """Aceita `"Tarifa"` ou `("R01", "Reajuste negado")`.

    O PAR EXISTE POR CAUSA DOS RISCOS. "R17" não diz a ninguém qual risco é, e
    a aba é justamente onde a pessoa descobre — mas a COLUNA aceita só o
    código. Juntar os dois num texto ("R01 — Reajuste negado") colocaria na
    lista suspensa um valor que a conferência recusa, o que é pior que não ter
    a dica: a pessoa escolheria da lista oferecida e veria a linha recusada.
    Então o valor fica na coluna A, sozinho, e a explicação na B.
    """
    if isinstance(item, str):
        return item, ""
    valor, *resto = item
    return str(valor), str(resto[0]) if resto else ""


def gerar(
    vocabularios: Mapping[str, Sequence[str | Sequence[str]]],
    linhas: Sequence[Mapping[str, str]] = (),
) -> bytes:
    """A aba `Subtemas` preenchida, mais uma aba por vocabulário.

    `linhas` é uma sequência de dicionários coluna → texto, na ordem em que
    devem sair. A chave é o NOME da coluna (`"Pilar (N1)"`), e não o campo:
    quem monta `linhas` é a rota, que lê do banco, e amarrar pelo nome mantém
    `FORMATO` como a única fonte da ordem das colunas.

    Cada valor de vocabulário é um texto ou um par `(valor, explicação)`.
    """
    try:
        from openpyxl import Workbook
        from openpyxl.comments import Comment
        from openpyxl.styles import Font
        from openpyxl.utils import get_column_letter
        from openpyxl.worksheet.datavalidation import DataValidation
    except ModuleNotFoundError as erro:  # pragma: no cover - dependência declarada
        raise RegraViolada("Leitura de planilha indisponível neste servidor.") from erro

    pasta = Workbook()
    pasta.remove(pasta.active)  # a Workbook nasce com uma aba "Sheet" que ninguém pediu
    negrito = Font(bold=True)

    principal = pasta.create_sheet(ABA_PRINCIPAL)
    principal.append([coluna.nome for coluna in FORMATO])

    # A DICA VAI DE COMENTÁRIO NA CÉLULA, nunca de linha extra acima do
    # cabeçalho: o leitor procura o cabeçalho nas 20 primeiras linhas e casa os
    # nomes contra `FORMATO`, e uma célula sobrando na linha 1 é exatamente a
    # divergência que a descrição única existe para evitar.
    for indice, coluna in enumerate(FORMATO, start=1):
        celula = principal.cell(row=1, column=indice)
        celula.font = negrito
        if coluna.dica:
            celula.comment = Comment(coluna.dica, "Painel Reputacional")

    for linha in linhas:
        principal.append([linha.get(coluna.nome, "") for coluna in FORMATO])

    principal.freeze_panes = "A2"
    for indice, coluna in enumerate(FORMATO, start=1):
        principal.column_dimensions[get_column_letter(indice)].width = max(
            16, min(44, len(coluna.nome) + 8)
        )

    # -- uma aba por vocabulário, todas FECHADAS -------------------------------
    # A importação não cria pilar nem risco: ela ESCOLHE do que existe. A
    # proteção da aba torna isso visível a quem abre o arquivo — embora o
    # servidor recuse valor novo mesmo que alguém desproteja a aba, porque o
    # controle é a conferência, não o Excel.
    ultima = 1 + len(linhas) + FOLGA_DE_LINHAS
    nomes_das_colunas = [coluna.nome for coluna in FORMATO]

    for chave in VOCABULARIOS:
        pares = [_valor_e_nota(item) for item in vocabularios.get(chave, ())]
        valores = [valor for valor, _ in pares]
        rotulo = ROTULO_DA_ABA[chave]
        aba = pasta.create_sheet(rotulo)
        aba.append([rotulo, "O que é"])
        aba["A1"].font = negrito
        aba["B1"].font = negrito
        for valor, nota in pares:
            aba.append([valor, nota])
        aba.column_dimensions["A"].width = 30
        aba.column_dimensions["B"].width = 60
        aba.protection.sheet = True

        coluna_do_formato = next(
            (coluna for coluna in FORMATO if coluna.vocabulario == chave), None
        )
        if coluna_do_formato is None or not valores:
            # VOCABULÁRIO VAZIO NÃO PODE QUEBRAR O ARQUIVO. Uma base nova pode
            # não ter risco cadastrado, e o modelo ainda tem de abrir — a aba
            # sai presente e vazia, sem suspensa ligada a um intervalo nulo,
            # que o Excel rejeita na abertura.
            continue

        letra = get_column_letter(nomes_das_colunas.index(coluna_do_formato.nome) + 1)
        suspensa = DataValidation(
            type="list",
            formula1=f"'{rotulo}'!$A$2:$A${1 + len(valores)}",
            allow_blank=True,
            showDropDown=False,
        )
        # `showErrorMessage` DESLIGADO na coluna de riscos, de propósito: ela
        # aceita "R01, R07", que nenhuma lista de valor único casa. A suspensa
        # ali é consulta ("quais códigos existem?"), não trava — quem recusa
        # código inexistente é a conferência, que sabe separar por vírgula.
        suspensa.showErrorMessage = chave != "riscos"
        suspensa.errorTitle = "Valor fora da lista"
        suspensa.error = f"Escolha um dos valores da aba {rotulo}."
        principal.add_data_validation(suspensa)
        suspensa.add(f"{letra}2:{letra}{ultima}")

    arquivo = io.BytesIO()
    pasta.save(arquivo)
    return arquivo.getvalue()
