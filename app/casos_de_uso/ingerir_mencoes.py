"""A planilha do fornecedor entrando no índice.

O que este módulo acrescenta ao `dominio/ingestao_score.py`: abrir o `.xlsx`
(openpyxl), achar a aba, casar o cabeçalho com o mapeamento gravado e trocar o
mês inteiro no banco. A regra de o que vira menção NÃO mora aqui.

UM ARQUIVO PODE ALIMENTAR VÁRIAS FONTES, e o upload trata todas de uma vez. O
export da Clipei vira a lente Imprensa inteira e, recortado por público
investidor, a lente Mercado; o da Approach traz Social Listening e Community
Management em abas diferentes do mesmo `.xlsx`. Importar por uma fonte só
deixaria a irmã com o mês anterior, e duas lentes passariam a ler versões
diferentes do mesmo arquivo — uma divergência que ninguém veria na tela, porque
cada lente mostraria um número plausível. Quem diz quais fontes andam juntas é
o campo `arquivo` do mapeamento, que é cadastro.

SUBSTITUIÇÃO POR MÊS, e não acréscimo. O fornecedor reenvia o arquivo quando
corrige uma classificação — e o mesmo post reenviado, somado, contaria duas
vezes. Cada mês presente no arquivo é apagado e regravado; um mês que o arquivo
não traz fica intacto, porque o export de junho não é uma afirmação sobre maio.

O PREÇO DISSO É UM ARQUIVO PARCIAL PODER ENCOLHER UM MÊS: um export baixado
antes do fechamento substitui o mês cheio pelo pedaço. Não dá para o servidor
distinguir "o fornecedor reclassificou e sobraram menos" de "baixaram cedo
demais" — os dois chegam como um arquivo com menos linhas. O que dá é NÃO
DEIXAR ISSO PASSAR CALADO: o resumo diz quantas menções o mês tinha antes de a
troca acontecer, ao lado de quantas entraram.

POR QUE O AGREGADO SE REFAZ AQUI. `score_mes_fonte` guarda as quatro somas de
que toda régua precisa. Como a substituição deixa o banco com exatamente as
menções do arquivo para aqueles meses, as somas do arquivo JÁ SÃO as somas do
mês — não há por que relê-las do banco para conferir a própria escrita.

UMA CONSEQUÊNCIA A CARREGAR: `soma_log` e `soma_cargo` congelam, no instante da
ingestão, as fórmulas de `dominio/score.py` (`peso_do_engajamento` e
`peso_do_cargo`). Mudar uma delas NÃO reescreve o passado. É de propósito — o
número que a diretoria citou em julho continua sendo o que ela citou —, e o
grão para reprocessar, se um dia for preciso, está inteiro em `mencao`.
"""

from __future__ import annotations

import io
from collections.abc import Iterator, Mapping
from dataclasses import dataclass, field
from datetime import date

from sqlalchemy import delete, func, select
from sqlalchemy.orm import Session

from app.banco.tabelas_score import Mencao, ScoreFonte, ScoreMesFonte
from app.dominio.erros import RegraViolada
from app.dominio.ingestao_score import Leitura, Mapeamento, ler_planilha, somar

#: Abrir uma planilha de 10 mil linhas é barato; abrir um arquivo de 200 MB
#: enviado por engano não é. O maior dos quatro exports de hoje tem 1,4 MB.
TAMANHO_MAXIMO = 40 * 1024 * 1024

#: A assinatura de um ZIP — e um `.xlsx` é um ZIP. Conferir os quatro primeiros
#: bytes recusa o `.xls` antigo, o CSV renomeado e o PDF trocado por engano
#: ANTES de entregar o conteúdo ao parser de XML, que é a parte que um arquivo
#: hostil tenta alcançar. Não substitui o `defusedxml` (declarado no
#: `pyproject.toml`, que o openpyxl usa sozinho quando presente): é a primeira
#: das duas barreiras, não a única.
ASSINATURA_ZIP = b"PK\x03\x04"


@dataclass(frozen=True, slots=True)
class Resumo:
    """O que a ingestão fez com UMA fonte — o que a tela mostra depois."""

    fonte: str
    nome: str
    linhas: int
    ingeridas: int
    descartes: Mapping[str, int]
    meses: tuple[date, ...]
    #: Quantas menções esta fonte tinha nestes meses ANTES da substituição.
    #: É o que deixa visível um export parcial encolhendo um mês fechado.
    antes: int = 0
    avisos: Mapping[str, int] = field(default_factory=dict)


def _linhas_da_aba(
    conteudo: bytes, mapeamento: Mapeamento
) -> Iterator[Mapping[str, object]]:
    """As linhas da planilha como dicionários cabeçalho → célula."""
    try:
        from openpyxl import load_workbook
    except ModuleNotFoundError as erro:  # pragma: no cover - dependência declarada
        raise RegraViolada("Leitura de planilha indisponível neste servidor.") from erro

    # `read_only` não carrega a planilha inteira na memória, e `data_only` traz
    # o VALOR de uma célula com fórmula — sem ele, a coluna de engajamento
    # calculada chegaria como a string "=SOMA(...)".
    try:
        planilha = load_workbook(io.BytesIO(conteudo), read_only=True, data_only=True)
    except Exception as erro:
        # Um `.xls` antigo, um CSV renomeado ou um upload truncado chegam aqui.
        # Sem a tradução, o openpyxl devolveria 500 e a pessoa leria "erro
        # interno" para um arquivo que ela mesma pode trocar.
        raise RegraViolada(
            "Não consegui abrir o arquivo como planilha .xlsx. "
            "Se ele veio em .xls ou .csv, salve como .xlsx e envie de novo."
        ) from erro
    try:
        if mapeamento.aba and mapeamento.aba in planilha.sheetnames:
            aba = planilha[mapeamento.aba]
        elif mapeamento.aba:
            raise RegraViolada(
                f"A planilha não tem a aba {mapeamento.aba!r}. "
                f"Abas encontradas: {', '.join(planilha.sheetnames)}."
            )
        else:
            aba = planilha[planilha.sheetnames[0]]

        linhas = aba.iter_rows(values_only=True)
        try:
            cabecalho = [
                str(celula).strip() if celula is not None else ""
                for celula in next(linhas)
            ]
        except StopIteration:
            raise RegraViolada("A planilha está vazia.") from None

        faltando = sorted(mapeamento.colunas_necessarias - set(cabecalho))
        if faltando:
            raise RegraViolada(
                "A planilha não tem as colunas que o cadastro desta fonte espera: "
                f"{', '.join(faltando)}."
            )

        for linha in linhas:
            yield dict(zip(cabecalho, linha, strict=False))
    finally:
        planilha.close()


def _quanto_havia(sessao: Session, fonte: ScoreFonte, meses: list[date]) -> int:
    total = sessao.scalar(
        select(func.count())
        .select_from(Mencao)
        .where(Mencao.fonte_id == fonte.id, Mencao.mes.in_(meses))
    )
    return int(total or 0)


def _regravar(sessao: Session, fonte: ScoreFonte, leitura: Leitura) -> None:
    """Troca os meses que o arquivo traz — menções e agregado."""
    meses = list(leitura.meses)
    sessao.execute(
        delete(Mencao).where(Mencao.fonte_id == fonte.id, Mencao.mes.in_(meses))
    )
    sessao.execute(
        delete(ScoreMesFonte).where(
            ScoreMesFonte.fonte_id == fonte.id, ScoreMesFonte.mes.in_(meses)
        )
    )

    sessao.add_all(
        Mencao(
            fonte_id=fonte.id,
            mes=mencao.mes,
            data=mencao.data,
            sentimento=mencao.sentimento,
            tier=mencao.tier,
            engajamento=mencao.engajamento,
            cargo=mencao.cargo,
            atributo=mencao.atributo,
            veiculo=mencao.veiculo,
            publico_alvo=mencao.publico_alvo,
            tema_texto=mencao.tema_texto,
            unidade_texto=mencao.unidade_texto,
            teor=mencao.teor,
            acionavel=mencao.acionavel,
            autor=mencao.autor,
        )
        for mencao in leitura.mencoes
    )
    sessao.add_all(
        ScoreMesFonte(
            fonte_id=fonte.id,
            mes=soma.mes,
            sentimento=soma.sentimento,
            tier=soma.tier,
            mencoes=soma.mencoes,
            soma_log=soma.soma_log,
            soma_engajamento=soma.soma_engajamento,
            soma_cargo=soma.soma_cargo,
        )
        for soma in somar(leitura.mencoes)
    )
    sessao.flush()


def _mapeamento_de(fonte: ScoreFonte) -> Mapeamento:
    try:
        return Mapeamento.de_json(fonte.mapeamento_colunas)
    except ValueError as erro:
        # O mapeamento é cadastro, não entrada de usuário: um mapeamento
        # quebrado é problema de configuração da fonte, e a mensagem precisa
        # dizer isso, senão quem subiu o arquivo procura o defeito no arquivo.
        raise RegraViolada(
            f"O mapeamento de colunas de {fonte.nome} está inválido: {erro}"
        ) from erro


def _ingerir_uma(sessao: Session, fonte: ScoreFonte, conteudo: bytes) -> Resumo:
    """Lê a planilha pelo mapeamento desta fonte e substitui os meses dela."""
    mapeamento = _mapeamento_de(fonte)
    leitura = ler_planilha(_linhas_da_aba(conteudo, mapeamento), mapeamento)
    if not leitura.mencoes:
        raise RegraViolada(
            f"Nenhuma linha da planilha virou menção em {fonte.nome}. "
            f"Lidas {leitura.linhas}; descartes: {dict(leitura.descartes)}."
        )

    antes = _quanto_havia(sessao, fonte, list(leitura.meses))
    _regravar(sessao, fonte, leitura)
    return Resumo(
        fonte=fonte.codigo,
        nome=fonte.nome,
        linhas=leitura.linhas,
        ingeridas=len(leitura.mencoes),
        descartes=leitura.descartes,
        meses=leitura.meses,
        antes=antes,
        avisos={motivo: total for motivo, total in leitura.avisos.items() if total},
    )


def fontes_do_mesmo_arquivo(sessao: Session, fonte: ScoreFonte) -> list[ScoreFonte]:
    """As fontes que leem o export que esta fonte lê — ela inclusive.

    Sem `arquivo` no mapeamento, a fonte anda sozinha: é o caso de um
    fornecedor que entrega um arquivo só dele.
    """
    grupo = _mapeamento_de(fonte).arquivo
    if not grupo:
        return [fonte]

    irmas = sessao.scalars(
        select(ScoreFonte)
        .where(
            ScoreFonte.ativo.is_(True),
            ScoreFonte.interna.is_(False),
            ScoreFonte.mapeamento_colunas["arquivo"].astext == grupo,
        )
        .order_by(ScoreFonte.ordem)
    ).all()
    # A fonte pedida entra mesmo que o filtro não a alcance (inativa, por
    # exemplo, ela já foi recusada antes): quem subiu escolheu esta.
    return list(irmas) if fonte in irmas else [fonte, *irmas]


def ingerir(sessao: Session, fonte: ScoreFonte, conteudo: bytes) -> list[Resumo]:
    """Lê o export e substitui os meses dele em TODAS as fontes que o leem."""
    if fonte.interna:
        raise RegraViolada(
            f"{fonte.nome} é fonte interna: o dado já está neste banco e não se importa."
        )
    if not fonte.ativo:
        raise RegraViolada(f"{fonte.nome} está inativa. Reative-a antes de importar.")
    if len(conteudo) > TAMANHO_MAXIMO:
        raise RegraViolada("A planilha passa de 40 MB.")
    if not conteudo.startswith(ASSINATURA_ZIP):
        raise RegraViolada(
            "O arquivo não é uma planilha .xlsx. "
            "Se ele veio em .xls ou .csv, salve como .xlsx e envie de novo."
        )

    return [
        _ingerir_uma(sessao, irma, conteudo)
        for irma in fontes_do_mesmo_arquivo(sessao, fonte)
    ]
