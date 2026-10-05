"""A carga inicial das Lentes, a partir da planilha do pacote de produção.

O QUE ESTE MÓDULO É. O pacote de produção das Lentes (handoff de out/2026) veio
com 29.414 itens reais numa planilha de uma aba por lente, nas 40 colunas do
padrão Aegea. Isto lê uma dessas abas e põe as menções no banco, no vocabulário
que o índice já usa.

O QUE ELE NÃO É: o importador mensal. Aquele é `casos_de_uso/ingerir_mencoes`,
lê UMA aba por fonte pelo mapeamento gravado em `score_fonte`, e é o caminho que
a coordenação usa todo mês — com o checklist de bloqueios que o padrão pede.
Aqui a planilha tem uma aba por LENTE, com as fontes misturadas na coluna
`fonte`, e é carga de uma vez. Forçar as duas pelo mesmo caminho faria o fluxo
mensal carregar um parâmetro que só a carga inicial usa.

O QUE OS DOIS DIVIDEM é o que importa: `MencaoLida` (o vocabulário do índice),
`somar` (as quatro somas do agregado) e `regravar` (a substituição por mês). A
régua de o que vira menção é uma só, e é a do domínio.

UMA LENTE POR VEZ. `PACOTE` tem uma linha por lente implementada, e o dono do
produto pediu a Sociedade digital primeiro. As outras entram quando a etapa
delas chegar — acrescentar a linha aqui é o que a carga precisa saber.
"""

from __future__ import annotations

import io
from collections import defaultdict
from dataclasses import dataclass, field
from datetime import date

from openpyxl import load_workbook
from sqlalchemy import select
from sqlalchemy.orm import Session

from app.banco.tabelas_score import ScoreFonte
from app.casos_de_uso.ingerir_mencoes import ASSINATURA_ZIP, TAMANHO_MAXIMO, regravar
from app.dominio.erros import RegraViolada
from app.dominio.ingestao_score import (
    SENTIMENTOS,
    Leitura,
    MencaoLida,
    achatar,
    para_data,
    para_inteiro,
)


@dataclass(frozen=True, slots=True)
class AbaDoPacote:
    """Onde mora a lente na planilha, e como os nomes de fonte se traduzem."""

    aba: str
    #: Nome na coluna `fonte` da planilha → `score_fonte.codigo` deste banco. É
    #: tradução de cadastro, não de código: a planilha diz "Approach SL" e a
    #: fonte se chama `approach_sl` aqui.
    fontes: dict[str, str]


#: As lentes que a carga sabe ler, uma por etapa do rollout.
PACOTE: dict[str, AbaDoPacote] = {
    "sociedade": AbaDoPacote(
        aba="item_sociedade_digital",
        fontes={"Bites": "bites", "Approach SL": "approach_sl"},
    ),
}

#: Coluna da planilha → campo de `MencaoLida`. É aqui que o padrão Aegea e o
#: vocabulário do índice se encontram, e as cinco primeiras linhas são a razão
#: de este mapa existir: os dois chamam a mesma coisa por nomes diferentes.
CAMPOS: dict[str, str] = {
    "classificacao": "sentimento",
    "veiculo_rede": "veiculo",
    "tema": "tema_texto",
    "empresa_citada": "unidade_texto",
    "cargo_autor": "cargo",
    # Daqui para baixo o nome é o mesmo nos dois lados.
    "data": "data",
    "autor": "autor",
    "engajamento": "engajamento",
    "publico_alvo": "publico_alvo",
    "atributo": "atributo",
    "teor": "teor",
    "id_fonte": "id_fonte",
    "uf": "uf",
    "subtema": "subtema",
    "perfil_autor": "perfil_autor",
    "titulo_texto": "titulo_texto",
    "link": "link",
    "peso_tier": "peso_tier",
}


@dataclass(frozen=True, slots=True)
class LeituraDoPacote:
    """O que a aba rendeu, separado pela fonte que mandou cada linha."""

    mencoes: tuple[MencaoLida, ...]
    #: Código da fonte → as menções dela. É por fonte que o agregado e a
    #: substituição por mês acontecem.
    por_fonte: dict[str, tuple[MencaoLida, ...]]
    descartes: dict[str, int] = field(default_factory=dict)
    #: (código da fonte, mês) → quantas linhas chegaram sem classificação.
    nao_classificadas: dict[tuple[str, date], int] = field(default_factory=dict)


@dataclass(frozen=True, slots=True)
class ResumoDaCarga:
    """O que a carga fez — é o que o CLI imprime e o teste confere."""

    lente: str
    gravadas: int
    por_fonte: dict[str, int]
    meses: tuple[date, ...]
    descartes: dict[str, int]


def _mes_do_padrao(valor: object) -> date | None:
    """`2026-06` → 1º de junho de 2026.

    O MÊS VEM DA COLUNA `mes`, e não da `data`: 86% dos itens da Sociedade
    chegam sem data, e derivar o mês dela jogaria todos eles para fora da conta.
    """
    if isinstance(valor, date):
        return date(valor.year, valor.month, 1)
    texto = str(valor or "").strip()
    if len(texto) < 7:
        return None
    try:
        return date(int(texto[:4]), int(texto[5:7]), 1)
    except ValueError:
        return None


def _texto(valor: object) -> str | None:
    """O valor como texto, ou nulo quando a célula não diz nada.

    `-` CONTA COMO VAZIO porque é o que a planilha usa para "não se aplica", e
    guardá-lo como texto faria um traço virar uma categoria na decomposição —
    uma linha chamada "-" no corte por tema.
    """
    if valor is None:
        return None
    texto = str(valor).strip()
    return texto or None if texto != "-" else None


def ler_aba_da_lente(conteudo: bytes, lente: str) -> LeituraDoPacote:
    """A aba daquela lente, lida no vocabulário do índice."""
    if lente not in PACOTE:
        tem = ", ".join(sorted(PACOTE))
        raise RegraViolada(f"A carga não conhece a lente {lente!r}. Tem: {tem}.")
    descricao = PACOTE[lente]

    if len(conteudo) > TAMANHO_MAXIMO:
        raise RegraViolada("A planilha do pacote passa do tamanho que se aceita ler.")
    if not conteudo.startswith(ASSINATURA_ZIP):
        raise RegraViolada("O arquivo não é um .xlsx — os primeiros bytes não são de ZIP.")

    pasta = load_workbook(io.BytesIO(conteudo), read_only=True, data_only=True)
    if descricao.aba not in pasta.sheetnames:
        raise RegraViolada(
            f"A planilha não tem a aba {descricao.aba!r}, que é a da lente {lente!r}."
        )
    folha = pasta[descricao.aba]

    linhas = folha.iter_rows(values_only=True)
    cabecalho = [achatar(str(celula or "")) for celula in next(linhas, ())]
    onde = {nome: posicao for posicao, nome in enumerate(cabecalho)}

    #: SEM ESTAS DUAS NÃO HÁ LEITURA POSSÍVEL: sem `mes` não se sabe de que mês é
    #: a linha, e sem `fonte` não se sabe de quem ela é — e é por fonte que a
    #: substituição acontece.
    for obrigatoria in ("mes", "fonte"):
        if obrigatoria not in onde:
            raise RegraViolada(
                f"A aba {descricao.aba!r} não tem a coluna {obrigatoria!r}."
            )

    por_fonte: dict[str, list[MencaoLida]] = defaultdict(list)
    descartes: dict[str, int] = defaultdict(int)
    nao_classificadas: dict[tuple[str, date], int] = defaultdict(int)

    for linha in linhas:
        #: A LINHA VIRA UM DICIONÁRIO POR NOME DE COLUNA, e não uma função que
        #: fecha sobre a variável do laço: closure sobre `linha` é o defeito que o
        #: `B023` do ruff aponta, e ele tem razão — basta alguém guardar a função
        #: para usá-la depois do laço e ela passa a ler a última linha de todas.
        celula = {
            nome: linha[posicao]
            for nome, posicao in onde.items()
            if posicao < len(linha)
        }.get

        nome_da_fonte = _texto(celula("fonte"))
        codigo = descricao.fontes.get(nome_da_fonte or "")
        if codigo is None:
            descartes["fonte desconhecida"] += 1
            continue

        mes = _mes_do_padrao(celula("mes"))
        if mes is None:
            descartes["mês ilegível"] += 1
            continue

        sentimento = SENTIMENTOS.get(achatar(str(celula("classificacao") or "")))
        if sentimento is None:
            #: VOLUME SEM LEITURA é um estado próprio, e a tela o desenha
            #: diferente de "não houve nada". Por isso conta, em vez de só sair.
            descartes["sem classificação"] += 1
            nao_classificadas[(codigo, mes)] += 1
            continue

        valores: dict[str, object] = {}
        for coluna, campo in CAMPOS.items():
            if coluna in ("classificacao", "data", "engajamento", "peso_tier"):
                continue
            valores[campo] = _texto(celula(coluna))

        por_fonte[codigo].append(
            MencaoLida(
                mes=mes,
                sentimento=sentimento,
                data=para_data(celula("data")),
                engajamento=para_inteiro(celula("engajamento")),
                peso_tier=float(para_inteiro(celula("peso_tier")) or 1),
                **valores,  # type: ignore[arg-type]
            )
        )

    pasta.close()
    todas = tuple(mencao for lista in por_fonte.values() for mencao in lista)
    return LeituraDoPacote(
        mencoes=todas,
        por_fonte={codigo: tuple(lista) for codigo, lista in por_fonte.items()},
        descartes=dict(descartes),
        nao_classificadas=dict(nao_classificadas),
    )


def carregar(sessao: Session, conteudo: bytes, lente: str) -> ResumoDaCarga:
    """Lê a aba e troca, no banco, os meses que ela traz — por fonte.

    SUBSTITUIÇÃO E NÃO ACRÉSCIMO, que é o que torna a carga reexecutável:
    consertar a planilha e rodar de novo deixa o banco com exatamente o que o
    arquivo diz. Somar duplicaria o mês e a nota continuaria plausível — o pior
    tipo de erro, porque não aparece na tela.
    """
    leitura = ler_aba_da_lente(conteudo, lente)

    gravadas = 0
    por_fonte: dict[str, int] = {}
    meses: set[date] = set()
    for codigo, mencoes in sorted(leitura.por_fonte.items()):
        fonte = sessao.scalars(
            select(ScoreFonte).where(ScoreFonte.codigo == codigo)
        ).one_or_none()
        if fonte is None:
            raise RegraViolada(
                f"A fonte {codigo!r} não está cadastrada em `score_fonte`. "
                "A carga não cria fonte: fornecedor novo é cadastro, não importação."
            )
        da_fonte = {
            mes: total
            for (fonte_do_par, mes), total in leitura.nao_classificadas.items()
            if fonte_do_par == codigo
        }
        regravar(
            sessao,
            fonte,
            Leitura(
                mencoes=mencoes,
                descartes={},
                linhas=len(mencoes),
                nao_classificadas=da_fonte,
            ),
        )
        gravadas += len(mencoes)
        por_fonte[codigo] = len(mencoes)
        meses.update(mencao.mes for mencao in mencoes)

    return ResumoDaCarga(
        lente=lente,
        gravadas=gravadas,
        por_fonte=por_fonte,
        meses=tuple(sorted(meses)),
        descartes=leitura.descartes,
    )
