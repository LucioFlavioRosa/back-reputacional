"""Bytes de um `.xlsx` viram linhas de células. UM leitor para a plataforma.

POR QUE UM SÓ: a plataforma tinha cinco leitores e três comportamentos — um
estimava a memória e caía num caminho lento, outro estimava e RECUSAVA, os
demais abriam direto. O mesmo arquivo era lido de jeitos diferentes conforme a
tela por onde subia, e um export de 649 KB chegou a ser recusado por "memória"
numa tela que a outra teria lido sem reclamar.

O MECANISMO: calamine primeiro, openpyxl de reserva, e as células sempre na
mesma forma — é o que permite a tudo rio abaixo não saber qual leitor trouxe a
célula.

A DIFERENÇA ENTRE ELES NÃO É SÓ VELOCIDADE. O openpyxl obedece à dimensão
DECLARADA no cabeçalho do arquivo e preenche a linha inteira com `None` até a
última coluna; o calamine devolve a área realmente usada. No export de bonds
isso deu 3,8 milhões de células contra 121 mil, e 36,6 MB contra 2,3 MB.

MAS "ÁREA USADA" NÃO É "POUCA CÉLULA": basta um valor no canto inferior direito
para o retângulo do calamine ficar tão grande quanto o do openpyxl. Este módulo
é mais barato no caso comum, não imune ao caso ruim — ver
`ingerir_mencoes.MEMORIA_POR_CELULA_VAZIA`.

A POLÍTICA NÃO MORA AQUI. Este módulo lê; quem chama decide o que fazer quando
não cabe — guardar ou não a leitura, recusar ou degradar.
"""

from __future__ import annotations

import io

from app.dominio.erros import RegraViolada

#: O QUE O EXCEL ESCREVE NUMA CÉLULA QUE DEU ERRO (`<c t="e"><v>#N/A</v></c>`).
#:
#: OS DOIS LEITORES DISCORDAM: o openpyxl devolve a string (`'#N/A'`), o
#: calamine devolve vazio. Sem acerto, um ticker com `#N/A` no catálogo dos
#: bonds some pelo caminho rápido e vira `PapelLido(ticker="#N/A")` pelo de
#: reserva — dois caminhos discordando sobre o mesmo arquivo.
#:
#: O ACERTO É PELO VAZIO, e é forçado: o calamine não diz QUAL erro era, então
#: não há como fazer o rápido parecer com o antigo. É também o sentido certo —
#: célula que deu erro não tem valor, e tratá-la como o texto "#N/A" é deixar
#: lixo de planilha virar dado.
ERROS_DO_EXCEL = frozenset(
    {
        "#N/A",
        "#DIV/0!",
        "#VALUE!",
        "#REF!",
        "#NAME?",
        "#NUM!",
        "#NULL!",
        "#SPILL!",
        "#CALC!",
        "#GETTING_DATA",
    }
)


def sem_erro_do_excel(linha: tuple) -> tuple:
    """A linha do openpyxl com as células de erro viradas `None`.

    É o lado openpyxl de `ERROS_DO_EXCEL` — ver lá por que o acerto é neste
    sentido. Devolve a MESMA tupla quando não há erro nenhum, que é o caso de
    quase toda linha: comparar é mais barato que reconstruir.
    """
    if not any(type(c) is str and c in ERROS_DO_EXCEL for c in linha):
        return linha
    return tuple(None if type(c) is str and c in ERROS_DO_EXCEL else c for c in linha)


def como_o_openpyxl(celula: object) -> object:
    """Uma célula do calamine na forma que o openpyxl entregava.

    AS DUAS DIFERENÇAS, resolvidas aqui para que nada rio abaixo precise saber
    qual leitor trouxe a célula:

    * VAZIO vem como `""`, e não `None` — `linha.get(coluna) is None` em
      qualquer lugar passaria a ser falso;
    * INTEIRO vem como `float`, e a coluna `ID` da Clipei é numérica:
      `id_fonte` gravaria `"12345.0"` no lugar de `"12345"`, e é por esse campo
      que se compara matéria com matéria.
    """
    if celula == "":
        return None
    if type(celula) is float and celula.is_integer():
        return int(celula)
    return celula


def celulas_pelo_calamine(conteudo: bytes, aba: str | None = None) -> list[tuple]:
    """A aba lida pelo calamine, com as células na forma do openpyxl.

    `skip_empty_area=False` porque `to_python()` PULA as linhas vazias do topo
    por padrão: o leitor de menções procura o cabeçalho nas 20 primeiras linhas
    e o de bonds lê as cinco do topo por índice, e pular uma mudaria o que os
    dois acham.

    LEVANTA `RegraViolada` para o que é recado à pessoa (aba que falta, arquivo
    vazio). Qualquer outra exceção é falha DESTE leitor, e quem chama pode
    tentar o openpyxl.
    """
    from python_calamine import CalamineWorkbook

    livro = CalamineWorkbook.from_filelike(io.BytesIO(conteudo))
    nomes = livro.sheet_names
    if aba and aba not in nomes:
        raise RegraViolada(
            f"A planilha não tem a aba {aba!r}. Abas encontradas: {', '.join(nomes)}."
        )
    if not nomes:
        raise RegraViolada("A planilha está vazia.")
    folha = livro.get_sheet_by_name(aba or nomes[0])
    return [
        tuple(como_o_openpyxl(celula) for celula in linha)
        for linha in folha.to_python(skip_empty_area=False)
    ]


def abrir_pelo_openpyxl(conteudo: bytes, aba: str | None = None):
    """Abre o arquivo pelo leitor antigo e devolve `(planilha, folha)`.

    QUEM ABRE FECHA: a planilha volta junto porque `read_only` segura o ZIP, e
    quem chama precisa de `planilha.close()` num `finally`.

    É O CAMINHO DE RESERVA, e o único que materializa a dimensão declarada (ver
    o cabeçalho). Quem o usa para guardar a leitura inteira deve medir antes.
    """
    try:
        from openpyxl import load_workbook
    except ModuleNotFoundError as erro:  # pragma: no cover - dependência declarada
        raise RegraViolada("Leitura de planilha indisponível neste servidor.") from erro

    # `read_only` não carrega a planilha inteira na memória de uma vez, e
    # `data_only` traz o VALOR de uma célula com fórmula — sem ele, a coluna de
    # engajamento calculada chegaria como a string "=SOMA(...)".
    try:
        planilha = load_workbook(io.BytesIO(conteudo), read_only=True, data_only=True)
    except Exception as erro:
        # Um `.xls` antigo, um CSV renomeado ou um upload truncado chegam aqui.
        # Sem a tradução, o openpyxl devolveria 500 e a pessoa leria "erro
        # interno" para um arquivo que ela mesma pode trocar.
        raise RegraViolada(
            "Não consegui abrir o arquivo como planilha .xlsx. "
            "Se ele veio em .xls ou .csv, salve como .xlsx e envie de novo."
        ) from erro

    if aba and aba in planilha.sheetnames:
        return planilha, planilha[aba]
    if aba:
        abas = ", ".join(planilha.sheetnames)
        planilha.close()
        raise RegraViolada(f"A planilha não tem a aba {aba!r}. Abas encontradas: {abas}.")
    return planilha, planilha[planilha.sheetnames[0]]
